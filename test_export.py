"""Export integrity tests using real files and a non-Origin fake binding."""

from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import spectra_to_origin as sto


def _write_spectrum(path: Path, rows: list[tuple[str, str]], header: str | None = None) -> Path:
    lines = ([f"# {header}"] if header is not None else [])
    lines.extend(f"{x}\t{y}" for x, y in rows)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


class FakeGraphLayer:
    def __init__(self) -> None:
        self.obj = type("Data", (), {"DataPlots": []})()

    def add_plot(self, *_args, **_kwargs):
        self.obj.DataPlots.append(object())
        return object()

    def axis(self, _axis: str):
        return type("Axis", (), {"title": ""})()

    def rescale(self) -> None:
        pass

    def lt_exec(self, _command: str) -> None:
        pass


class FakeGraph:
    def __init__(self) -> None:
        self.layer = FakeGraphLayer()

    def __getitem__(self, _index: int):
        return self.layer


class FakeSheet:
    def __init__(self, name: str = "") -> None:
        self.name = name
        self.lname = ""

    def from_list2(self, *_args) -> None:
        pass

    def set_labels(self, *_args) -> None:
        pass

    def cols_axis(self, *_args) -> None:
        pass


class FakeBook:
    def __init__(self) -> None:
        self.lname = ""
        self.sheets = [FakeSheet()]

    def __getitem__(self, index: int):
        return self.sheets[index]

    def add_sheet(self, name: str):
        sheet = FakeSheet(name)
        self.sheets.append(sheet)
        return sheet


class FakeOrigin:
    def __init__(self, save_result=False) -> None:
        self.book = FakeBook()
        self.save_result = save_result
        self.exit_calls = 0
        self.save_path: Path | None = None
        self.opened_path: str | None = None

    def set_show(self, _show: bool) -> None:
        pass

    def new(self, _new: bool) -> None:
        pass

    def pages(self, page_type: str):
        return [self.book] if page_type == "w" else []

    def new_graph(self, **_kwargs):
        return FakeGraph()

    def save(self, _path: str):
        self.save_path = Path(_path)
        if self.save_result:
            Path(_path).write_bytes(b"fresh-opju")
        return self.save_result

    def open(self, path: str):
        self.opened_path = path
        return True

    def exit(self) -> None:
        self.exit_calls += 1


class ExportIntegrityTests(unittest.TestCase):
    def test_decimal_equal_x_formats_share_grid_without_tolerance(self) -> None:
        first = sto.Spectrum(Path("a.txt"), ("1", "2.50", "3e0"), ("1", "2", "3"), "a", "", None)
        equivalent = sto.Spectrum(Path("b.txt"), ("1.0", "2.500", "3.000"), ("4", "5", "6"), "b", "", None)
        different = sto.Spectrum(Path("c.txt"), ("1.0000000000000001", "2.5", "3"), ("7", "8", "9"), "c", "", None)
        self.assertTrue(sto.shared_x_grid([first, equivalent]))
        self.assertFalse(sto.shared_x_grid([first, different]))
        self.assertEqual(sto.infer_layout([first, equivalent]), "XYYY")
        self.assertEqual(sto.infer_layout([first, different]), "XYXY")

    def test_auto_layout_is_inferred_independently_for_each_group(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            files = [
                _write_spectrum(root / "same_a.txt", [("1", "2"), ("2", "3")]),
                _write_spectrum(root / "same_b.txt", [("1.0", "4"), ("2.0", "5")]),
                _write_spectrum(root / "other_a.txt", [("1", "6"), ("2", "7")]),
                _write_spectrum(root / "other_b.txt", [("10", "8"), ("20", "9")]),
            ]
            spectra = sto.parse_spectra(files)
            groups = [("same", spectra[:2]), ("other", spectra[2:])]
            specs = sto.build_sheet_specs(groups, "auto")
            self.assertEqual([spec.layout for spec in specs], ["XYYY", "XYXY"])
            with self.assertRaisesRegex(ValueError, "X 网格不一致"):
                sto.build_sheet_specs(groups, "XYYY")

    def test_explicit_xyyy_mismatch_fails_before_origin_or_output(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            a = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            b = _write_spectrum(root / "b.txt", [("10", "4"), ("20", "5")])
            output = root / "out.opju"
            with mock.patch.object(sto, "origin_process_running") as process_check, mock.patch.object(
                sto, "_load_originpro", side_effect=AssertionError("Origin must not be loaded")
            ) as load_origin:
                with self.assertRaisesRegex(ValueError, "X 网格不一致"):
                    sto.export_origin_project([a, b], output, layout="XYYY")
            process_check.assert_not_called()
            load_origin.assert_not_called()
            self.assertFalse(output.exists())

    def test_group_exports_round_trip_to_xlsx_and_every_csv_with_safe_names(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            a = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            b = _write_spectrum(root / "b.txt", [("10", "4"), ("20", "5")])
            spectra = sto.parse_spectra([a, b])
            groups = [("../unsafe", [spectra[0]]), ("../unsafe", [spectra[1]])]
            labels = sto.AxisLabels("Q", "nm^-1", "Counts", "cts")

            xlsx, csv_result = sto.export_spectra(
                [a, b], root / "series.xlsx", groups=groups, axis_labels=labels
            )

            self.assertEqual(csv_result, root / "series_csv")
            csv_paths = sorted(csv_result.glob("*.csv"))
            self.assertEqual(len(csv_paths), 2)
            self.assertEqual(len({path.name.casefold() for path in csv_paths}), 2)
            self.assertTrue(all(path.parent == csv_result for path in csv_paths))
            for path in csv_paths:
                self.assertNotIn("..", path.name)
            with csv_paths[0].open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[0][0], "Q")
            self.assertEqual(rows[1][0], "nm^-1")

            book = load_workbook(xlsx, data_only=False)
            self.assertEqual(book.sheetnames, [".._unsafe", ".._unsafe_2", "readme"])
            self.assertEqual(book[book.sheetnames[0]]["A1"].value, "Q")
            self.assertEqual(book[book.sheetnames[0]]["A4"].value, 1)
            self.assertEqual(book["readme"]["B6"].value, "Q (nm^-1)")
            book.close()

    def test_metadata_that_looks_like_formula_is_stored_as_text(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "formula.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.Spectrum(source, ("1", "2"), ("2", "3"), "=2+2", "=1+1", None)
            xlsx, _csv = sto.export_spectra(
                [source], root / "safe.xlsx", groups=[("spectra", [spectrum])], write_csv_copy=False
            )
            book = load_workbook(xlsx, data_only=False)
            sheet = book["spectra"]
            self.assertEqual(sheet["B1"].value, "=2+2")
            self.assertEqual(sheet["B1"].data_type, "s")
            self.assertEqual(sheet["B3"].value, "=1+1")
            self.assertEqual(sheet["B3"].data_type, "s")
            book.close()

    def test_xlsx_csv_collision_with_input_fails_before_writing(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            input_csv = _write_spectrum(root / "out.csv", [("1", "2"), ("2", "3")])
            xlsx = root / "out.xlsx"
            before = input_csv.read_bytes()
            with self.assertRaisesRegex(ValueError, "输入文件冲突"):
                sto.export_spectra([input_csv], xlsx)
            self.assertFalse(xlsx.exists())
            self.assertEqual(input_csv.read_bytes(), before)

    def test_csv_save_dialog_suffix_produces_distinct_xlsx_and_csv(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "scan.txt", [("1", "2"), ("2", "3")])
            xlsx, csv_path = sto.export_spectra([source], root / "result.csv")
            self.assertEqual(xlsx, root / "result.xlsx")
            self.assertEqual(csv_path, root / "result.csv")
            self.assertNotEqual(xlsx, csv_path)
            workbook = load_workbook(xlsx, data_only=True)
            self.assertEqual(workbook["spectra"]["A4"].value, 1)
            workbook.close()
            self.assertTrue(csv_path.is_file())

    def test_failed_xlsx_write_keeps_previous_file_and_cleans_temporary_file(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            output = root / "out.xlsx"
            output.write_bytes(b"old-workbook")
            with mock.patch.object(sto.Workbook, "save", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(OSError, "disk full"):
                    sto.export_spectra([source], output, write_csv_copy=False)
            self.assertEqual(output.read_bytes(), b"old-workbook")
            self.assertEqual(list(root.glob(".out_*.tmp.xlsx")), [])

    def test_failed_csv_write_keeps_previous_sidecar(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            xlsx = root / "out.xlsx"
            csv_path = root / "out.csv"
            csv_path.write_bytes(b"old-csv")
            with mock.patch.object(sto, "_write_csv_contents", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(OSError, "disk full"):
                    sto.export_spectra([source], xlsx)
            self.assertTrue(xlsx.is_file())
            self.assertEqual(csv_path.read_bytes(), b"old-csv")
            self.assertEqual(list(root.glob(".out_*.tmp.csv")), [])

    def test_multi_csv_replace_removes_old_groups_and_refuses_user_files(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            sources = [
                _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")]),
                _write_spectrum(root / "b.txt", [("3", "4"), ("4", "5")]),
            ]
            spectra = sto.parse_spectra(sources)
            _, csv_dir = sto.export_spectra(
                sources, root / "series.xlsx", groups=[("old-a", [spectra[0]]), ("old-b", [spectra[1]])]
            )
            self.assertEqual({path.name for path in csv_dir.glob("*.csv")}, {"old-a.csv", "old-b.csv"})
            _, csv_dir = sto.export_spectra(
                sources, root / "series.xlsx", groups=[("new", [spectra[0]]), ("new", [spectra[1]])]
            )
            self.assertEqual({path.name for path in csv_dir.glob("*.csv")}, {"new.csv", "new_2.csv"})
            unrelated = csv_dir / "notes.txt"
            unrelated.write_text("keep", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "清单之外"):
                sto.export_spectra(
                    sources, root / "series.xlsx", groups=[("latest-a", [spectra[0]]), ("latest-b", [spectra[1]])]
                )
            self.assertEqual(unrelated.read_text(encoding="utf-8"), "keep")
            self.assertEqual({path.name for path in csv_dir.glob("*.csv")}, {"new.csv", "new_2.csv"})

    def test_multi_csv_directory_may_not_contain_any_input(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            target = root / "series_csv"
            target.mkdir()
            source = _write_spectrum(target / "not-a-group.csv", [("1", "2"), ("2", "3")])
            other = _write_spectrum(root / "b.txt", [("3", "4"), ("4", "5")])
            spectra = sto.parse_spectra([source, other])
            with self.assertRaisesRegex(ValueError, "包含输入文件"):
                sto.export_spectra(
                    [source, other],
                    root / "series.xlsx",
                    groups=[("one", [spectra[0]]), ("two", [spectra[1]])],
                )
            self.assertTrue(source.exists())
            self.assertFalse((root / "series.xlsx").exists())

    def test_failed_multi_csv_install_restores_previous_directory(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            sources = [
                _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")]),
                _write_spectrum(root / "b.txt", [("3", "4"), ("4", "5")]),
            ]
            spectra = sto.parse_spectra(sources)
            groups = [("old-a", [spectra[0]]), ("old-b", [spectra[1]])]
            _, csv_dir = sto.export_spectra(sources, root / "series.xlsx", groups=groups)
            previous = {path.name: path.read_bytes() for path in csv_dir.glob("*.csv")}
            real_replace = sto.os.replace

            def fail_stage_install(source: str | Path, destination: str | Path) -> None:
                if Path(source).name.startswith(".series_csv_stage_") and Path(destination) == csv_dir:
                    raise OSError("cannot install staged CSV directory")
                real_replace(source, destination)

            with mock.patch.object(sto.os, "replace", side_effect=fail_stage_install):
                with self.assertRaisesRegex(OSError, "cannot install staged"):
                    sto.export_spectra(
                        sources,
                        root / "series.xlsx",
                        groups=[("new-a", [spectra[0]]), ("new-b", [spectra[1]])],
                    )
            self.assertEqual({path.name: path.read_bytes() for path in csv_dir.glob("*.csv")}, previous)
            self.assertEqual(list(root.glob(".series_csv_stage_*")), [])

    def test_failed_multi_csv_rollback_preserves_backup_and_reports_path(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            sources = [
                _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")]),
                _write_spectrum(root / "b.txt", [("3", "4"), ("4", "5")]),
            ]
            spectra = sto.parse_spectra(sources)
            _, csv_dir = sto.export_spectra(
                sources,
                root / "series.xlsx",
                groups=[("old-a", [spectra[0]]), ("old-b", [spectra[1]])],
            )
            old_files = {path.name: path.read_bytes() for path in csv_dir.glob("*.csv")}
            real_replace = sto.os.replace

            def fail_install_and_restore(source: str | Path, destination: str | Path) -> None:
                source_name = Path(source).name
                if Path(destination) == csv_dir and (
                    source_name.startswith(".series_csv_stage_")
                    or source_name.startswith(".series_csv_backup_")
                ):
                    raise OSError("replace blocked")
                real_replace(source, destination)

            with mock.patch.object(sto.os, "replace", side_effect=fail_install_and_restore):
                with self.assertRaisesRegex(sto.OriginExportError, "旧目录保留在") as caught:
                    sto.export_spectra(
                        sources,
                        root / "series.xlsx",
                        groups=[("new-a", [spectra[0]]), ("new-b", [spectra[1]])],
                    )
            message = str(caught.exception)
            backup = Path(message.split("旧目录保留在 ", 1)[1].split("；", 1)[0])
            self.assertTrue(backup.is_dir())
            self.assertEqual({path.name: path.read_bytes() for path in backup.glob("*.csv")}, old_files)

    def test_user_file_added_while_staging_is_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            sources = [
                _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")]),
                _write_spectrum(root / "b.txt", [("3", "4"), ("4", "5")]),
            ]
            spectra = sto.parse_spectra(sources)
            _, csv_dir = sto.export_spectra(
                sources,
                root / "series.xlsx",
                groups=[("old-a", [spectra[0]]), ("old-b", [spectra[1]])],
            )
            old_files = {path.name: path.read_bytes() for path in csv_dir.glob("*.csv")}
            write_csv = sto._write_csv_contents
            injected = False

            def add_user_file_during_staging(path: Path, table: sto.SheetTable) -> None:
                nonlocal injected
                if not injected:
                    (csv_dir / "notes.txt").write_text("keep", encoding="utf-8")
                    injected = True
                write_csv(path, table)

            with mock.patch.object(sto, "_write_csv_contents", side_effect=add_user_file_during_staging):
                with self.assertRaisesRegex(ValueError, "清单之外"):
                    sto.export_spectra(
                        sources,
                        root / "series.xlsx",
                        groups=[("new-a", [spectra[0]]), ("new-b", [spectra[1]])],
                    )
            self.assertEqual((csv_dir / "notes.txt").read_text(encoding="utf-8"), "keep")
            self.assertEqual({path.name: path.read_bytes() for path in csv_dir.glob("*.csv")}, old_files)
            self.assertEqual(list(root.glob(".series_csv_backup_*")), [])

    def test_origin_running_is_refused_before_binding_or_file_write(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.parse_spectrum(source)
            target = root / "out.opju"
            with mock.patch.object(sto, "origin_process_running", return_value=True), mock.patch.object(
                sto, "_load_originpro", side_effect=AssertionError("must not attach to user Origin")
            ) as load_origin:
                with self.assertRaisesRegex(sto.OriginExportError, "Origin 已运行"):
                    sto.write_origin_project([spectrum], target)
            load_origin.assert_not_called()
            self.assertFalse(target.exists())

    def test_fake_origin_save_failure_preserves_old_opju_and_closes_session(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.parse_spectrum(source)
            target = root / "out.opju"
            target.write_bytes(b"previous-project")
            origin = FakeOrigin(save_result=False)
            previous_tempdir = tempfile.tempdir
            tempfile.tempdir = str(root)
            try:
                with mock.patch.object(sto, "origin_process_running", return_value=False), mock.patch.object(
                    sto, "_load_originpro", return_value=origin
                ):
                    with self.assertRaisesRegex(sto.OriginExportError, r"没有写出 \.opju 文件"):
                        sto.write_origin_project([spectrum], target, keep_open=True)
            finally:
                tempfile.tempdir = previous_tempdir
            self.assertEqual(target.read_bytes(), b"previous-project")
            self.assertEqual(origin.exit_calls, 1)
            self.assertEqual(list(root.glob("spectra_opju_*")), [])

    def test_fake_origin_success_replaces_old_opju_atomically(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.parse_spectrum(source)
            target = root / "out.opju"
            target.write_bytes(b"previous-project")
            origin = FakeOrigin(save_result=True)
            previous_tempdir = tempfile.tempdir
            tempfile.tempdir = str(root)
            try:
                with mock.patch.object(sto, "origin_process_running", return_value=False), mock.patch.object(
                    sto, "_load_originpro", return_value=origin
                ):
                    result = sto.write_origin_project([spectrum], target, keep_open=True)
            finally:
                tempfile.tempdir = previous_tempdir
            self.assertEqual(result, target)
            self.assertEqual(target.read_bytes(), b"fresh-opju")
            self.assertEqual(list(root.glob("spectra_opju_*")), [])

    def test_locked_origin_temp_source_does_not_report_failure_and_is_cleaned_after_exit(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.parse_spectrum(source)
            target = root / "out.opju"
            target.write_bytes(b"previous-project")
            origin = FakeOrigin(save_result=True)
            real_unlink = sto.Path.unlink
            previous_tempdir = tempfile.tempdir
            tempfile.tempdir = str(root)

            def unlink_after_origin_closes(path: Path, *args, **kwargs):
                if path == origin.save_path and origin.exit_calls == 0:
                    raise PermissionError("Origin holds the temporary save open")
                return real_unlink(path, *args, **kwargs)

            try:
                with mock.patch.object(sto, "origin_process_running", return_value=False), mock.patch.object(
                    sto, "_load_originpro", return_value=origin
                ), mock.patch.object(sto.Path, "unlink", autospec=True, side_effect=unlink_after_origin_closes):
                    result = sto.write_origin_project([spectrum], target)
            finally:
                tempfile.tempdir = previous_tempdir
            self.assertEqual(result, target)
            self.assertEqual(target.read_bytes(), b"fresh-opju")
            self.assertEqual(origin.exit_calls, 1)
            self.assertFalse(origin.save_path.exists())
            self.assertEqual(list(root.glob("spectra_opju_*")), [])

    def test_keep_open_switches_to_final_project_path(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            spectrum = sto.parse_spectrum(source)
            target = root / "out.opju"
            origin = FakeOrigin(save_result=True)
            previous_tempdir = tempfile.tempdir
            tempfile.tempdir = str(root)
            try:
                with mock.patch.object(sto, "origin_process_running", return_value=False), mock.patch.object(
                    sto, "_load_originpro", return_value=origin
                ):
                    sto.write_origin_project([spectrum], target, keep_open=True)
            finally:
                tempfile.tempdir = previous_tempdir
            self.assertEqual(origin.opened_path, str(target))

    def test_excel_dimension_limit_is_checked_before_creating_output(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw).resolve()
            source = _write_spectrum(root / "a.txt", [("1", "2"), ("2", "3")])
            target = root / "too_many_rows.xlsx"
            with mock.patch.object(sto, "EXCEL_MAX_ROWS", 4):
                with self.assertRaisesRegex(ValueError, "超过 Excel 上限"):
                    sto.export_spectra([source], target, write_csv_copy=False)
            self.assertFalse(target.exists())


if __name__ == "__main__":
    unittest.main()
