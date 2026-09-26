"""Unit tests for the versioned generic-import contract; no Origin process is used."""

from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from origin_bridge import planning
from origin_bridge.models import DataColumn, DataImportError, DataTable


READ_OPTIONS = {
    "sheet": None,
    "header": True,
    "skip_rows": 0,
    "delimiter": "whitespace",
    "encoding": "utf-8",
    "missing_values": ["NA"],
    "formula_policy": "cached",
}


class GenericPlanningTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.source = self.root / "measurements.csv"
        self.source.write_text("placeholder input", encoding="utf-8")
        self.digest = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.output = self.root / "result.opju"

    def table(
        self,
        columns: tuple[DataColumn, ...] | None = None,
        *,
        name: str = "measurements",
        sheet: str | None = None,
        options: dict | None = None,
        source_hash: str | None = None,
        source: Path | None = None,
    ) -> DataTable:
        if columns is None:
            columns = (
                DataColumn("Time", "number", ("0", "1", "2"), "s"),
                DataColumn("Signal", "number", ("10", None, "12"), "count"),
                DataColumn("Error", "number", ("0.2", "0", "0.4"), "count"),
                DataColumn("Sample", "text", ("A", "B", "C")),
            )
        replay = dict(READ_OPTIONS)
        replay["sheet"] = sheet
        if options is not None:
            replay = options
        return DataTable(
            source=source or self.source,
            source_sheet=sheet,
            name=name,
            columns=columns,
            source_hash=source_hash or self.digest,
            read_options=replay,
        )

    @contextmanager
    def reader(self, *tables: DataTable):
        with mock.patch.object(planning, "_discover_files", return_value=[self.source]) as discover, mock.patch.object(
            planning, "_read_tables", return_value=list(tables)
        ) as read:
            yield discover, read

    def test_inspection_reports_types_missing_ranges_and_plot_suggestion(self):
        table = self.table()
        with self.reader(table) as (discover, read):
            result = planning.inspect_inputs([self.source])

        self.assertEqual(result["schema_version"], 1)
        self.assertEqual(len(result["tables"]), 1)
        inspected = result["tables"][0]
        self.assertEqual(inspected["source"]["sha256"], self.digest)
        self.assertEqual(inspected["source"]["options"], READ_OPTIONS)
        self.assertEqual(inspected["n_rows"], 3)
        self.assertEqual(inspected["columns"][1]["kind"], "number")
        self.assertEqual(inspected["columns"][1]["missing"], 1)
        self.assertEqual(inspected["columns"][1]["sample"], ["10", None, "12"])
        self.assertEqual(inspected["columns"][1]["range"], {"min": "10", "max": "12"})
        self.assertEqual(inspected["suggested_plot"]["kind"], "line")
        self.assertEqual(inspected["suggested_plot"]["x"], 0)
        self.assertEqual(inspected["suggested_plot"]["y"], [1])
        self.assertEqual(inspected["suggested_plot"]["y_error"], {"Signal": 2})
        json.dumps(result)
        discover.assert_called_once()
        read.assert_called_once()

    def test_create_plan_binds_each_workbook_sheet_with_actual_read_options(self):
        workbook = self.root / "book.xlsx"
        workbook.write_bytes(b"fake workbook")
        digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
        options_a = dict(READ_OPTIONS, sheet="Results")
        options_b = dict(READ_OPTIONS, sheet="Summary")
        columns = (DataColumn("Category", "text", ("A", "B")), DataColumn("Value", "number", ("1", "2")))
        tables = [
            self.table(columns, name="book", source=workbook, source_hash=digest, sheet="Results", options=options_a),
            self.table(columns, name="book", source=workbook, source_hash=digest, sheet="Summary", options=options_b),
        ]
        with mock.patch.object(planning, "_discover_files", return_value=[workbook]), mock.patch.object(
            planning, "_read_tables", return_value=tables
        ):
            plan = planning.create_plan([workbook], self.root / "result", format="opju")

        self.assertEqual(plan["schema_version"], 1)
        self.assertEqual(plan["output"]["path"], str(self.output))
        self.assertEqual([item["source"]["options"]["sheet"] for item in plan["tables"]], ["Results", "Summary"])
        self.assertEqual([item["name"] for item in plan["tables"]], ["book", "book_2"])
        self.assertEqual([item["plot"]["kind"] for item in plan["tables"]], ["column", "column"])
        self.assertFalse(self.output.exists())
        json.dumps(plan)

    def test_real_multi_sheet_workbook_round_trips_through_plan_without_export(self):
        workbook_path = self.root / "real.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "Temperature"
        first.append(["Time (s)", "Temperature (K)"])
        first.append([0, 298.15])
        first.append([1, 299.15])
        second = workbook.create_sheet("Voltage")
        second.append(["Time (s)", "Voltage (V)"])
        second.append([0, 1.2])
        second.append([1, 1.3])
        workbook.save(workbook_path)

        inspection = planning.inspect_inputs([workbook_path])
        self.assertEqual([table["source"]["sheet"] for table in inspection["tables"]], ["Temperature", "Voltage"])
        plan = planning.create_plan([workbook_path], self.output)
        self.assertEqual([table["source"]["options"]["sheet"] for table in plan["tables"]], ["Temperature", "Voltage"])
        prepared = planning.prepare_plan(plan)
        self.assertEqual([table.table.source_sheet for table in prepared.tables], ["Temperature", "Voltage"])
        self.assertFalse(self.output.exists())

    def valid_plan(self, *, table: DataTable | None = None) -> dict:
        source_table = table or self.table()
        return {
            "schema_version": 1,
            "output": {
                "path": str(self.output),
                "format": "opju",
                "overwrite": False,
                "keep_open": False,
            },
            "tables": [
                {
                    "source": {
                        "path": str(source_table.source),
                        "sha256": source_table.source_hash,
                        "options": source_table.read_options,
                    },
                    "name": "Processed",
                    "plot": {
                        "kind": "line",
                        "x": "Time",
                        "y": ["Signal"],
                        "y_error": {"Signal": "Error"},
                        "title": "Signal over time",
                        "x_label": "Time (s)",
                        "y_label": "Signal (count)",
                    },
                    "column_labels": {
                        "Time": {"name": "Elapsed", "unit": "s"},
                        "Signal": {"name": "Counts", "unit": "count"},
                    },
                }
            ],
        }

    def test_prepare_plan_resolves_selectors_and_keeps_all_typed_columns(self):
        table = self.table()
        with mock.patch.object(planning, "_read_tables", return_value=[table]):
            prepared = planning.prepare_plan(self.valid_plan())

        self.assertEqual(len(prepared.tables), 1)
        prepared_table = prepared.tables[0].table
        self.assertEqual(len(prepared_table.columns), 4)
        self.assertEqual(prepared_table.columns[0].name, "Elapsed")
        self.assertEqual(prepared_table.columns[1].name, "Counts")
        self.assertEqual(prepared_table.columns[1].original_name, "Signal")
        self.assertEqual(prepared_table.columns[1].kind, "number")
        self.assertEqual(prepared_table.columns[3].kind, "text")
        self.assertEqual(prepared.tables[0].plot.x, 0)
        self.assertEqual(prepared.tables[0].plot.y, (1,))
        self.assertEqual(prepared.tables[0].plot.y_error, ((1, 2),))

        description = planning.describe_prepared(prepared)
        self.assertTrue(description["valid"])
        self.assertEqual(description["tables"][0]["columns"][1]["source_name"], "Signal")
        self.assertEqual(description["tables"][0]["columns"][1]["name"], "Counts")
        self.assertEqual(description["graphs"][0]["y"], [1])
        json.dumps(description)

    def test_relative_paths_resolve_from_base_dir(self):
        plan = self.valid_plan()
        plan["output"]["path"] = "result.opju"
        plan["tables"][0]["source"]["path"] = "measurements.csv"
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            prepared = planning.prepare_plan(plan, base_dir=self.root)
        self.assertEqual(prepared.output, self.output)

    def test_plan_rejects_unknown_fields_and_versions(self):
        plan = self.valid_plan()
        plan["surprise"] = True
        with self.assertRaisesRegex(DataImportError, "Unknown plan field"):
            planning.prepare_plan(plan)

        plan = self.valid_plan()
        plan["schema_version"] = 2
        with self.assertRaisesRegex(DataImportError, "Unsupported import plan schema_version"):
            planning.prepare_plan(plan)

        plan = self.valid_plan()
        plan["tables"][0]["source"]["options"]["command"] = "run arbitrary code"
        with self.assertRaisesRegex(DataImportError, "Unknown source.options field"):
            planning.prepare_plan(plan)

    def test_plan_rejects_source_hash_drift_and_never_writes_output(self):
        plan = self.valid_plan()
        table = self.table(source_hash="0" * 64)
        with mock.patch.object(planning, "_read_tables", return_value=[table]):
            with self.assertRaisesRegex(DataImportError, "Source changed since inspection"):
                planning.prepare_plan(plan)
        self.assertFalse(self.output.exists())

    def test_plan_rejects_overwrite_by_default_and_accepts_explicit_overwrite(self):
        self.output.write_bytes(b"preserve")
        plan = self.valid_plan()
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "overwrite is false"):
                planning.prepare_plan(plan)
            plan["output"]["overwrite"] = True
            prepared = planning.prepare_plan(plan)
        self.assertTrue(prepared.overwrite)
        self.assertEqual(self.output.read_bytes(), b"preserve")

    def test_plan_rejects_input_output_conflict(self):
        source = self.root / "conflict.opju"
        source.write_text("source", encoding="utf-8")
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        table = self.table(source=source, source_hash=digest)
        plan = self.valid_plan(table=table)
        plan["output"]["path"] = str(source)
        with mock.patch.object(planning, "_read_tables", return_value=[table]):
            with self.assertRaisesRegex(DataImportError, "conflicts with an input"):
                planning.prepare_plan(plan)

    def test_multi_sheet_plan_requires_explicit_sheet_names(self):
        workbook = self.root / "book.xlsx"
        workbook.write_bytes(b"fake workbook")
        digest = hashlib.sha256(workbook.read_bytes()).hexdigest()
        table = self.table(source=workbook, source_hash=digest)
        plan = self.valid_plan(table=table)
        second = dict(plan["tables"][0])
        second["name"] = "Processed 2"
        plan["tables"].append(second)
        with mock.patch.object(planning, "_read_tables", return_value=[table]):
            with self.assertRaisesRegex(DataImportError, "sheet must identify each worksheet"):
                planning.prepare_plan(plan)

    def test_plan_rejects_non_numeric_y_categorical_line_x_and_negative_error(self):
        table = self.table()
        invalid_plans = []

        plan = self.valid_plan()
        plan["tables"][0]["plot"]["y"] = ["Sample"]
        plan["tables"][0]["plot"]["y_error"] = {}
        invalid_plans.append((plan, "must be numeric"))

        plan = self.valid_plan()
        plan["tables"][0]["plot"]["x"] = "Sample"
        plan["tables"][0]["plot"]["y_error"] = {}
        invalid_plans.append((plan, "requires a numeric or datetime X"))

        negative_columns = list(table.columns)
        negative_columns[2] = DataColumn("Error", "number", ("0.2", "-0.1", "0.4"), "count")
        negative_table = self.table(tuple(negative_columns))
        plan = self.valid_plan(table=negative_table)
        invalid_plans.append((plan, "negative value"))

        for plan, expected in invalid_plans:
            with self.subTest(expected=expected), mock.patch.object(planning, "_read_tables", return_value=[
                negative_table if expected == "negative value" else table
            ]):
                with self.assertRaisesRegex(DataImportError, expected):
                    planning.prepare_plan(plan)

    def test_y_error_keys_are_source_y_names_and_error_columns_cannot_overlap_y(self):
        plan = self.valid_plan()
        plan["tables"][0]["plot"]["y_error"] = {"1": "Error"}
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "must name one of the selected Y columns"):
                planning.prepare_plan(plan)

        plan = self.valid_plan()
        plan["tables"][0]["plot"]["y_error"] = {"Signal": "Sample"}
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "must be numeric"):
                planning.prepare_plan(plan)

    def test_output_column_labels_must_be_unique(self):
        plan = self.valid_plan()
        plan["tables"][0]["column_labels"]["Error"] = {"name": "Counts", "unit": "count"}
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "must be unique"):
                planning.prepare_plan(plan)

    def test_plot_suggestions_cover_category_row_index_and_text_only_tables(self):
        category = self.table(
            (DataColumn("Condition", "text", ("A", "B", "A")), DataColumn("Value", "number", ("1", "2", "3")))
        )
        row_index = self.table((DataColumn("Value", "number", ("1", "2", "3")),))
        text_only = self.table((DataColumn("Label", "text", ("a", "b", "c")),))
        with mock.patch.object(planning, "_discover_files", return_value=[self.source]), mock.patch.object(
            planning, "_read_tables", side_effect=[[category], [row_index], [text_only]]
        ):
            self.assertEqual(planning.inspect_inputs([self.source])["tables"][0]["suggested_plot"]["kind"], "column")
            self.assertIsNone(planning.inspect_inputs([self.source])["tables"][0]["suggested_plot"]["x"])
            self.assertEqual(planning.inspect_inputs([self.source])["tables"][0]["suggested_plot"]["kind"], "none")

    def test_mixed_y_units_use_neutral_label_and_warning(self):
        mixed = self.table(
            (
                DataColumn("Time", "number", ("0", "1"), "s"),
                DataColumn("Temperature", "number", ("300", "301"), "K"),
                DataColumn("Voltage", "number", ("1", "2"), "V"),
            )
        )
        with self.reader(mixed):
            inspected = planning.inspect_inputs([self.source])["tables"][0]
        self.assertEqual(inspected["suggested_plot"]["y_label"], "Value")
        self.assertTrue(any("mixed units" in warning for warning in inspected["warnings"]))

    def test_auto_error_pairing_requires_matching_units_and_explicit_mismatch_fails(self):
        mismatched = self.table(
            (
                DataColumn("Strain", "number", ("0", "1"), ""),
                DataColumn("Stress", "number", ("10", "20"), "MPa"),
                DataColumn("Error(kPa)", "number", ("1", "2"), "kPa"),
            )
        )
        with self.reader(mismatched):
            suggested = planning.inspect_inputs([self.source])["tables"][0]
        self.assertEqual(suggested["suggested_plot"]["y"], [1])
        self.assertEqual(suggested["suggested_plot"]["y_error"], {})
        self.assertTrue(any("was not assigned as Y error" in warning for warning in suggested["warnings"]))

        plan = self.valid_plan(table=mismatched)
        plan["tables"][0]["plot"]["x"] = "Strain"
        plan["tables"][0]["plot"]["y"] = ["Stress"]
        plan["tables"][0]["plot"]["y_error"] = {"Stress": "Error(kPa)"}
        plan["tables"][0]["column_labels"] = {}
        with mock.patch.object(planning, "_read_tables", return_value=[mismatched]):
            with self.assertRaisesRegex(DataImportError, "does not match Y unit"):
                planning.prepare_plan(plan)

    def test_column_label_unit_mismatch_is_validated_without_conversion(self):
        plan = self.valid_plan()
        plan["tables"][0]["column_labels"]["Error"] = {"name": "Error", "unit": "kPa"}
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "no unit conversion is applied"):
                planning.prepare_plan(plan)

    def test_explicit_mixed_y_plot_reports_shared_axis_warning(self):
        mixed = self.table(
            (
                DataColumn("Time", "number", ("0", "1"), "s"),
                DataColumn("Temperature", "number", ("300", "301"), "K"),
                DataColumn("Voltage", "number", ("1", "2"), "V"),
            )
        )
        plan = self.valid_plan(table=mixed)
        plan["tables"][0]["plot"].update(x="Time", y=["Temperature", "Voltage"], y_error={})
        plan["tables"][0]["column_labels"] = {}
        with mock.patch.object(planning, "_read_tables", return_value=[mixed]):
            prepared = planning.prepare_plan(plan)
        self.assertTrue(any("mixed units" in warning for warning in prepared.warnings))

    def test_datetime_ranges_compare_timezone_instants_and_keep_source_text(self):
        first = "2024-01-01T00:00:00+14:00"
        second = "2023-12-31T23:00:00-12:00"
        dated = self.table(
            (
                DataColumn("Timestamp", "datetime", (first, second)),
                DataColumn("Value", "number", ("1", "2")),
            )
        )
        with self.reader(dated):
            inspected = planning.inspect_inputs([self.source])["tables"][0]
        self.assertEqual(inspected["columns"][0]["range"], {"min": first, "max": second})

    def test_reader_sheet_option_requires_a_sheet_name(self):
        plan = self.valid_plan()
        plan["tables"][0]["source"]["options"]["sheet"] = 0
        with mock.patch.object(planning, "_read_tables", return_value=[self.table()]):
            with self.assertRaisesRegex(DataImportError, "sheet must be null or a non-empty sheet name"):
                planning.prepare_plan(plan)


if __name__ == "__main__":
    unittest.main()
