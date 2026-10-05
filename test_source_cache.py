"""Content invalidation, bounded reuse and source verification regressions."""
from __future__ import annotations

import copy
import os
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from origin_bridge import cli, planning, source_cache
from origin_bridge.exporter import execute_import
from origin_bridge.models import DataImportError
from origin_bridge.source_cache import SourceCache


class SourceCacheTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.source = self.root / "data.tsv"
        self.source.write_bytes(b"Time\tValue\n0\t1\n1\t2\n")
        self.output = self.root / "out.xlsx"

    def parser(self):
        return mock.patch.object(source_cache, "_parse_tables", wraps=source_cache._parse_tables)

    def test_independent_operations_rehash_but_reuse_parsing_and_ranges(self):
        cache = SourceCache()
        reads = []
        original = Path.read_bytes

        def read(path):
            reads.append(path)
            return original(path)

        with self.parser() as parse, mock.patch.object(Path, "read_bytes", read), mock.patch.object(
            planning, "_column_range", wraps=planning._column_range
        ) as ranges, mock.patch.object(planning, "_validate_table_shape", wraps=planning._validate_table_shape) as shapes:
            first = planning.inspect_inputs([self.source], cache=cache)
            second = planning.inspect_inputs([self.source], cache=cache)
            plan = planning.create_plan([self.source], self.output, format="xlsx", cache=cache)
            self.assertEqual(shapes.call_count, 1, "unchanged summaries reuse source-shape validation")
            prepared = planning.prepare_plan(plan, cache=cache)
        self.assertEqual(first, second)
        self.assertEqual(parse.call_count, 1)
        self.assertEqual(ranges.call_count, 2)
        self.assertEqual(shapes.call_count, 2, "preparation still validates the source table")
        self.assertEqual(len(reads), 4, "each independent decision must verify fresh bytes")
        self.assertEqual(prepared.tables[0].table.columns[1].values, ("1", "2"))
        self.assertFalse(self.output.exists())

    def test_cli_import_reuses_its_plan_read(self):
        args = cli.build_parser().parse_args(["import", "-i", str(self.source), "-o", str(self.output)])
        with self.parser() as parse:
            receipt = cli._dispatch(args)
        self.assertEqual(parse.call_count, 1)
        self.assertTrue(receipt["ok"])
        self.assertTrue(self.output.exists())

    def test_cli_inspection_does_not_size_an_unused_cache(self):
        expected = {"ok": True, **planning.inspect_inputs([self.source])}
        args = cli.build_parser().parse_args(["inspect", "-i", str(self.source)])
        with self.parser() as parse, mock.patch.object(source_cache.sys, "getsizeof") as sizes:
            inspected = cli._dispatch(args)
        self.assertEqual(inspected, expected)
        sizes.assert_not_called()
        parse.assert_not_called()  # The ordinary reader parsed the source, without cache bookkeeping.

    def test_workbook_subset_aliases_and_per_sheet_options_match_uncached_replay(self):
        path = self.root / "varied.xlsx"
        book = Workbook()
        layouts = {
            "Header1": [["X", "Y"], [1, 10]],
            "Header2": [["X", "Y"], [2, 20]],
            "Values": [[3, 30], [4, 40]],
            "Skipped": [["metadata", "v2"], ["X", "Y"], [5, 50]],
        }
        for index, (name, rows) in enumerate(layouts.items()):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = name
            for row in rows:
                sheet.append(row)
        book.save(path)
        book.close()
        requests = (
            {"sheet": "Header1"}, {"sheet": "Header2"}, {"sheet": "Values", "header": False},
            {"sheet": "Skipped", "header": True, "skip_rows": 1, "missing_values": ["50"]},
        )
        plan = None
        for options in requests:
            part = planning.create_plan([path], self.output, format="xlsx", options=options)
            if plan is None:
                plan = part
            else:
                plan["tables"].extend(part["tables"])
        expected = planning.prepare_plan(plan)
        for warm in (False, True):
            with self.subTest(warm=warm), self.parser() as parse:
                cache = SourceCache()
                if warm:
                    cache.read(path)
                    subset = cache.read_sheets(path, ("Header2", "Header1"), {"header": True})
                    self.assertEqual([table.source_sheet for table in subset], ["Header2", "Header1"])
                    self.assertEqual(parse.call_count, 1, "selected sheets reuse their existing aliases")
                prepared = planning.prepare_plan(plan, cache=cache)
                self.assertEqual(prepared, expected)
                self.assertEqual(parse.call_count, 2 if warm else 3)
                self.assertEqual(prepared.tables[2].table.columns[0].values, ("3", "4"))
                self.assertEqual(prepared.tables[3].table.columns[1].values, (None,))

    def test_same_size_same_mtime_change_invalidates_and_rejects_old_plan(self):
        cache = SourceCache()
        stamp = self.source.stat()
        plan = planning.create_plan([self.source], self.output, format="xlsx", cache=cache)
        self.source.write_bytes(b"Time\tValue\n0\t3\n1\t4\n")
        os.utime(self.source, ns=(stamp.st_atime_ns, stamp.st_mtime_ns))
        self.assertEqual(self.source.stat().st_size, stamp.st_size)
        with self.assertRaisesRegex(DataImportError, "Source changed"):
            planning.prepare_plan(plan, cache=cache)
        summary = planning.inspect_inputs([self.source], cache=cache)
        self.assertEqual(summary["tables"][0]["columns"][1]["sample"], ["3", "4"])
        self.assertFalse(self.output.exists())

    def test_exporter_rechecks_sources_after_preparation(self):
        cache = SourceCache()
        with cache.operation():
            plan = planning.create_plan([self.source], self.output, format="xlsx", cache=cache)
            prepared = planning.prepare_plan(plan, cache=cache)
        self.source.write_bytes(b"Time\tValue\n0\t5\n1\t6\n")
        with self.assertRaisesRegex(DataImportError, "Source changed after planning"):
            execute_import(prepared)
        self.assertFalse(self.output.exists())

    def test_changed_options_and_returned_mutable_objects_cannot_corrupt_cache(self):
        cache = SourceCache()
        with self.parser() as parse:
            first = planning.inspect_inputs([self.source], cache=cache)
            first["tables"][0]["columns"][1]["sample"][0] = "forged"
            first["tables"][0]["suggested_plot"]["y"].clear()
            table = cache.read(self.source)[0]
            table.read_options["missing_values"].append("1")
            table.read_options["header"] = False
            second = planning.inspect_inputs([self.source], cache=cache)
            no_header = planning.inspect_inputs([self.source], options={"header": False}, cache=cache)
            with self.assertRaises(DataImportError):
                planning.inspect_inputs([self.source], options={"unknown": True}, cache=cache)
        self.assertEqual(second["tables"][0]["columns"][1]["sample"][0], "1")
        self.assertEqual(second["tables"][0]["suggested_plot"]["y"], [1])
        self.assertEqual(no_header["tables"][0]["n_rows"], 3)
        self.assertEqual(parse.call_count, 2)

    def test_eviction_and_oversized_sources_fall_back_to_correct_reads(self):
        other = self.root / "other.tsv"
        other.write_bytes(self.source.read_bytes())
        cache = SourceCache(max_entries=1)
        with self.parser() as parse:
            cache.read(self.source)
            cache.read(other)
            restored = cache.read(self.source)
        self.assertEqual(parse.call_count, 3)
        self.assertEqual(len(cache._entries), 1)
        self.assertEqual(restored[0].columns[1].values, ("1", "2"))
        small = SourceCache(max_bytes=1)
        with self.parser() as parse:
            self.assertEqual(small.read(self.source), small.read(self.source))
        self.assertEqual(parse.call_count, 2)
        self.assertFalse(small._entries)

    def test_missing_or_invalid_sources_do_not_cache_failures(self):
        cache = SourceCache()
        cache.read(self.source)
        self.source.unlink()
        with self.assertRaises(DataImportError):
            cache.read(self.source)
        self.assertFalse(cache._entries)
        self.source.write_bytes(b"")
        with self.assertRaises(DataImportError):
            cache.read(self.source)
        self.source.write_bytes(b"a,b\n1,2\n")
        self.assertEqual(cache.read(self.source)[0].n_rows, 1)

    def test_cold_multi_sheet_plan_reads_only_selected_sheets_once(self):
        path = self.root / "book.xlsx"
        book = Workbook()
        first = book.active
        first.title = "First"
        first.append(["X", "Y"])
        first.append([0, 1])
        second = book.create_sheet("Second")
        second.append(["X", "Y"])
        second.append([0, 2])
        unused = book.create_sheet("Unused")
        unused.append(["X", "Y"])
        unused.append([0, "=1+2"])  # A cached-formula read would fail if this were parsed.
        book.save(path)
        book.close()
        plan = planning.create_plan([path], self.output, format="xlsx", options={"sheet": "First"})
        other = planning.create_plan([path], self.output, format="xlsx", options={"sheet": "Second"})
        plan["tables"].extend(copy.deepcopy(other["tables"]))
        cache = SourceCache()
        with self.parser() as parse:
            prepared = planning.prepare_plan(plan, cache=cache)
        self.assertEqual(parse.call_count, 1)
        self.assertEqual([table.table.source_sheet for table in prepared.tables], ["First", "Second"])
        with self.assertRaisesRegex(DataImportError, "公式没有缓存结果"):
            cache.read(path)

        repeated = copy.deepcopy(plan)
        duplicate = copy.deepcopy(repeated["tables"][0])
        duplicate["name"] = "FirstAgain"
        repeated["tables"].insert(1, duplicate)
        with self.parser() as parse:
            prepared = planning.prepare_plan(repeated, cache=SourceCache(max_bytes=1))
        self.assertEqual(parse.call_count, 1)
        options = prepared.tables[0].table.read_options
        expected = copy.deepcopy(prepared.tables[1].table.read_options)
        options["header"] = False
        options["missing_values"].append("1")
        self.assertEqual(prepared.tables[1].table.read_options, expected)

    def test_many_and_oversized_workbooks_do_not_reparse_evicted_preloads(self):
        paths = []
        for index in range(10):
            path = self.root / f"book{index}.xlsx"
            book = Workbook()
            for sheet_index in range(3):
                sheet = book.active if sheet_index == 0 else book.create_sheet()
                sheet.title = f"Sheet{sheet_index}"
                sheet.append(["X", "Y"])
                sheet.append([index, sheet_index])
            book.save(path)
            book.close()
            paths.append(path)
        plan = planning.create_plan(paths, self.output, format="xlsx")
        for cache in (SourceCache(max_entries=2), SourceCache(max_bytes=1)):
            with self.subTest(budget=cache.max_bytes, entries=cache.max_entries), self.parser() as parse:
                prepared = planning.prepare_plan(plan, cache=cache)
                self.assertEqual(parse.call_count, 10, "load each workbook once, even beyond the LRU budget")
                self.assertEqual(len(prepared.tables), 30)
                for index, table in enumerate(prepared.tables):
                    self.assertEqual(table.table.columns[0].values, (str(index // 3),))

    def test_parallel_operations_do_not_share_unverified_task_state(self):
        cache = SourceCache()
        with self.parser() as parse, ThreadPoolExecutor(max_workers=3) as pool:
            results = list(pool.map(lambda _: planning.inspect_inputs([self.source], cache=cache), range(3)))
        self.assertEqual(parse.call_count, 1)
        self.assertTrue(all(result == results[0] for result in results))
        self.assertFalse(cache._checked)


if __name__ == "__main__":
    unittest.main()
