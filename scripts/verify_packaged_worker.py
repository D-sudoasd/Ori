"""Exercise the Windows GUI executable's real bundled multiprocessing worker."""
import multiprocessing as mp
import multiprocessing.spawn as spawn
import argparse
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook  # noqa: E402
from origin_bridge.planning import create_plan  # noqa: E402
from origin_bridge.worker import import_worker  # noqa: E402


def main(executable=None):
    executable = str(Path(executable or ROOT / "dist" / "SpectraToOrigin.exe").resolve())
    with tempfile.TemporaryDirectory() as raw:
        output = Path(raw).resolve() / "from-frozen-worker.xlsx"
        plan = create_plan([ROOT / "examples_general" / "stress_strain.tsv"], output, format="xlsx")
        original_prepare = spawn.get_preparation_data

        def prepare(name):
            data = original_prepare(name)
            # The child must import its own bundled modules, not the source checkout.
            for key in ("sys_path", "init_main_from_name", "init_main_from_path"):
                data.pop(key, None)
            return data

        spawn.get_preparation_data = prepare
        mp.set_executable(executable)
        sys.executable = sys._base_executable = executable
        sys.frozen = True
        spawn.WINEXE = True
        context = mp.get_context("spawn")
        receive, send = context.Pipe(duplex=False)
        process = context.Process(target=import_worker, args=(send, plan))
        process.start()
        send.close()
        result = None
        deadline = time.monotonic() + 45
        try:
            while time.monotonic() < deadline:
                if receive.poll(0.1):
                    try:
                        message = receive.recv()
                    except EOFError:
                        break
                    if message["type"] == "result":
                        result = message["result"]
                        break
                elif not process.is_alive():
                    break
            process.join(5)
            assert process.exitcode == 0, process.exitcode
            assert result and result["ok"], result
            workbook = load_workbook(output)
            try:
                sheet = workbook.worksheets[0]
                assert sheet.cell(3, 2).value == 120
                assert sheet.cell(6, 2).value is None
                assert sheet.cell(7, 4).value == "Ti-A"
                assert "Provenance" in workbook.sheetnames
            finally:
                workbook.close()
            print("Frozen GUI worker imported typed data and missing cells successfully")
        finally:
            if process.is_alive():
                process.terminate()  # This harness only exports XLSX; it never starts Origin.
                process.join(3)
            receive.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--executable", type=Path, help="要验证的 GUI executable；默认 dist/SpectraToOrigin.exe")
    main(parser.parse_args().executable)
