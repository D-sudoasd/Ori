"""Replay common CLI/MCP imports and count work, without writing repository data.

Pass --repo to replay another checkout with the same interpreter and fixtures.
Counts include real parsing and verification; timings are descriptive only.
"""
from __future__ import annotations

import argparse
import asyncio
import contextlib
import hashlib
import inspect
import io
import json
import statistics
import sys
import tempfile
import time
from collections import Counter
from pathlib import Path
from unittest.mock import patch


def measure(label, sources, action):
    from openpyxl import load_workbook
    from origin_bridge import exporter, planning, readers

    counts = Counter(tool_calls=0, source_reads=0, parse_calls=0,
                     workbook_opens=0, range_calls=0, conversion_checks=0,
                     final_source_checks=0)
    source_paths = {path.resolve() for path in sources}
    real_read = Path.read_bytes
    real_hash = exporter._sha256_file

    def read_bytes(path):
        if path.resolve() in source_paths:
            counts["source_reads"] += 1
        return real_read(path)

    def hash_file(path):
        if path.resolve() in source_paths:
            counts["source_reads"] += 1
            counts["final_source_checks"] += 1
        return real_hash(path)

    def count_call(key, fn):
        def counted(*args, **kwargs):
            counts[key] += 1
            return fn(*args, **kwargs)
        return counted

    def open_book(filename, *args, **kwargs):
        if isinstance(filename, io.BytesIO):
            counts["workbook_opens"] += 1
        return load_workbook(filename, *args, **kwargs)

    with contextlib.ExitStack() as stack:
        stack.enter_context(patch.object(Path, "read_bytes", read_bytes))
        stack.enter_context(patch.object(exporter, "_sha256_file", hash_file))
        stack.enter_context(patch("openpyxl.load_workbook", open_book))
        for name in ("_read_excel", "_read_delimited_rows", "_read_json", "_read_jsonl"):
            stack.enter_context(patch.object(readers, name, count_call("parse_calls", getattr(readers, name))))
        stack.enter_context(patch.object(planning, "_column_range", count_call("range_calls", planning._column_range)))
        stack.enter_context(patch.object(exporter, "_validate_prepared", count_call("conversion_checks", exporter._validate_prepared)))
        start = time.perf_counter()
        output = action(counts)
        elapsed = time.perf_counter() - start
    book = load_workbook(output, read_only=True, data_only=True)
    try:
        # Exclude Provenance paths/hashes; the actual cell values must agree.
        data = {sheet.title: list(sheet.values) for sheet in book if sheet.title != "Provenance"}
    finally:
        book.close()
    digest = hashlib.sha256(json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")).hexdigest()
    return {"task": label, **counts, "elapsed_s": round(elapsed, 4), "data_sha256": digest,
            "table_rows": {name: len(values) for name, values in data.items()}}


def replay(root, rows):
    from openpyxl import Workbook
    from origin_bridge import cli, planning
    from origin_bridge.mcp_server import build_server

    build_server()  # Exclude the one-time optional SDK import from operation timings.

    source = root / "stress.tsv"
    source.write_text("Strain\tStress (MPa)\n" + "".join(f"{i / 1000}\t{i * 2}\n" for i in range(rows)), encoding="utf-8")
    workbook = root / "multi.xlsx"
    book = Workbook()
    for index in range(4):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = f"Run{index + 1}"
        sheet.append(["Time (s)", "Signal (V)"])
        for row in range(rows):
            sheet.append([row, row + index])
    book.save(workbook)
    book.close()

    def direct(path, target):
        def run(counts):
            counts["tool_calls"] += 1
            receipt = cli._dispatch(cli.build_parser().parse_args(
                ["import", "-i", str(path), "-o", str(target)]))
            assert receipt["ok"], receipt
            return target
        return run

    results = [
        measure("cli_tsv_import", [source], direct(source, root / "tsv.xlsx")),
        measure("cli_four_sheet_import", [workbook], direct(workbook, root / "sheets.xlsx")),
    ]

    target = root / "replayed.xlsx"
    plan_path = root / "saved-plan.json"
    plan_path.write_text(json.dumps(planning.create_plan([workbook], target, format="xlsx")), encoding="utf-8")

    def replay_plan(counts):
        counts["tool_calls"] += 1
        result = cli._dispatch(cli.build_parser().parse_args(["execute", str(plan_path)]))
        assert result["ok"], result
        return target

    results.append(measure("cli_saved_four_sheet_plan", [workbook], replay_plan))

    for checks in (True, False):
        target = root / ("mcp_four.xlsx" if checks else "mcp_two.xlsx")

        def mcp_run(counts):
            server = build_server()

            async def call(name, arguments):
                counts["tool_calls"] += 1
                result = await server.call_tool(name, arguments)
                return result.structured_content

            async def run():
                if checks:
                    await call("inspect_data", {"paths": [str(source)]})
                arguments = {"paths": [str(source)], "output_path": str(target), "output_format": "xlsx"}
                tool = server._tool_manager.get_tool("create_import_plan")
                combined = not checks and "include_inspection" in inspect.signature(tool.fn).parameters
                if combined:
                    arguments["include_inspection"] = True
                planned = await call("create_import_plan", arguments)
                plan = planned["plan"] if combined else planned
                if checks:
                    await call("validate_import_plan", {"plan": plan})
                receipt = await call("execute_import_plan", {"plan": plan, "confirm": True})
                assert receipt["ok"], receipt
            asyncio.run(run())
            return target

        results.append(measure("mcp_four_tools" if checks else "mcp_plan_execute", [source], mcp_run))
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--rows", type=int, default=2000)
    parser.add_argument("--repeats", type=int, default=1)
    args = parser.parse_args()
    if args.rows < 1 or args.repeats < 1:
        parser.error("--rows and --repeats must be positive")
    sys.path.insert(0, str(args.repo.resolve()))
    trials = []
    for _ in range(args.repeats):
        with tempfile.TemporaryDirectory(prefix="ori-agent-workflow-") as temp:
            trials.append(replay(Path(temp), args.rows))
    results = trials[0]
    for index, result in enumerate(results):
        expected = {key: value for key, value in result.items() if key != "elapsed_s"}
        for trial in trials:
            assert {key: value for key, value in trial[index].items() if key != "elapsed_s"} == expected
        result["elapsed_s"] = round(statistics.median(trial[index]["elapsed_s"] for trial in trials), 4)
    print(json.dumps({"rows": args.rows, "repeats": args.repeats, "results": results}, ensure_ascii=False))


if __name__ == "__main__":
    main()
