"""Generic typed exporter tests; the Origin path uses an isolated fake API."""

from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import load_workbook

from origin_bridge import exporter
from origin_bridge.models import DataColumn, DataImportError, DataTable, PlotSpec, PlannedTable, PreparedImport


class FakeOriginColumn:
    def __init__(self):
        self.format = None
        self.categories = None

    def SetDataFormat(self, value):
        self.format = value

    def SetCategMapSortCategories(self, mode, categories):
        self.categories = (mode, list(categories))


class FakeSheet:
    def __init__(self, name="Sheet1"):
        self.name = name
        self.lname = name
        self._cols = 1
        self.obj = [FakeOriginColumn()]
        self.data = {}
        self.labels = {}
        self.roles = []
        self.date_formats = {}

    @property
    def cols(self):
        return self._cols

    @cols.setter
    def cols(self, value):
        self._cols = value
        self.obj = [FakeOriginColumn() for _ in range(value)]

    @property
    def rows(self):
        return len(self.data.get(0, []))

    def from_list(self, index, values, lname="", units="", comments=""):
        self.data[index] = list(values)
        self.labels[index] = {"L": lname, "U": units, "C": comments}

    def to_list(self, index):
        return list(self.data[index])

    def get_label(self, index, label):
        return self.labels[index].get(label, "")

    def cols_axis(self, *args):
        self.roles.append(args)

    def as_date(self, index, fmt):
        self.date_formats[index] = fmt


class FakeBook:
    def __init__(self):
        self.lname = ""
        self.sheets = [FakeSheet()]

    def __getitem__(self, index):
        return self.sheets[index]

    def __iter__(self):
        return iter(self.sheets)

    def add_sheet(self, name):
        sheet = FakeSheet(name)
        self.sheets.append(sheet)
        return sheet


class FakeAxis:
    def __init__(self):
        self.title = ""


class FakeLayer:
    def __init__(self):
        self.obj = SimpleNamespace(DataPlots=[])
        self.axes = {"x": FakeAxis(), "y": FakeAxis()}
        self.plot_args = []

    def add_plot(self, sheet, **kwargs):
        self.plot_args.append(kwargs)
        self.obj.DataPlots.append(dict(kwargs))
        if "colyerr" in kwargs:
            self.obj.DataPlots.append({"error_for": kwargs["coly"]})
        return object()

    def axis(self, name):
        return self.axes[name]

    def rescale(self):
        pass


class FakeGraph:
    def __init__(self, title):
        self.lname = title
        self.layer = FakeLayer()

    def __getitem__(self, index):
        assert index == 0
        return self.layer


class FakeOrigin:
    def __init__(self):
        self.book = FakeBook()
        self.graphs = []
        self.sysvar = {"DSO": 2415018.5}
        self.visible = None

    def set_show(self, value):
        self.visible = value

    def new(self, _):
        self.book = FakeBook()
        self.graphs = []

    def pages(self, page_type):
        return [self.book] if page_type == "w" else list(self.graphs)

    def new_graph(self, lname, hidden=False):
        graph = FakeGraph(lname)
        self.graphs.append(graph)
        return graph


class GenericExporterTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name).resolve()
        self.source = self.root / "source.json"
        self.source.write_text('{"source":"fixture"}', encoding="utf-8")
        self.source_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()

        def isolated_lock():
            lock = exporter._OriginExportLock()
            lock.path = self.root / "export.lock"
            return lock

        patcher = mock.patch.object(exporter, "_origin_export_lock", side_effect=isolated_lock)
        patcher.start()
        self.addCleanup(patcher.stop)

    def planned(self, name="Measurements", columns=None, plot=None):
        if columns is None:
            columns = (
                DataColumn("Elapsed", "number", ("0", "1", "2"), "s"),
                DataColumn("Sample", "text", ("A", "B", "A")),
                DataColumn("Reading", "number", ("1.25", None, "3.5"), "N"),
                DataColumn("Uncertainty", "number", ("0.1", "0", "0.2"), "N"),
                DataColumn("When", "datetime", ("2026-01-01T12:00:00+08:00", None, "2026-01-03T00:00:00Z")),
            )
        table = DataTable(
            source=self.source,
            source_sheet="Sweep",
            name=name,
            columns=columns,
            source_hash=self.source_hash,
            read_options={"sheet": "Sweep", "header": True},
        )
        return PlannedTable(table, name, plot or PlotSpec())

    def prepared(self, output, *, fmt="xlsx", tables=None, overwrite=False):
        return PreparedImport(
            tables=tuple(tables or [self.planned()]),
            output=output,
            format=fmt,
            overwrite=overwrite,
        )

    def test_xlsx_writes_all_columns_as_typed_values_and_records_physical_sheet_name(self):
        planned = self.planned(name="Provenance", plot=PlotSpec(kind="line", x=0, y=(2,)))
        output = self.root / "typed.xlsx"
        receipt = exporter.execute_import(self.prepared(output, tables=[planned]))

        self.assertTrue(receipt["ok"])
        self.assertEqual(receipt["tables"][0]["output_sheet"], "Provenance_2")
        self.assertEqual(receipt["tables"][0]["rows"], 3)
        self.assertIsNone(receipt["tables"][0]["graph"], "XLSX output must not claim an Origin graph was created")
        workbook = load_workbook(output, data_only=False)
        try:
            sheet = workbook["Provenance_2"]
            self.assertEqual([sheet.cell(1, index).value for index in range(1, 6)],
                             ["Elapsed", "Sample", "Reading", "Uncertainty", "When"])
            self.assertEqual(sheet["A2"].value, 0)
            self.assertEqual(sheet["B2"].value, "A")
            self.assertEqual(sheet["C3"].value, None)
            self.assertEqual(sheet["E2"].value.isoformat(), "2026-01-01T04:00:00")
            self.assertEqual(workbook["Provenance"]["A2"].value, "table")
            note = json.loads(workbook["Provenance"]["N2"].value)
            self.assertEqual(note["output_sheet"], "Provenance_2")
            self.assertEqual(note["warnings"][0].split(":")[0], "Provenance/When")
        finally:
            workbook.close()

    def test_xlsx_formula_like_text_is_stored_as_text_and_long_cells_rejected(self):
        planned = self.planned(columns=(DataColumn("Entry", "text", ("=1+1", "+cmd", "@SUM(A1:A2)", "")),))
        output = self.root / "formula.xlsx"
        receipt = exporter.execute_import(self.prepared(output, tables=[planned]))
        self.assertTrue(any("empty string(s)" in warning for warning in receipt["warnings"]))
        workbook = load_workbook(output, data_only=False)
        try:
            sheet = workbook["Measurements"]
            self.assertEqual(sheet["A2"].value, "=1+1")
            self.assertEqual(sheet["A2"].data_type, "s")
            self.assertIsNone(sheet["A5"].value)
        finally:
            workbook.close()

        huge = self.planned(columns=(DataColumn("Entry", "text", ("x" * 32_768,)),))
        rejected = self.root / "too_long.xlsx"
        with self.assertRaisesRegex(DataImportError, "cell limit"):
            exporter.execute_import(self.prepared(rejected, tables=[huge]))
        self.assertFalse(rejected.exists())

    def test_no_clobber_is_atomic_when_destination_appears_during_export(self):
        output = self.root / "race.xlsx"
        original_writer = exporter._write_xlsx

        def writer(stage, tables, *args, **kwargs):
            original_writer(stage, tables, *args, **kwargs)
            output.write_bytes(b"external writer owns this")

        with mock.patch.object(exporter, "_write_xlsx", side_effect=writer):
            with self.assertRaises(FileExistsError):
                exporter.execute_import(self.prepared(output, overwrite=False))
        self.assertEqual(output.read_bytes(), b"external writer owns this")

    def test_numeric_binary64_rounding_warns_and_underflow_is_rejected(self):
        rounded = self.planned(columns=(DataColumn("ID", "number", ("9007199254740993",)),))
        result = exporter.execute_import(self.prepared(self.root / "rounded.xlsx", tables=[rounded]))
        self.assertTrue(any("rounded to binary64" in warning for warning in result["warnings"]))
        underflow = self.planned(columns=(DataColumn("Tiny", "number", ("1e-400",)),))
        with self.assertRaisesRegex(DataImportError, "underflows to zero"):
            exporter.execute_import(self.prepared(self.root / "underflow.xlsx", tables=[underflow]))

    def test_origin_maps_column_categories_dates_plots_and_errors_without_originpro(self):
        origin = FakeOrigin()
        support = SimpleNamespace(
            OriginExportError=RuntimeError,
            origin_process_running=lambda: False,
            _load_originpro=lambda: origin,
            _save_origin_project=lambda _op, path, pending: (path.write_bytes(b"fake opju") or path),
            close_origin_app=lambda _op, started: None,
            _cleanup_origin_temp_dirs=lambda directories: None,
        )
        # Path.write_bytes returns an int, so keep the fake saver explicit.
        def save(_op, path, pending_cleanup):
            path.write_bytes(b"fake opju")
            return path
        support._save_origin_project = save
        plot = PlotSpec(kind="column", x=1, y=(2,), y_error=((2, 3),), title="Force by sample")
        output = self.root / "typed.opju"
        planned = self.planned(plot=plot)
        with mock.patch.object(exporter, "_origin_support", return_value=support), \
             mock.patch.object(exporter, "_origin_format_codes", return_value={"number": 1, "text": 2, "datetime": 3}), \
             mock.patch.object(exporter, "_origin_custom_category_sort", return_value="custom"):
            receipt = exporter.execute_import(self.prepared(output, fmt="opju", tables=[planned]))

        self.assertTrue(receipt["ok"])
        self.assertEqual(receipt["tables"][0]["output_sheet"], "Measurements")
        sheet = origin.book[0]
        self.assertEqual([sheet.obj[index].format for index in range(5)], [1, 2, 1, 1, 3])
        self.assertEqual(sheet.obj[1].categories, ("custom", ["A", "B"]))
        self.assertEqual(sheet.get_label(0, "U"), "s")
        self.assertEqual(sheet.get_label(2, "U"), "N")
        self.assertAlmostEqual(sheet.to_list(4)[0], exporter._origin_date_number(
            exporter._datetime_value("2026-01-01T12:00:00+08:00", "T", "When", 1), 2415018.5))
        layer = origin.graphs[0][0]
        self.assertEqual(layer.plot_args, [{"coly": 2, "colx": 1, "type": "c", "colyerr": 3}])
        self.assertEqual(len(layer.obj.DataPlots), 2)
        self.assertEqual(layer.axis("x").title, "Sample")
        self.assertEqual(layer.axis("y").title, "Reading (N)")
        self.assertEqual(origin.book.sheets[-1].lname, "Import Provenance")
        self.assertNotEqual(origin.book.sheets[0].lname, origin.book.sheets[-1].lname)

    def test_failed_origin_save_preserves_existing_target(self):
        origin = FakeOrigin()
        support = SimpleNamespace(
            OriginExportError=RuntimeError,
            origin_process_running=lambda: False,
            _load_originpro=lambda: origin,
            _save_origin_project=lambda *args: (_ for _ in ()).throw(RuntimeError("save failed")),
            close_origin_app=lambda _op, started: None,
            _cleanup_origin_temp_dirs=lambda directories: None,
        )
        output = self.root / "keep.opju"
        output.write_bytes(b"old target")
        with mock.patch.object(exporter, "_origin_support", return_value=support), \
             mock.patch.object(exporter, "_origin_format_codes", return_value={"number": 1, "text": 2, "datetime": 3}):
            with self.assertRaisesRegex(RuntimeError, "save failed"):
                exporter.execute_import(self.prepared(output, fmt="opju", overwrite=True))
        self.assertEqual(output.read_bytes(), b"old target")
        with exporter._origin_export_lock():
            pass

    def test_xlsx_and_origin_warn_about_native_datetime_precision_limits(self):
        planned = self.planned(columns=(
            DataColumn("Timestamp", "datetime", ("2026-01-01T12:00:00.123456789",)),
        ))
        xlsx_warning = exporter._table_conversion_warnings(planned, "xlsx")
        opju_warning = exporter._table_conversion_warnings(planned, "opju")
        self.assertTrue(any("milliseconds" in item for item in xlsx_warning))
        self.assertTrue(any("microseconds" in item for item in opju_warning))


if __name__ == "__main__":
    unittest.main()
