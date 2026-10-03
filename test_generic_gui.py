"""GUI contracts and spawn tests, with opt-in real Tk integration; no Origin."""

from __future__ import annotations

import json
import multiprocessing
import os
import sys
import tempfile
import time
import types
import unittest
from pathlib import Path
from unittest import mock
import tkinter as tk

from openpyxl import load_workbook

from origin_bridge import planning
from origin_bridge.gui import (
    _collect_supported,
    _default_output_for,
    _plan_table_key,
    _read_options,
    _source_key,
    _validate_plan_save_target,
)
from origin_bridge.worker import execute_import_request, import_worker


class Value:
    def __init__(self, value):
        self.value = value

    def get(self):
        return self.value


class YList:
    def __init__(self, values: list[str]):
        self.values = values

    def curselection(self):
        return tuple(range(len(self.values)))

    def get(self, index: int) -> str:
        return self.values[index]


class GenericGuiContractTests(unittest.TestCase):
    def test_read_options_default_to_empty_string_as_missing_value(self):
        options = _read_options("自动", "0", "自动识别")

        self.assertEqual(options["missing_values"], [""])
        self.assertEqual(options["header"], "auto")
        self.assertEqual(options["delimiter"], "auto")

    def test_empty_cells_do_not_turn_numeric_columns_into_text(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "values.csv"
            source.write_text("time,signal\n0,1\n1,\n2,3\n", encoding="utf-8")

            table = planning.inspect_inputs([source], options=_read_options("自动", "0", "自动识别"))["tables"][0]

        self.assertEqual([column["kind"] for column in table["columns"]], ["number", "number"])
        self.assertEqual(table["columns"][1]["missing"], 1)
        self.assertEqual(table["columns"][1]["sample"], ["1", None, "3"])

    def test_folder_expansion_uses_reader_extensions_and_natural_order(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            first = folder / "sample10.ndjson"
            second = folder / "sample2.jsonl"
            ignored = folder / "notes.md"
            first.write_text('{"x":10}\n', encoding="utf-8")
            second.write_text('{"x":2}\n', encoding="utf-8")
            ignored.write_text("notes", encoding="utf-8")

            found, rejected = _collect_supported([folder, second])

        self.assertEqual([path.name for path in found], ["sample2.jsonl", "sample10.ndjson"])
        self.assertEqual(rejected, [])

    def test_explicit_missing_input_is_not_silently_ignored(self):
        with tempfile.TemporaryDirectory() as raw:
            with self.assertRaisesRegex(FileNotFoundError, "does not exist|不存在"):
                _collect_supported([Path(raw) / "missing.csv"])

    def test_default_workbook_path_does_not_overwrite_an_input_workbook(self):
        with tempfile.TemporaryDirectory() as raw:
            source = Path(raw) / "study.xlsx"
            source.touch()

            output = _default_output_for([source], "xlsx")

        self.assertNotEqual(output.resolve(), source.resolve())
        self.assertEqual(output.name, "study_imported.xlsx")

    def test_saved_plan_cannot_replace_a_source_or_planned_output(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            plan = {
                "output": {"path": "result.xlsx"},
                "tables": [{"source": {"path": "source.json"}}],
            }
            with self.assertRaisesRegex(ValueError, "source.json"):
                _validate_plan_save_target(plan, folder / "source.json", folder)
            with self.assertRaisesRegex(ValueError, "result.xlsx"):
                _validate_plan_save_target(plan, folder / "result.xlsx", folder)

    def test_relative_plan_table_identity_uses_the_plan_directory(self):
        source = {"path": "data.csv", "options": {"sheet": None}}
        table = {"source": source}
        with tempfile.TemporaryDirectory() as raw:
            base = Path(raw)
            expected = str((base / "data.csv").resolve())
            self.assertEqual(_source_key(source, base)[0], expected.casefold())
            self.assertEqual(_plan_table_key(table, base), (expected.casefold(), None))

    def test_plot_error_mapping_uses_original_y_column_name_and_none_resets_axes(self):
        from origin_bridge.gui import GeneralDataApp

        app = GeneralDataApp.__new__(GeneralDataApp)
        app._current_key = ("source.csv", None)
        app._current_table = {
            "columns": [
                {"index": 0, "name": "Time"},
                {"index": 1, "name": "Signal"},
                {"index": 2, "name": "Uncertainty"},
            ]
        }
        app._plots = {}
        app.y_list = YList(["1: Signal (number)"])
        app.error_choice_var = Value("2: Uncertainty (number)")
        app.plot_kind_var = Value("line")
        app.x_choice_var = Value("0: Time (number)")
        app.title_var = Value("Signal")
        app.x_label_var = Value("Time")
        app.y_label_var = Value("Signal")

        app._save_current_plot()
        self.assertEqual(app._plots[app._current_key]["y_error"], {"Signal": 2})

        app.plot_kind_var = Value("none")
        app._save_current_plot()
        self.assertEqual(app._plots[app._current_key]["x"], None)
        self.assertEqual(app._plots[app._current_key]["y"], [])
        self.assertEqual(app._plots[app._current_key]["y_error"], {})

    def test_gui_plot_choices_build_a_plan_accepted_by_the_import_validator(self):
        from origin_bridge.gui import GeneralDataApp

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "signal.csv"
            source.write_text("Time,Signal,Error\n0,10,0.2\n1,12,0.3\n", encoding="utf-8")
            table = planning.inspect_inputs([source], options=_read_options("自动", "0", "自动识别"))["tables"][0]
            app = GeneralDataApp.__new__(GeneralDataApp)
            app._current_key = _source_key(table["source"])
            app._current_table = table
            app._plots = {}
            app._column_labels = {}
            app._loaded_plan = None
            app._plan_base_dir = None
            app._paths = [source]
            app._tables = [table]
            app.y_list = YList(["1: Signal (number)"])
            app.error_choice_var = Value("2: Error (number)")
            app.plot_kind_var = Value("line")
            app.x_choice_var = Value("0: Time (number)")
            app.title_var = Value("Signal over time")
            app.x_label_var = Value("Time")
            app.y_label_var = Value("Signal")
            app.header_var = Value("自动")
            app.skip_rows_var = Value("0")
            app.delimiter_var = Value("自动识别")
            app.keep_open_var = Value(False)

            plan = app._effective_plan(folder / "result.xlsx", "xlsx", False)
            prepared = planning.prepare_plan(plan)

        self.assertEqual(plan["tables"][0]["plot"]["y_error"], {"Signal": 2})
        self.assertEqual(prepared.tables[0].plot.y_error, ((1, 2),))

    def test_worker_prepares_then_exports_without_importing_tk(self):
        calls = []
        prepared = object()
        planning_module = types.ModuleType("origin_bridge.planning")
        planning_module.prepare_plan = lambda plan, base_dir=None: calls.append(("prepare", plan, base_dir)) or prepared
        exporter_module = types.ModuleType("origin_bridge.exporter")
        exporter_module.execute_import = lambda value: calls.append(("export", value)) or {
            "ok": True,
            "output": {"path": "result.xlsx", "format": "xlsx", "size_bytes": 10, "sha256": "a" * 64},
            "tables": [],
            "warnings": [],
        }
        plan = {"schema_version": 1}
        base_dir = Path("plan-folder")

        with mock.patch.dict(sys.modules, {
            "origin_bridge.planning": planning_module,
            "origin_bridge.exporter": exporter_module,
        }):
            result = execute_import_request(plan, base_dir)

        self.assertTrue(result["ok"])
        self.assertEqual(calls, [("prepare", plan, base_dir), ("export", prepared)])
        self.assertFalse(hasattr(sys.modules["origin_bridge.worker"], "tk"))
        json.dumps(result)

    def test_spawn_worker_exports_xlsx_without_opening_a_gui(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "worker.csv"
            output = folder / "worker.xlsx"
            source.write_text("Time,Signal\n0,4\n1,5\n", encoding="utf-8")
            plan = planning.create_plan(
                [source], output, format="xlsx", options=_read_options("自动", "0", "自动识别")
            )
            context = multiprocessing.get_context("spawn")
            recv_conn, send_conn = context.Pipe(duplex=False)
            process = context.Process(target=import_worker, args=(send_conn, plan, None))
            process.start()
            send_conn.close()
            messages = []
            deadline = time.monotonic() + 30
            try:
                while time.monotonic() < deadline:
                    try:
                        if recv_conn.poll(0.1):
                            messages.append(recv_conn.recv())
                        elif not process.is_alive():
                            break
                    except (EOFError, OSError):
                        break
                process.join(timeout=2)
                self.assertFalse(process.is_alive(), "spawn worker did not finish")
            finally:
                if process.is_alive():
                    process.terminate()
                    process.join(timeout=2)
                recv_conn.close()
            exit_code = process.exitcode
            process.close()
            results = [message["result"] for message in messages if message.get("type") == "result"]
            self.assertEqual(exit_code, 0, messages)
            self.assertTrue(results and results[-1]["ok"], messages)
            self.assertTrue(output.is_file() and output.stat().st_size > 0)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook["worker"].cell(3, 2).value, 5)
            finally:
                workbook.close()

    @unittest.skipUnless(os.environ.get("SPECTRA_TEST_GUI") == "1", "set SPECTRA_TEST_GUI=1 for the Tk and spawn-export integration check")
    def test_real_tk_remains_responsive_during_spawned_xlsx_export(self):
        from origin_bridge import gui
        from origin_bridge.gui import GeneralDataApp

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "signal.csv"
            source.write_text("Time,Signal\n0,10\n1,\n2,12\n", encoding="utf-8")
            plan_folder = folder / "saved-plans"
            plan_folder.mkdir()
            plan_path = plan_folder / "signal_plan.json"
            output = folder / "signal_result.xlsx"
            try:
                app = GeneralDataApp([source])
            except tk.TclError as exc:
                self.skipTest(f"Tk display is unavailable: {exc}")
            process = None
            try:
                app.root.withdraw()
                deadline = time.monotonic() + 15
                while app._busy and time.monotonic() < deadline:
                    app.root.update()
                    time.sleep(0.01)
                self.assertFalse(app._busy, "source inspection did not finish")
                self.assertEqual(app._tables[0]["columns"][1]["kind"], "number")
                self.assertEqual(app._tables[0]["columns"][1]["sample"], ["10", None, "12"])
                self.assertEqual(len(app.sample_tree.get_children()), 3)
                second_preview = app.sample_tree.item(app.sample_tree.get_children()[1], "values")
                self.assertEqual(second_preview[1], "")

                app.format_var.set("xlsx")
                app.plot_kind_var.set("none")
                with (
                    mock.patch.object(gui.filedialog, "asksaveasfilename", return_value=str(plan_path)),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                ):
                    app.save_plan()
                    showerror.assert_not_called()
                saved_plan = json.loads(plan_path.read_text(encoding="utf-8"))
                saved_plan["tables"][0]["source"]["path"] = os.path.relpath(source, plan_folder)
                saved_plan["output"]["path"] = "future.xlsx"
                plan_path.write_text(json.dumps(saved_plan), encoding="utf-8")
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", return_value=str(plan_path)),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                ):
                    app.load_plan()
                    deadline = time.monotonic() + 15
                    while app._busy and time.monotonic() < deadline:
                        app.root.update()
                        time.sleep(0.01)
                    self.assertFalse(app._busy, "relative-path plan inspection did not finish")
                    showerror.assert_not_called()
                self.assertEqual(app._paths, [source.resolve()])
                self.assertEqual(app._plan_base_dir, plan_folder.resolve())
                self.assertEqual(app.plot_kind_var.get(), "none")

                ticks = []

                def heartbeat():
                    ticks.append(time.monotonic())
                    if app._busy:
                        app.root.after(10, heartbeat)

                with (
                    mock.patch.object(gui.filedialog, "asksaveasfilename", return_value=str(output)),
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(gui.messagebox, "showwarning"),
                ):
                    app.root.after(10, heartbeat)
                    app.export_data()
                    self.assertTrue(app._busy)
                    deadline = time.monotonic() + 45
                    while app._busy and time.monotonic() < deadline:
                        app.root.update()
                        time.sleep(0.01)
                    app.root.update()
                    self.assertFalse(app._busy, "spawned importer did not finish")
                    self.assertGreaterEqual(len(ticks), 2, "Tk callbacks stopped while the worker exported")
                    showerror.assert_not_called()
                self.assertEqual(str(app.export_button.cget("state")), "normal")
                self.assertTrue(output.is_file() and output.stat().st_size > 0)
                workbook = load_workbook(output, read_only=True, data_only=True)
                try:
                    sheet = workbook["signal"]
                    self.assertEqual(sheet.cell(3, 2).value, None)
                    self.assertEqual(sheet.cell(4, 2).value, 12)
                    self.assertIn("Provenance", workbook.sheetnames)
                finally:
                    workbook.close()
            finally:
                process = app._import_process
                if process is not None and process.is_alive():
                    process.terminate()
                    process.join(timeout=3)
                app.root.destroy()


class _Var:
    def __init__(self, value=""):
        self.value = value

    def set(self, value):
        self.value = value

    def get(self):
        return self.value


class _Pipe:
    def __init__(self, messages):
        self.messages = list(messages)
        self.batch = 0
        self.max_batch = 0
        self.recv_count = 0

    def poll(self, _timeout=0):
        return bool(self.messages)

    def recv(self):
        if not self.messages:
            raise EOFError
        self.batch += 1
        self.recv_count += 1
        return self.messages.pop(0)

    def close(self):
        self.messages.clear()

    def note_batch(self):
        self.max_batch = max(self.max_batch, self.batch)
        self.batch = 0


class _Process:
    def __init__(self, alive=True, exit_code=0):
        self.alive = alive
        self.exitcode = exit_code
        self.closed = False

    def is_alive(self):
        return self.alive

    def join(self, timeout=None):
        self.alive = False

    def close(self):
        self.closed = True


def _contains_data_table(value):
    from origin_bridge.models import DataTable

    if isinstance(value, DataTable):
        return True
    if isinstance(value, dict):
        return any(_contains_data_table(item) for item in value.values())
    if isinstance(value, (list, tuple)):
        return any(_contains_data_table(item) for item in value)
    return False


class SummaryStoreTests(unittest.TestCase):
    def _options(self, header="自动", skip="0", delimiter="自动识别"):
        return _read_options(header, skip, delimiter)

    def test_incremental_reads_ignore_size_and_mtime_and_keep_going_after_a_bad_source(self):
        from origin_bridge import planning as planning_module
        from origin_bridge.source_summary import SummaryStore

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            first = folder / "first.csv"
            second = folder / "second.csv"
            changed = folder / "changed.csv"
            bad = folder / "bad.csv"
            wide = folder / "wide.csv"
            payload_a = b"a,b\n1,2\n"
            payload_b = b"a,b\n3,4\n"
            self.assertEqual(len(payload_a), len(payload_b))
            first.write_bytes(b"a,b\n1,2\n")
            second.write_bytes(b"a,b\n5,6\n")
            changed.write_bytes(payload_a)
            bad.write_bytes(b"")
            wide.write_text("x,y\n" + "\n".join(f"{index},{index}" for index in range(30)) + "\n", encoding="utf-8")
            stamp = 1_700_000_000
            os.utime(changed, (stamp, stamp))
            options = self._options()
            calls = []
            real_inspect = planning_module.inspect_inputs

            def spy(paths, options=None):
                calls.append([Path(path) for path in paths])
                return real_inspect(paths, options=options)

            store = SummaryStore()
            with mock.patch.object(planning_module, "inspect_inputs", side_effect=spy):
                first_result = store.inspect_path(first, options)
                second_result = store.inspect_path(second, options)
                bad_result = store.inspect_path(bad, options)
                wide_result = store.inspect_path(wide, options)
                changed_first = store.inspect_path(changed, options)
                misses = store.inspect_calls
                again = store.inspect_path(first, options)
                self.assertEqual(store.inspect_calls, misses)
                size_before = changed.stat().st_size
                mtime_before = changed.stat().st_mtime
                changed.write_bytes(payload_b)
                os.utime(changed, (stamp, stamp))
                self.assertEqual(changed.stat().st_size, size_before)
                self.assertEqual(changed.stat().st_mtime, mtime_before)
                rewritten = store.inspect_path(changed, options)
                reread = store.inspect_path(changed, options)
                shifted = store.inspect_path(first, self._options(header="无标题行"))

            self.assertTrue(all(len(batch) == 1 for batch in calls))
            self.assertTrue(first_result["ok"] and second_result["ok"] and wide_result["ok"] and changed_first["ok"])
            self.assertEqual(changed_first["tables"][0]["columns"][0]["sample"][0], "1")
            self.assertFalse(bad_result["ok"])
            self.assertIn("空", bad_result["error"])
            self.assertTrue(again["reused"])
            self.assertEqual(store.cache_hits, 2)
            self.assertFalse(rewritten["reused"])
            self.assertEqual(rewritten["tables"][0]["columns"][0]["sample"][0], "3")
            self.assertTrue(reread["reused"])
            self.assertEqual(store.inspect_calls, misses + 2)
            self.assertEqual([column["name"] for column in shifted["tables"][0]["columns"]], ["Column 1", "Column 2"])
            self.assertFalse(shifted["reused"])
            self.assertEqual(wide_result["tables"][0]["n_rows"], 30)
            self.assertLessEqual(len(wide_result["tables"][0]["columns"][0]["sample"]), 5)
            self.assertNotIn("values", wide_result["tables"][0]["columns"][0])
            self.assertFalse(_contains_data_table(wide_result))
            self.assertFalse(_contains_data_table(store._slots))


class ImportPollBudgetTests(unittest.TestCase):
    def _app(self):
        from origin_bridge.gui import GeneralDataApp

        app = GeneralDataApp.__new__(GeneralDataApp)
        app.status_var = _Var()
        app._busy = True
        app._import_messages = []
        app._import_result = None
        app._import_exit_waits = 0
        app._import_eof = False
        app._set_busy = lambda busy: setattr(app, "_busy", busy)
        app.root = types.SimpleNamespace(after=lambda *_args, **_kwargs: None)
        return app

    def test_progress_stream_is_bounded_per_poll_and_final_result_is_kept(self):
        from origin_bridge import gui

        app = self._app()
        messages = [{"type": "progress", "text": f"step {index}"} for index in range(100)]
        messages.append({"type": "result", "result": {"ok": True, "output": {"path": "out.xlsx", "size_bytes": 12}, "warnings": []}})
        pipe = _Pipe(messages)
        process = _Process(alive=True)
        app._import_recv = pipe
        app._import_process = process
        with mock.patch.object(gui.messagebox, "showinfo") as showinfo, mock.patch.object(gui.messagebox, "showerror") as showerror:
            while pipe.messages and process.alive:
                app._poll_import()
                pipe.note_batch()
                self.assertLessEqual(pipe.max_batch, 20)
            self.assertEqual(pipe.recv_count, 101)
            self.assertEqual(app._import_messages, [])
            self.assertTrue(app._import_result["ok"])
            showinfo.assert_not_called()
            process.alive = False
            app._poll_import()
            showerror.assert_not_called()
            showinfo.assert_called_once()
        self.assertFalse(app._busy)
        self.assertIn("out.xlsx", app.status_var.get())

    def test_missing_result_after_process_exit_reports_the_exit_code(self):
        from origin_bridge import gui

        app = self._app()
        app._import_recv = _Pipe([])
        app._import_process = _Process(alive=False, exit_code=3)
        with mock.patch.object(gui.messagebox, "showerror") as showerror, mock.patch.object(gui.messagebox, "showinfo") as showinfo:
            for _ in range(5):
                if not app._busy:
                    break
                app._poll_import()
            showinfo.assert_not_called()
            showerror.assert_called_once()
            self.assertIn("3", showerror.call_args.args[1])
        self.assertFalse(app._busy)
        self.assertEqual(app._import_messages, [])


def _pump(app, seconds: float) -> None:
    deadline = time.monotonic() + seconds
    while app._busy and time.monotonic() < deadline:
        app.root.update()
        time.sleep(0.005)
    app.root.update()


class RealTkSourcePerformanceTests(unittest.TestCase):
    @unittest.skipUnless(os.environ.get("SPECTRA_TEST_GUI") == "1", "set SPECTRA_TEST_GUI=1 for the Tk and spawn-export integration check")
    def test_real_tk_reuses_summaries_and_keeps_going_after_one_bad_source(self):
        from origin_bridge import gui
        from origin_bridge.gui import GeneralDataApp

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            good = folder / "signal.csv"
            other = folder / "other.csv"
            extra = folder / "extra.csv"
            bad = folder / "bad.csv"
            payload_a = b"Time,Signal\n0,1\n"
            payload_b = b"Time,Signal\n9,8\n"
            self.assertEqual(len(payload_a), len(payload_b))
            good.write_bytes(payload_a)
            other.write_bytes(b"Time,Signal\n2,3\n")
            extra.write_bytes(b"Time,Signal\n4,5\n")
            os.utime(good, (1_700_000_000, 1_700_000_000))
            try:
                app = GeneralDataApp()
            except tk.TclError as exc:
                self.skipTest(f"Tk display is unavailable: {exc}")
            try:
                app.root.withdraw()
                with mock.patch.object(gui.messagebox, "showerror") as showerror:
                    app._add_paths([folder / "missing.csv"])
                    _pump(app, 10)
                    self.assertFalse(app._busy)
                    showerror.assert_called()
                self.assertEqual(app._paths, [])
                self.assertEqual(str(app.export_button.cget("state")), "normal")

                app._add_paths([good])
                _pump(app, 15)
                self.assertFalse(app._busy, "first source inspect did not finish")
                self.assertEqual(app._summaries.inspect_calls, 1)
                self.assertEqual(app._tables[0]["columns"][1]["sample"], ["1"])
                app.title_var.set("KEEP ME")
                app._save_current_plot()

                app._add_paths([other])
                _pump(app, 15)
                self.assertEqual(app._summaries.inspect_calls, 2)
                self.assertEqual(len(app._tables), 2)

                hits = app._summaries.cache_hits
                app.reload_inputs()
                _pump(app, 15)
                self.assertEqual(app._summaries.inspect_calls, 2)
                self.assertGreaterEqual(app._summaries.cache_hits, hits + 2)
                self.assertEqual(app.title_var.get(), "KEEP ME")

                mtime = good.stat().st_mtime
                size = good.stat().st_size
                good.write_bytes(payload_b)
                os.utime(good, (1_700_000_000, 1_700_000_000))
                self.assertEqual(good.stat().st_size, size)
                self.assertEqual(good.stat().st_mtime, mtime)
                app.reload_inputs()
                _pump(app, 15)
                self.assertEqual(app._summaries.inspect_calls, 3)
                self.assertEqual(app._tables[0]["columns"][1]["sample"], ["8"])
                self.assertNotEqual(app.title_var.get(), "KEEP ME")

                app.title_var.set("KEEP HEADER")
                app._save_current_plot()
                app.header_var.set("无标题行")
                app.reload_inputs()
                _pump(app, 15)
                self.assertGreaterEqual(app._summaries.inspect_calls, 5)
                names = [column["name"] for column in app._tables[0]["columns"]]
                self.assertEqual(names, ["Column 1", "Column 2"])
                self.assertNotEqual(app.title_var.get(), "KEEP HEADER")

                bad.write_bytes(b"")
                app.header_var.set("自动")
                app._add_paths([bad, extra])
                _pump(app, 15)
                self.assertFalse(app._busy)
                values = [app.table_tree.item(iid, "values") for iid in app.table_tree.get_children()]
                bad_rows = [row for row in values if Path(row[0]).name == "bad.csv"]
                extra_rows = [row for row in values if Path(row[0]).name == "extra.csv"]
                self.assertTrue(bad_rows and "错误" in bad_rows[0][2] and bad_rows[0][4])
                self.assertTrue(extra_rows and extra_rows[0][4] == "就绪")
                self.assertTrue(any(table["source"]["path"].endswith("extra.csv") for table in app._tables))
                self.assertFalse(any(table["source"]["path"].endswith("bad.csv") for table in app._tables))
                with self.assertRaisesRegex(ValueError, "不完整"):
                    app._effective_plan(folder / "blocked.xlsx", "xlsx", False)
                self.assertEqual(str(app.export_button.cget("state")), "normal")
                print(
                    "MEASURE incremental "
                    f"inspect_calls={app._summaries.inspect_calls} "
                    f"cache_hits={app._summaries.cache_hits} "
                    f"cache_misses={app._summaries.cache_misses}"
                )
            finally:
                thread = getattr(app, "_thread", None)
                if thread is not None and thread.is_alive():
                    thread.join(timeout=3)
                try:
                    app.root.destroy()
                except tk.TclError:
                    pass

    @unittest.skipUnless(os.environ.get("SPECTRA_TEST_GUI") == "1", "set SPECTRA_TEST_GUI=1 for the Tk and spawn-export integration check")
    def test_real_tk_background_scan_and_chunked_refresh_stay_responsive_for_1000_sources(self):
        from origin_bridge import gui
        from origin_bridge.gui import GeneralDataApp
        from origin_bridge.readers import discover_files

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw) / "sources"
            folder.mkdir()
            for index in range(1000):
                (folder / f"sample{index}.csv").write_text("x,y\n1,2\n", encoding="utf-8")
            (folder / "notes.md").write_text("ignore", encoding="utf-8")
            try:
                app = GeneralDataApp()
            except tk.TclError as exc:
                self.skipTest(f"Tk display is unavailable: {exc}")
            samples = []
            after_ids: list[str] = []

            def heartbeat():
                samples.append((time.monotonic(), len(app.table_tree.get_children())))
                if app._busy:
                    after_ids.append(app.root.after(10, heartbeat))

            try:
                app.root.withdraw()
                warmup = time.monotonic() + 0.5
                while time.monotonic() < warmup:
                    app.root.update()
                    time.sleep(0.01)
                after_ids.append(app.root.after(10, heartbeat))
                with (
                    mock.patch.object(gui.messagebox, "showwarning") as showwarning,
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(gui.messagebox, "showinfo"),
                ):
                    app._add_paths([folder, folder / "notes.md"])
                    self.assertTrue(app._busy)
                    self.assertLess(len(app.table_tree.get_children()), 1000)
                    app._on_close()
                    self.assertTrue(app.root.winfo_exists())
                    deadline = time.monotonic() + 180
                    while app._busy and time.monotonic() < deadline:
                        app.root.update()
                        time.sleep(0.005)
                    app.root.update()
                    self.assertFalse(app._busy, "1000-source inspect did not finish")
                    showerror.assert_not_called()
                    showwarning.assert_called()
                names = [Path(app.table_tree.item(iid, "values")[0]).name for iid in app.table_tree.get_children()]
                self.assertEqual(names, [path.name for path in discover_files([folder])])
                self.assertEqual(len(names), 1000)
                self.assertEqual(app._summaries.inspect_calls, 1000)
                gaps = [samples[index][0] - samples[index - 1][0] for index in range(1, len(samples))]
                jumps = [samples[index][1] - samples[index - 1][1] for index in range(1, len(samples))]
                max_gap = max(gaps) if gaps else 999
                max_jump = max(jumps) if jumps else 1000
                partial = any(0 < count < 1000 for _stamp, count in samples)
                print(
                    "MEASURE tk1000 "
                    f"heartbeat_max_gap_s={max_gap:.4f} "
                    f"max_row_jump={max_jump} "
                    f"heartbeats={len(samples)} "
                    f"inspect_calls={app._summaries.inspect_calls} "
                    f"cache_hits={app._summaries.cache_hits}"
                )
                self.assertGreaterEqual(len(samples), 5)
                self.assertTrue(partial, "tree jumped to the full source list between heartbeats")
                self.assertLess(max_jump, 80, f"one heartbeat observed {max_jump} new rows")
                self.assertLess(max_gap, 0.25, f"Tk heartbeat stalled for {max_gap:.3f}s")
                self.assertEqual(str(app.export_button.cget("state")), "normal")
            finally:
                for after_id in after_ids:
                    try:
                        app.root.after_cancel(after_id)
                    except tk.TclError:
                        pass
                thread = getattr(app, "_thread", None)
                if thread is not None and thread.is_alive():
                    thread.join(timeout=3)
                try:
                    app.root.destroy()
                except tk.TclError:
                    pass


if __name__ == "__main__":
    unittest.main()
