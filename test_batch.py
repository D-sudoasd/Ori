"""Batch naming, isolation, resume, PDF, residency, and conversion reuse."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import Workbook, load_workbook

from origin_bridge import exporter
from origin_bridge.batch import (
    build_batch_request,
    execute_batch,
    load_batch_record,
    log_path_for,
    read_batch_log,
    retry_batch_tasks,
    save_batch_record,
    stable_output_path,
    task_id_for,
)
from origin_bridge.exporter import origin_export_lock
from origin_bridge.models import DataColumn, DataImportError, DataTable, PlannedTable, PlotSpec, PreparedImport
from origin_bridge.planning import create_plan, prepare_plan
from origin_bridge.session import OriginSession, OriginSessionLost, _save_graph_pdf, origin_pids, validate_pdf_file
from test_generic_exporter import FakeOrigin

ROOT = Path(__file__).resolve().parent
PDF = b"%PDF-1.4\n1 0 obj << /Type /Page >> endobj\nstartxref\n0\n%%EOF\n"


class FakeSession:
    def __init__(self, *, lose=False, pdf_error=False):
        self.starts = 0
        self.stops = 0
        self.owned_pids = [71001]
        self.lose = lose
        self.pdf_error = pdf_error
        self.pdf_requests: list[bool] = []
        self._on = False

    def start(self) -> None:
        self.starts += 1
        self._on = True

    def close(self) -> None:
        if self._on:
            self.stops += 1
        self._on = False

    def healthy(self) -> bool:
        return self._on

    def write_project(self, stage_dir, tables, prepared, converted, *, pdf=False, conversion_notes=None):
        del prepared, converted, conversion_notes
        if not self._on:
            raise OriginSessionLost("not running")
        if self.lose:
            self.lose = False
            raise OriginSessionLost("injected exit")
        self.pdf_requests.append(bool(pdf))
        stage = Path(stage_dir)
        opju = stage / "candidate.opju"
        opju.write_bytes(b"OPJU" + tables[0].name.encode("utf-8"))
        pdfs = []
        if pdf:
            if self.pdf_error:
                pdfs.append({
                    "ok": False, "staged": None, "table": tables[0].name,
                    "graph": "g", "error": "pdf boom",
                })
            else:
                for index, item in enumerate(tables, start=1):
                    if item.plot.kind == "none":
                        continue
                    target = stage / f"graph_{index:02d}.pdf"
                    target.write_bytes(PDF)
                    pdfs.append({
                        "ok": True, "staged": str(target), "table": item.name,
                        "graph": item.plot.title or item.name, "pages": 1,
                    })
        return {"opju": opju, "pdfs": pdfs, "warnings": [], "session_lost": False}


class BatchCoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ori_batch_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()

    def write_csv(self, path: Path, text: str = "x,y\n0,1\n1,2\n") -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def export(self, inputs, output, **kwargs):
        factory = kwargs.pop("session_factory", None)
        progress = kwargs.pop("progress", None)
        cancel = kwargs.pop("cancel_event", None)
        resume = kwargs.pop("resume", None)
        request = build_batch_request(list(inputs), output, **kwargs)
        return execute_batch(
            request,
            progress=progress,
            cancel_event=cancel,
            session_factory=factory,
            resume=resume,
        )

    def test_names_collisions_chinese_and_generated_files_are_not_inputs(self):
        long_source = Path("C:/data") / (("名" * 180) + ".csv")
        long_output = stable_output_path(long_source, self.root / "unused", ".xlsx")
        self.assertEqual(len(long_output.stem.split("__")[0]), 160)
        self.assertTrue(long_output.name.startswith("名"))
        self.assertEqual(len(task_id_for(long_source)), 16)

        left = self.write_csv(self.root / "left" / "same.csv", "x,y\n0,11\n")
        right = self.write_csv(self.root / "right" / "same.csv", "x,y\n0,22\n")
        chinese = self.write_csv(self.root / "扫描" / "光谱.csv", "x,y\n0,33\n")
        data = self.root / "scan"
        data.mkdir()
        sample = self.write_csv(data / "sample.csv", "x,y\n0,44\n")
        output = data / "out"
        output.mkdir()
        planted = output / "keep__0123456789ab.xlsx"
        planted.write_bytes(b"keep-me")
        loose = data / "loose__abcdefabcdef.opju"
        loose.write_bytes(b"loose")
        hidden = self.write_csv(output / "hidden.csv", "x,y\n0,99\n")

        collision = self.export([left.parent, right.parent, chinese.parent], self.root / "named", format="xlsx", plot={"kind": "none"})
        self.assertTrue(collision["ok"], collision)
        paths = [Path(item["output"]["path"]).name for item in collision["tasks"]]
        self.assertEqual(len(paths), 3)
        self.assertEqual(len(set(paths)), 3)
        self.assertTrue(any("光谱" in name for name in paths))
        self.assertTrue(any(name.startswith("same__") for name in paths))

        scanned = self.export([data], output, format="xlsx", plot={"kind": "none"})
        self.assertTrue(scanned["ok"], scanned)
        self.assertEqual([Path(item["source"]).name for item in scanned["tasks"]], ["sample.csv"])
        self.assertEqual(planted.read_bytes(), b"keep-me")
        self.assertEqual(loose.read_bytes(), b"loose")
        self.assertEqual(hidden.read_text(encoding="utf-8"), "x,y\n0,99\n")
        self.assertTrue(sample.is_file())

    def test_excel_sheets_share_one_output_and_unified_is_one_workbook(self):
        book_path = self.root / "book.xlsx"
        book = Workbook()
        alpha = book.active
        alpha.title = "Alpha"
        alpha.append(["x", "y"])
        alpha.append([1, 2])
        beta = book.create_sheet("Beta")
        beta.append(["x", "y"])
        beta.append([3, 4])
        book.save(book_path)
        book.close()

        per_file = self.export([book_path], self.root / "excel", format="xlsx", plot={"kind": "none"})
        self.assertEqual(len(per_file["tasks"]), 1, per_file)
        self.assertEqual(per_file["tasks"][0]["status"], "succeeded")
        saved = load_workbook(per_file["tasks"][0]["output"]["path"], data_only=False)
        try:
            found = [row for sheet in saved.worksheets for row in sheet.iter_rows(values_only=True)]
        finally:
            saved.close()
        self.assertIn((1, 2), found)
        self.assertIn((3, 4), found)

        first = self.write_csv(self.root / "uni" / "a.csv", "x,y\n0,5\n")
        second = self.write_csv(self.root / "uni" / "b.csv", "x,y\n0,6\n")
        unified = self.export(
            [first.parent], self.root / "unified",
            layout="unified", format="xlsx", plot={"kind": "none"}, output_name="combined.xlsx",
        )
        self.assertTrue(unified["ok"], unified)
        self.assertEqual(len(unified["tasks"]), 1)
        self.assertTrue(any("memory" in warning for warning in unified["warnings"]))
        combined = load_workbook(unified["tasks"][0]["output"]["path"])
        try:
            rows = [row for sheet in combined.worksheets for row in sheet.iter_rows(values_only=True)]
        finally:
            combined.close()
        self.assertIn((0, 5), rows)
        self.assertIn((0, 6), rows)

    def test_existing_output_is_not_overwritten_and_resume_skips_only_matching_bytes(self):
        source = self.write_csv(self.root / "once.csv")
        output = self.root / "once-out"
        first = self.export([source], output, format="xlsx", plot={"kind": "none"})
        self.assertTrue(first["ok"], first)
        target = Path(first["tasks"][0]["output"]["path"])
        original = target.read_bytes()
        record_bytes = Path(first["record_path"]).read_bytes()

        skipped = self.export([source], output, format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertTrue(skipped["ok"], skipped)
        self.assertEqual(skipped["counts"]["skipped"], 1)
        self.assertEqual(target.read_bytes(), original)
        self.assertEqual(Path(first["record_path"]).read_bytes(), record_bytes)

        blocked = self.export([source], output, format="xlsx", plot={"kind": "none"})
        self.assertFalse(blocked["ok"])
        self.assertEqual(blocked["counts"]["blocked"], 1)
        self.assertEqual(target.read_bytes(), original)

    def test_resume_sees_same_size_content_change_replaced_empty_and_missing_record(self):
        source = self.write_csv(self.root / "drift.csv", "x,y\n1,10\n2,20\n")
        output = self.root / "drift-out"
        first = self.export([source], output, format="xlsx", plot={"kind": "none"})
        target = Path(first["tasks"][0]["output"]["path"])
        original = target.read_bytes()
        stamp = source.stat().st_mtime_ns
        source_bytes = source.read_bytes()
        mutated = source_bytes.replace(b"1,10", b"1,11").replace(b"2,20", b"2,19")
        self.assertEqual(len(mutated), len(source_bytes))
        self.assertNotEqual(mutated, source_bytes)
        source.write_bytes(mutated)
        os.utime(source, ns=(stamp, stamp))
        self.assertEqual(source.stat().st_size, len(source_bytes))

        changed = self.export([source], output, format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertEqual(changed["tasks"][0]["status"], "blocked", changed)
        self.assertNotEqual(changed["tasks"][0]["status"], "skipped")
        self.assertEqual(target.read_bytes(), original)

        source.write_text("x,y\n1,10\n2,20\n", encoding="utf-8")
        os.utime(source, ns=(stamp, stamp))
        replaced = target.read_bytes()
        target.write_bytes(b"replaced-output")
        rewritten = self.export([source], output, format="xlsx", plot={"kind": "none"}, overwrite=True, resume=True)
        self.assertEqual(rewritten["tasks"][0]["status"], "succeeded", rewritten)
        self.assertNotEqual(target.read_bytes(), b"replaced-output")
        self.assertTrue(zip_is_xlsx(target))

        emptied = self.export([source], self.root / "empty-out", format="xlsx", plot={"kind": "none"})
        empty_target = Path(emptied["tasks"][0]["output"]["path"])
        empty_target.write_bytes(b"")
        empty_again = self.export([source], self.root / "empty-out", format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertEqual(empty_again["tasks"][0]["status"], "blocked", empty_again)
        self.assertEqual(empty_target.read_bytes(), b"")

        installed = self.export([source], self.root / "manifest-out", format="xlsx", plot={"kind": "none"})
        installed_target = Path(installed["tasks"][0]["output"]["path"])
        installed_bytes = installed_target.read_bytes()
        record = Path(installed["record_path"])
        record.unlink()
        backup = Path(str(record) + ".bak")
        if backup.exists():
            backup.unlink()
        missing = self.export([source], self.root / "manifest-out", format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertEqual(missing["tasks"][0]["status"], "blocked", missing)
        self.assertEqual(installed_target.read_bytes(), installed_bytes)
        self.assertNotEqual(replaced, b"")

    def test_corrupt_record_recovery_and_unsaved_success(self):
        first = self.write_csv(self.root / "rec" / "a.csv", "x,y\n0,1\n")
        second = self.write_csv(self.root / "rec" / "b.csv", "x,y\n0,2\n")
        output = self.root / "rec-out"
        done = self.export([first.parent], output, format="xlsx", plot={"kind": "none"})
        self.assertTrue(done["ok"], done)
        record_path = Path(done["record_path"])
        backup_path = Path(str(record_path) + ".bak")
        primary = load_batch_record(record_path)
        backup = load_batch_record(backup_path)
        self.assertEqual(len(primary["tasks"]), 2)
        self.assertEqual(len(backup["tasks"]), 1)
        first_id = done["tasks"][0]["task_id"]
        self.assertIn(first_id, backup["tasks"])
        bytes_before = {Path(item["output"]["path"]).read_bytes() for item in done["tasks"]}
        record_path.write_bytes(b"{")
        recovered = self.export([first.parent], output, format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertTrue(any("backup" in warning for warning in recovered["warnings"]))
        statuses = {item["task_id"]: item["status"] for item in recovered["tasks"]}
        self.assertEqual(statuses[first_id], "skipped")
        other_id = done["tasks"][1]["task_id"]
        self.assertEqual(statuses[other_id], "blocked")
        self.assertEqual({Path(item["output"]["path"]).read_bytes() for item in done["tasks"]}, bytes_before)

        both = self.export([first], self.root / "both-out", format="xlsx", plot={"kind": "none"})
        both_target = Path(both["tasks"][0]["output"]["path"])
        both_bytes = both_target.read_bytes()
        both_record = Path(both["record_path"])
        both_record.write_bytes(b"{")
        Path(str(both_record) + ".bak").write_bytes(b"{")
        ruined = self.export([first], self.root / "both-out", format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertTrue(any("corrupt" in warning for warning in ruined["warnings"]))
        self.assertEqual(ruined["tasks"][0]["status"], "blocked")
        self.assertEqual(both_target.read_bytes(), both_bytes)

        lone = self.write_csv(self.root / "unsaved.csv")
        calls = {"n": 0}
        real_save = save_batch_record

        def fail_second(path, record):
            calls["n"] += 1
            if calls["n"] >= 2:
                raise OSError("disk full")
            real_save(path, record)

        with mock.patch("origin_bridge.batch.save_batch_record", side_effect=fail_second):
            failed_save = self.export([lone], self.root / "unsaved-out", format="xlsx", plot={"kind": "none"})
        self.assertEqual(failed_save["tasks"][0]["status"], "failed", failed_save)
        self.assertTrue(any("record was not saved" in warning for warning in failed_save["warnings"]))
        self.assertTrue(Path(failed_save["tasks"][0]["output"]["path"]).is_file())
        self.assertEqual(load_batch_record(failed_save["record_path"]).get("tasks"), {})

    def test_failure_stays_on_one_task_and_cancellation_keeps_finished_output(self):
        good = self.write_csv(self.root / "iso" / "good.csv", "x,y\n0,7\n")
        empty = self.root / "iso" / "empty.csv"
        empty.write_bytes(b"")
        isolated = self.export([good.parent], self.root / "iso-out", format="xlsx", plot={"kind": "none"})
        by_name = {Path(item["source"]).name: item for item in isolated["tasks"]}
        self.assertEqual(by_name["good.csv"]["status"], "succeeded")
        self.assertEqual(by_name["empty.csv"]["status"], "failed")
        self.assertIn("空", by_name["empty.csv"]["error"]["message"])
        self.assertTrue(Path(by_name["good.csv"]["output"]["path"]).is_file())
        self.assertFalse(isolated["ok"])

        tiny = self.write_csv(self.root / "tiny.csv", "x,y\n0,1e-400\n")
        rejected = self.export([tiny], self.root / "tiny-out", format="xlsx", plot={"kind": "none"})
        self.assertEqual(rejected["tasks"][0]["status"], "failed")
        self.assertIn("underflow", rejected["tasks"][0]["error"]["message"])
        self.assertIsNone(rejected["tasks"][0]["output"])

        paths = [self.write_csv(self.root / "cancel" / f"{index}.csv") for index in range(3)]
        cancel = threading.Event()

        def progress(event):
            if event.get("type") == "task" and event.get("status") == "succeeded":
                cancel.set()

        stopped = self.export(
            [paths[0].parent], self.root / "cancel-out", format="xlsx", plot={"kind": "none"},
            progress=progress, cancel_event=cancel,
        )
        self.assertTrue(stopped["cancelled"])
        self.assertGreaterEqual(stopped["counts"]["succeeded"], 1)
        self.assertGreaterEqual(stopped["counts"]["cancelled"], 1)
        self.assertTrue(any(Path(item["output"]["path"]).is_file() for item in stopped["tasks"] if item["status"] == "succeeded"))

        def broken(_event):
            raise RuntimeError("callback failed on purpose")

        warned = self.export([good], self.root / "callback-out", format="xlsx", plot={"kind": "none"}, progress=broken)
        self.assertTrue(warned["ok"], warned)
        self.assertTrue(any("progress callback failed" in warning for warning in warned["warnings"]))

    def test_plot_override_bad_column_and_pdf_outcomes(self):
        good = self.write_csv(self.root / "plot" / "good.csv", "x,y\n0,1\n1,2\n")
        bad = self.write_csv(self.root / "plot" / "bad.csv", "x,y\n0,1\n1,2\n")
        selected = self.export(
            [good.parent], self.root / "plot-out", format="xlsx",
            plot={"kind": "none"},
            task_overrides=[{
                "source": bad,
                "plot": {"kind": "line", "x": "x", "y": ["missing"]},
            }],
        )
        by_name = {Path(item["source"]).name: item for item in selected["tasks"]}
        self.assertEqual(by_name["good.csv"]["status"], "succeeded")
        self.assertEqual(by_name["good.csv"]["resolved_plot"][0]["kind"], "none")
        self.assertEqual(by_name["bad.csv"]["status"], "failed")
        self.assertIn("missing", by_name["bad.csv"]["error"]["message"])

        explicit = self.export(
            [good], self.root / "index-out", format="xlsx",
            plot={"kind": "line", "x": 0, "y": [1]},
        )
        self.assertEqual(explicit["tasks"][0]["status"], "succeeded", explicit)
        self.assertEqual(explicit["tasks"][0]["resolved_plot"][0]["x"], 0)
        self.assertEqual(explicit["tasks"][0]["resolved_plot"][0]["y"], [1])

        sessions: list[FakeSession] = []

        def factory():
            session = FakeSession()
            sessions.append(session)
            return session

        pdf_ok = self.export(
            [good], self.root / "pdf-ok", format="opju", pdf=True,
            plot={"kind": "line", "x": "x", "y": ["y"], "title": "Good graph"},
            session_factory=factory,
        )
        self.assertTrue(pdf_ok["ok"], pdf_ok)
        self.assertEqual(pdf_ok["tasks"][0]["pdf_status"], "ok")
        pdf_path = Path(pdf_ok["tasks"][0]["pdfs"][0]["path"])
        self.assertIn("__g01_", pdf_path.name)
        self.assertGreaterEqual(validate_pdf_file(pdf_path)["pages"], 1)
        self.assertEqual(pdf_ok["origin"]["starts"], 1)
        self.assertEqual(pdf_ok["origin"]["stops"], 1)
        again = self.export(
            [good], self.root / "pdf-ok", format="opju", pdf=True,
            plot={"kind": "line", "x": "x", "y": ["y"], "title": "Good graph"},
            session_factory=factory, resume=True,
        )
        self.assertEqual(again["counts"]["skipped"], 1)

        failed_pdf = self.export(
            [good], self.root / "pdf-bad", format="opju", pdf=True,
            plot={"kind": "line", "x": "x", "y": ["y"]},
            session_factory=lambda: FakeSession(pdf_error=True),
        )
        self.assertFalse(failed_pdf["ok"])
        self.assertEqual(failed_pdf["tasks"][0]["status"], "succeeded")
        self.assertEqual(failed_pdf["tasks"][0]["pdf_status"], "failed")
        self.assertEqual(failed_pdf["counts"]["pdf_failed"], 1)
        self.assertTrue(Path(failed_pdf["tasks"][0]["output"]["path"]).is_file())
        self.assertEqual(list(Path(failed_pdf["output_dir"]).glob("*.pdf")), [])

        no_graph = self.export(
            [good], self.root / "pdf-none", format="opju", pdf=True,
            plot={"kind": "none"},
            session_factory=lambda: FakeSession(),
        )
        self.assertTrue(no_graph["ok"], no_graph)
        self.assertEqual(no_graph["tasks"][0]["pdf_status"], "not_applicable")
        self.assertEqual(no_graph["tasks"][0]["pdfs"], [])
        self.assertEqual(list(Path(no_graph["output_dir"]).glob("*.pdf")), [])
        with self.assertRaises(DataImportError):
            build_batch_request([good], self.root / "nope", format="xlsx", pdf=True)

    def test_one_origin_session_recycle_and_worker_loss(self):
        folder = self.root / "sess"
        paths = [self.write_csv(folder / f"{name}.csv") for name in ("a", "b", "c")]
        healthy_sessions: list[FakeSession] = []

        def healthy():
            session = FakeSession()
            healthy_sessions.append(session)
            return session

        batch = self.export(
            [folder], self.root / "healthy", format="opju", plot={"kind": "none"},
            session_factory=healthy,
        )
        self.assertTrue(batch["ok"], batch)
        self.assertEqual(batch["origin"]["starts"], 1)
        self.assertEqual(batch["origin"]["stops"], 1)
        self.assertEqual(len(healthy_sessions), 1)

        recycled_sessions: list[FakeSession] = []

        def recycled():
            session = FakeSession()
            recycled_sessions.append(session)
            return session

        recycled_result = self.export(
            [folder], self.root / "recycle", format="opju", plot={"kind": "none"},
            recycle_every=1, session_factory=recycled,
        )
        self.assertTrue(recycled_result["ok"], recycled_result)
        self.assertEqual(recycled_result["origin"]["starts"], 3)
        self.assertGreaterEqual(recycled_result["origin"]["recycled"], 2)
        self.assertEqual(len(recycled_sessions), 3)

        loss_sessions: list[FakeSession] = []

        def lossy():
            session = FakeSession(lose=not loss_sessions)
            loss_sessions.append(session)
            return session

        lost = self.export(
            [paths[0], paths[1]], self.root / "lost", format="opju", plot={"kind": "none"},
            origin_retries=1, session_factory=lossy,
        )
        self.assertTrue(lost["ok"], lost)
        self.assertGreaterEqual(lost["origin"]["starts"], 2)
        self.assertEqual(lost["counts"]["succeeded"], 2)

    def test_resident_cap_for_one_thousand_tasks(self):
        folder = self.root / "thousand"
        folder.mkdir()
        for index in range(999):
            (folder / f"f{index:04d}.csv").write_text("x,y\n0,1\n", encoding="utf-8")
        (folder / "zzz_empty.csv").write_bytes(b"")
        result = self.export(
            [folder], self.root / "thousand-out", format="xlsx", plot={"kind": "none"},
            prefetch=1, max_resident_tasks=2,
        )
        self.assertEqual(result["counts"]["succeeded"], 999, result["counts"])
        self.assertEqual(result["counts"]["failed"], 1)
        self.assertEqual(result["origin"]["starts"], 0)
        self.assertGreaterEqual(result["resident"]["max_full_tasks"], 2)
        self.assertLessEqual(result["resident"]["max_full_tasks"], result["resident"]["limit"])
        self.assertFalse(result["ok"])

    def test_lock_reentry_and_other_process_contention(self):
        nested = False
        with origin_export_lock():
            with origin_export_lock():
                nested = True
        self.assertTrue(nested)
        code = (
            "from origin_bridge.exporter import origin_export_lock\n"
            "try:\n"
            "    with origin_export_lock():\n"
            "        print('ACQUIRED')\n"
            "except Exception as exc:\n"
            "    print(type(exc).__name__ + ': ' + str(exc))\n"
        )
        env = os.environ.copy()
        env["PYTHONPATH"] = str(ROOT) + os.pathsep + env.get("PYTHONPATH", "")
        with origin_export_lock():
            proc = subprocess.run(
                [sys.executable, "-c", code], cwd=ROOT, env=env,
                capture_output=True, text=True, encoding="utf-8", timeout=30,
            )
        self.assertIn("already active", proc.stdout, proc.stderr)
        self.assertNotIn("ACQUIRED", proc.stdout)

    def test_spawn_cancel_stops_later_tasks(self):
        folder = self.root / "spawn"
        paths = [self.write_csv(folder / f"{index}.csv") for index in range(3)]
        script = self.root / "spawn_parent.py"
        script.write_text(
            "import json, sys\n"
            "from multiprocessing import get_context\n"
            "from origin_bridge.batch import build_batch_request\n"
            "from origin_bridge.worker import batch_worker\n"
            "\n"
            "def main() -> None:\n"
            "    inputs = sys.argv[1:-1]\n"
            "    output = sys.argv[-1]\n"
            "    request = build_batch_request(inputs, output, format='xlsx', plot={'kind': 'none'})\n"
            "    ctx = get_context('spawn')\n"
            "    recv, send = ctx.Pipe(duplex=False)\n"
            "    cancel = ctx.Event()\n"
            "    proc = ctx.Process(target=batch_worker, args=(send, request, cancel))\n"
            "    proc.start()\n"
            "    send.close()\n"
            "    result = None\n"
            "    seen_success = False\n"
            "    while True:\n"
            "        if not recv.poll(60):\n"
            "            proc.terminate()\n"
            "            raise SystemExit('spawn worker sent no message')\n"
            "        message = recv.recv()\n"
            "        if message['type'] == 'progress':\n"
            "            event = message['event']\n"
            "            if event.get('type') == 'task' and event.get('status') == 'succeeded' and not seen_success:\n"
            "                seen_success = True\n"
            "                cancel.set()\n"
            "        elif message['type'] == 'result':\n"
            "            result = message['result']\n"
            "            break\n"
            "    proc.join(60)\n"
            "    print(json.dumps({'seen_success': seen_success, 'result': result}))\n"
            "\n"
            "if __name__ == '__main__':\n"
            "    main()\n",
            encoding="utf-8",
        )
        env = os.environ.copy()
        env["PYTHONPATH"] = str(ROOT) + os.pathsep + env.get("PYTHONPATH", "")
        proc = subprocess.run(
            [sys.executable, str(script), *map(str, paths), str(self.root / "spawn-out")],
            cwd=ROOT, env=env, capture_output=True, text=True, encoding="utf-8", timeout=90,
        )
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        payload = json.loads(proc.stdout)
        self.assertTrue(payload["seen_success"])
        self.assertGreaterEqual(payload["result"]["counts"]["succeeded"], 1)
        self.assertGreaterEqual(payload["result"]["counts"]["cancelled"], 1)
        self.assertFalse(payload["result"]["ok"])

    def test_cli_batch_json_and_resume(self):
        source = self.write_csv(self.root / "cli.csv", "x,y\n0,4\n")
        output = self.root / "cli-out"
        first = run_cli("batch", "-i", source, "-o", output, "--format", "xlsx", "--plot", "none")
        self.assertEqual(first.returncode, 0, first.stderr)
        payload = json.loads(first.stdout)
        self.assertTrue(payload["ok"])
        target = Path(payload["tasks"][0]["output"]["path"])
        original = target.read_bytes()
        resumed = run_cli("batch", "-i", source, "-o", output, "--format", "xlsx", "--plot", "none", "--resume")
        self.assertEqual(resumed.returncode, 0, resumed.stdout + resumed.stderr)
        self.assertEqual(json.loads(resumed.stdout)["counts"]["skipped"], 1)
        self.assertEqual(target.read_bytes(), original)
        second = run_cli("batch", "-i", source, "-o", output, "--format", "xlsx", "--plot", "none")
        self.assertEqual(second.returncode, 4, second.stdout + second.stderr)
        blocked = json.loads(second.stdout)
        self.assertFalse(blocked["ok"])
        self.assertEqual(blocked["counts"]["blocked"], 1)
        self.assertEqual(target.read_bytes(), original)
        help_text = subprocess.run(
            [sys.executable, "-m", "origin_bridge", "batch", "--help"],
            cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=30,
        )
        self.assertEqual(help_text.returncode, 0, help_text.stderr)
        self.assertIn("--resume", help_text.stdout)

    def test_record_retry_and_corrupt_log_line(self):
        path = self.root / "record.json"
        save_batch_record(path, {
            "tasks": {
                "failed-one": {"task_id": "failed-one", "status": "failed", "source": "a"},
                "kept": {"task_id": "kept", "status": "succeeded", "source": "b"},
            }
        })
        retry_batch_tasks(path, force_task_ids=["kept"])
        loaded = load_batch_record(path)
        self.assertNotIn("failed-one", loaded["tasks"])
        self.assertNotIn("kept", loaded["tasks"])
        log = log_path_for(path)
        log.write_text('{"type":"batch"}\n{not-json\n', encoding="utf-8")
        events = read_batch_log(log)
        self.assertEqual(events[0]["type"], "batch")
        self.assertEqual(events[1]["type"], "corrupt_log_line")

    def test_conversion_and_warnings_are_reused_without_dropping_integrity(self):
        source = self.write_csv(self.root / "convert.csv", "x,y\n0,1\n1,2\n")
        output = self.root / "convert.xlsx"
        plan = create_plan([source], output, format="xlsx", overwrite=True)
        for entry in plan["tables"]:
            entry["plot"]["kind"] = "none"
            entry["plot"].update(x=None, y=[], y_error={})
        prepared = prepare_plan(plan)
        calls: list[str] = []
        real = exporter._convert_column

        def counted(column, table_name, output_format):
            calls.append(column.name)
            return real(column, table_name, output_format)

        with mock.patch.object(exporter, "_convert_column", side_effect=counted):
            receipt = exporter.execute_import(prepared)
        self.assertTrue(receipt["ok"])
        self.assertEqual(calls, ["x", "y"])
        calls.clear()
        with mock.patch.object(exporter, "_convert_column", side_effect=counted):
            exporter.validate_prepared_import(prepared)
        self.assertEqual(sorted(calls), ["x", "y"])
        bad = self.write_csv(self.root / "under.csv", "x,y\n0,1e-400\n")
        bad_plan = create_plan([bad], self.root / "under.xlsx", format="xlsx")
        with self.assertRaisesRegex(DataImportError, "underflows to zero"):
            exporter.validate_prepared_import(prepare_plan(bad_plan))

    def test_session_reset_drops_previous_project_markers(self):
        origin = FakeOrigin()
        support = SimpleNamespace(
            OriginExportError=RuntimeError,
            origin_process_running=lambda: False,
            _load_originpro=lambda: origin,
            close_origin_app=lambda _op, started: None,
            _cleanup_origin_temp_dirs=lambda directories: None,
        )

        def save(_op, path, pending_cleanup):
            del pending_cleanup
            path.write_bytes(b"fake opju")
            return path

        support._save_origin_project = save
        session = OriginSession(timeout_s=0)
        session.op = origin
        session.started = True
        session._alive = True
        session._live = False
        session._closed = False
        with mock.patch.object(exporter, "_origin_support", return_value=support), \
             mock.patch.object(exporter, "_origin_format_codes", return_value={"number": 1, "text": 2, "datetime": 3}):
            first = self.prepared_marker(111, "First marker")
            session.write_project(self.root / "stage-a", first.tables, first, converted_for(first), pdf=False)
            self.assertEqual(origin.book[0].to_list(0)[0], 111)
            self.assertEqual(origin.graphs[0].lname, "First marker")
            second = self.prepared_marker(222, "Second marker")
            session.write_project(self.root / "stage-b", second.tables, second, converted_for(second), pdf=False)
        self.assertEqual(origin.book[0].to_list(0)[0], 222)
        self.assertNotIn(111, origin.book[0].to_list(0))
        self.assertEqual([graph.lname for graph in origin.graphs], ["Second marker"])

    def prepared_marker(self, marker: int, title: str) -> PreparedImport:
        source = self.write_csv(self.root / f"marker_{marker}.csv", f"v\n{marker}\n")
        column = DataColumn("v", "number", (str(marker),))
        table = DataTable(
            source=source, source_sheet=None, name=f"M{marker}", columns=(column,),
            source_hash=hashlib.sha256(source.read_bytes()).hexdigest(),
        )
        planned = PlannedTable(table, table.name, PlotSpec(kind="line", y=(0,), title=title))
        return PreparedImport((planned,), self.root / f"{marker}.opju", "opju", True, False, ())

    def test_origin_pid_lookup_stays_stable_across_threads(self):
        if sys.platform != "win32":
            self.skipTest("Windows process snapshot")
        errors: list[BaseException] = []
        found: list[set[int]] = []

        def worker() -> None:
            try:
                for _ in range(40):
                    found.append(set(origin_pids()))
            except BaseException as exc:
                errors.append(exc)

        threads = [threading.Thread(target=worker) for _ in range(4)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()
        self.assertFalse(errors, errors)
        self.assertEqual(len(found), 160)
        self.assertTrue(all(isinstance(item, set) for item in found))

    def test_save_fig_accepts_only_a_new_valid_pdf_at_the_requested_path(self):
        source = self.write_csv(self.root / "pdf_src.csv")
        column = DataColumn("x", "number", ("0",))
        table = DataTable(
            source=source, source_sheet=None, name="T", columns=(column,),
            source_hash=hashlib.sha256(source.read_bytes()).hexdigest(),
        )
        planned = PlannedTable(table, "T", PlotSpec(kind="line", y=(0,), title="T"))
        dest = self.root / "graph.pdf"

        class Graph:
            def __init__(self, behavior: str):
                self.name = "Graph1"
                self.lname = "T"
                self.behavior = behavior
                self.obj = SimpleNamespace(GetName=lambda: self.name)

            def save_fig(self, path, replace=False):
                del replace
                target = Path(path)
                if self.behavior == "identity":
                    self.name = "Changed"
                    target.write_bytes(PDF)
                    return ""
                if self.behavior == "missing":
                    return ""
                if self.behavior == "mismatch":
                    target.write_bytes(PDF)
                    return str(target.with_name("other.pdf"))
                if self.behavior == "invalid":
                    target.write_bytes(b"this is not a pdf header, just long enough bytes")
                    return str(target)
                target.write_bytes(PDF)
                return ""

        dest.write_bytes(PDF)
        with self.assertRaisesRegex(DataImportError, "existing file"):
            _save_graph_pdf(Graph("valid"), dest, planned)
        dest.unlink()
        info = _save_graph_pdf(Graph("valid"), dest, planned)
        self.assertGreaterEqual(info["pages"], 1)
        dest.unlink()
        with self.assertRaisesRegex(DataImportError, "did not create"):
            _save_graph_pdf(Graph("missing"), dest, planned)
        with self.assertRaisesRegex(DataImportError, "not the requested path"):
            _save_graph_pdf(Graph("mismatch"), dest, planned)
        if dest.exists():
            dest.unlink()
        with self.assertRaisesRegex(DataImportError, "PDF header"):
            _save_graph_pdf(Graph("invalid"), dest, planned)
        if dest.exists():
            dest.unlink()
        with self.assertRaisesRegex(DataImportError, "identity changed"):
            _save_graph_pdf(Graph("identity"), dest, planned)


def converted_for(prepared: PreparedImport):
    _tables, _output, converted, _warnings, _notes = exporter._validate_prepared(prepared)
    return converted


def zip_is_xlsx(path: Path) -> bool:
    book = load_workbook(path)
    book.close()
    return True


def run_cli(*args):
    return subprocess.run(
        [sys.executable, "-m", "origin_bridge", *map(str, args)],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8", timeout=60,
    )


if __name__ == "__main__":
    unittest.main()
