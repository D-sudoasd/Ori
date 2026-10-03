"""Review fixes: ownership, record durability, PDF recovery, lock, and explicit modes.

These tests do not start Origin. Process-control checks use injected handles.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import threading
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import Workbook

from origin_bridge.batch import (
    build_batch_request,
    execute_batch,
    load_batch_record,
    log_path_for,
    read_batch_log,
    save_batch_record,
)
from origin_bridge.cli import build_parser
from origin_bridge.exporter import origin_export_lock
from origin_bridge.models import DataImportError
from origin_bridge.session import (
    OriginSession,
    OwnedProcess,
    _origin_image,
    claim_sentinel_holder,
    pid_creation_filetime,
    sentinel_holders,
    validate_pdf_file,
)
from test_batch import PDF, FakeSession, run_cli

ROOT = Path(__file__).resolve().parent


class ReviewFixTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ori_batch_review_")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()

    def write_csv(self, path: Path, text: str = "x,y\n0,1\n1,2\n") -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def workbook(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        left = book.active
        left.title = "Left"
        left.append(["x", "y"])
        left.append([0, 1])
        right = book.create_sheet("Right")
        right.append(["x", "y"])
        right.append([2, 3])
        book.save(path)
        book.close()
        return path

    def export(self, inputs, output, **kwargs):
        factory = kwargs.pop("session_factory", None)
        progress = kwargs.pop("progress", None)
        cancel = kwargs.pop("cancel_event", None)
        resume = kwargs.pop("resume", None)
        if resume is not None:
            kwargs["resume"] = resume
        request = build_batch_request(list(inputs), output, **kwargs)
        return execute_batch(
            request, progress=progress, cancel_event=cancel, session_factory=factory, resume=resume,
        )

    def snapshots(self, record: Path) -> dict[str, bytes | None]:
        paths = {
            "primary": record,
            "bak": Path(str(record) + ".bak"),
            "previous": record.with_name(record.name + ".previous"),
            "log": log_path_for(record),
        }
        found = {}
        for name, path in paths.items():
            found[name] = path.read_bytes() if path.is_file() else None
        return found

    def test_failed_discovery_keeps_the_previous_success_record(self):
        source = self.write_csv(self.root / "kept.csv", "x,y\n0,7\n")
        output = self.root / "kept-out"
        first = self.export([source], output, format="xlsx", plot={"kind": "none"})
        self.assertTrue(first["ok"], first)
        target = Path(first["tasks"][0]["output"]["path"])
        original = target.read_bytes()
        record = Path(first["record_path"])
        before = self.snapshots(record)
        self.assertIsNotNone(before["primary"])

        missing = self.root / "missing-input"
        empty = self.root / "empty-input"
        empty.mkdir()
        generated_only = self.root / "generated-only"
        generated_only.mkdir()
        planted = generated_only / "keep__0123456789ab.xlsx"
        planted.write_bytes(b"leave-me")
        for inputs, layout in (
            ([missing], "per_file"),
            ([empty], "per_file"),
            ([generated_only], "per_file"),
            ([missing], "unified"),
            ([empty], "unified"),
            ([generated_only], "unified"),
        ):
            with self.assertRaises(DataImportError):
                self.export(inputs, output, layout=layout, format="xlsx", plot={"kind": "none"})
            self.assertEqual(self.snapshots(record), before)
            self.assertEqual(target.read_bytes(), original)
            self.assertEqual(planted.read_bytes(), b"leave-me")

        resumed = self.export([source], output, format="xlsx", plot={"kind": "none"}, resume=True)
        self.assertEqual(resumed["counts"]["skipped"], 1)
        self.assertEqual(target.read_bytes(), original)
        self.assertEqual(Path(resumed["record_path"]).read_bytes(), before["primary"])

    def test_explicit_generated_name_is_imported_and_scans_warn(self):
        explicit = self.workbook(self.root / "named" / "sample__abcdefabcdef.xlsx")
        result = self.export([explicit], self.root / "explicit-out", format="xlsx", plot={"kind": "none"})
        self.assertTrue(result["ok"], result)
        self.assertEqual(len(result["tasks"]), 1)
        self.assertFalse(any("generated batch output name" in warning for warning in result["warnings"]))
        self.assertEqual(explicit.read_bytes()[:2], b"PK")

        folder = self.root / "scan"
        folder.mkdir()
        kept = self.write_csv(folder / "kept.csv", "x,y\n0,4\n")
        generated = folder / "other__abcdefabcdef.xlsx"
        generated.write_bytes(b"generated")
        scanned = self.export([folder], self.root / "scan-out", format="xlsx", plot={"kind": "none"})
        self.assertEqual([Path(item["source"]).name for item in scanned["tasks"]], ["kept.csv"])
        self.assertTrue(any("generated batch output name" in warning for warning in scanned["warnings"]))
        self.assertIn(str(generated.resolve()), "\n".join(scanned["warnings"]))
        self.assertEqual(generated.read_bytes(), b"generated")
        self.assertTrue(kept.is_file())

    def test_pdf_failure_is_visible_and_retry_keeps_the_project(self):
        source = self.write_csv(self.root / "pdf.csv", "x,y\n0,1\n1,2\n")
        output = self.root / "pdf-out"
        plot = {"kind": "line", "x": "x", "y": ["y"], "title": "Only"}
        events = []
        failed = self.export(
            [source], output, format="opju", pdf=True, plot=plot,
            session_factory=lambda: FakeSession(pdf_error=True), progress=events.append,
        )
        task = failed["tasks"][0]
        self.assertEqual(task["status"], "succeeded")
        self.assertIsNone(task["error"])
        self.assertEqual(task["pdf_status"], "failed")
        self.assertEqual(task["pdf_errors"][0]["message"], "pdf boom")
        self.assertTrue(any(event.get("pdf_errors") and event["pdf_errors"][0]["message"] == "pdf boom" for event in events))
        stored = load_batch_record(failed["record_path"])["tasks"][task["task_id"]]
        self.assertIsNone(stored["error"])
        self.assertEqual(stored["pdf_errors"][0]["message"], "pdf boom")
        self.assertTrue(stored["graph_identities"])
        opju = Path(task["output"]["path"])
        original = opju.read_bytes()
        logged = read_batch_log(log_path_for(failed["record_path"]))
        self.assertTrue(any("pdf boom" in json.dumps(event, ensure_ascii=False) for event in logged))

        resumed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: FakeSession(),
        )
        self.assertEqual(resumed["tasks"][0]["pdf_status"], "ok", resumed)
        self.assertEqual(opju.read_bytes(), original)
        self.assertTrue(resumed["tasks"][0]["pdfs"])

        sessions: list[FakeSession] = []

        def factory():
            session = FakeSession()
            sessions.append(session)
            return session

        opju.write_bytes(original)
        Path(resumed["tasks"][0]["pdfs"][0]["path"]).unlink()
        load_batch_record(resumed["record_path"])
        retried = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            retry_task_ids=[task["task_id"]], session_factory=factory,
        )
        self.assertEqual(retried["tasks"][0]["pdf_status"], "ok", retried)
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(len(sessions[0].recovered), 1)
        self.assertIsNone(retried["tasks"][0]["error"])

    def test_partial_pdf_recovery_and_refuses_replaced_inputs(self):
        source = self.workbook(self.root / "two.xlsx")
        output = self.root / "two-out"
        plot = {"kind": "line", "x": "x", "y": ["y"]}
        sessions: list[FakeSession] = []

        def failing():
            session = FakeSession(pdf_fail_indexes={2})
            sessions.append(session)
            return session

        first = self.export(
            [source], output, format="opju", pdf=True, plot=plot, session_factory=failing,
        )
        task = first["tasks"][0]
        self.assertEqual(task["status"], "succeeded")
        self.assertIsNone(task["error"])
        self.assertEqual(task["pdf_status"], "failed")
        self.assertEqual(len(task["pdfs"]), 1)
        self.assertEqual(task["pdfs"][0]["index"], 1)
        self.assertIn("pdf boom", task["pdf_errors"][0]["message"])
        opju = Path(task["output"]["path"])
        original = opju.read_bytes()
        first_pdf = Path(task["pdfs"][0]["path"])
        first_bytes = first_pdf.read_bytes()

        def recovering():
            session = FakeSession()
            sessions.append(session)
            return session

        filled = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=recovering,
        )
        self.assertEqual(filled["tasks"][0]["pdf_status"], "ok", filled)
        self.assertEqual(sorted(item["index"] for item in filled["tasks"][0]["pdfs"]), [1, 2])
        self.assertEqual([item["index"] for item in sessions[-1].recovered[0]], [2])
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(first_pdf.read_bytes(), first_bytes)

        before_sessions = len(sessions)
        skipped = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=recovering,
        )
        self.assertEqual(skipped["counts"]["skipped"], 1)
        self.assertEqual(len(sessions), before_sessions)
        self.assertEqual(opju.read_bytes(), original)

        second_pdf = Path(next(item["path"] for item in filled["tasks"][0]["pdfs"] if item["index"] == 2))
        second_bytes = second_pdf.read_bytes()
        second_pdf.unlink()
        deleted = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=recovering,
        )
        self.assertEqual([item["index"] for item in sessions[-1].recovered[0]], [2])
        self.assertEqual(first_pdf.read_bytes(), first_bytes)
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(deleted["tasks"][0]["pdf_status"], "ok", deleted)
        self.assertNotEqual(second_pdf.read_bytes(), b"")

        record = Path(filled["record_path"])
        receipt = record.read_bytes()
        stranger = PDF + b"\n% stranger\n"
        second_pdf.write_bytes(stranger)
        before_block = len(sessions)
        blocked = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=recovering,
        )
        self.assertEqual(blocked["counts"]["blocked"], 1)
        self.assertIn("pdf bytes changed", blocked["tasks"][0]["error"]["message"])
        self.assertEqual(blocked["tasks"][0]["pdf_status"], "failed")
        self.assertEqual(record.read_bytes(), receipt)
        self.assertEqual(second_pdf.read_bytes(), stranger)
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(load_batch_record(record)["tasks"][task["task_id"]]["status"], "succeeded")
        self.assertEqual(len(sessions), before_block)

        second_pdf.write_bytes(second_bytes)
        changed_book = Workbook()
        left = changed_book.active
        left.title = "Left"
        left.append(["x", "y"])
        left.append([0, 9])
        right = changed_book.create_sheet("Right")
        right.append(["x", "y"])
        right.append([2, 8])
        changed_book.save(source)
        changed_book.close()
        changed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=recovering,
        )
        self.assertEqual(changed["tasks"][0]["status"], "blocked", changed)
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(sessions[-1].recovered, [])

    def test_replaced_project_or_config_is_not_a_pdf_fill(self):
        source = self.write_csv(self.root / "guard.csv", "x,y\n0,1\n1,2\n")
        output = self.root / "guard-out"
        plot = {"kind": "line", "x": "x", "y": ["y"], "title": "Guard"}
        sessions: list[FakeSession] = []

        def factory():
            session = FakeSession()
            sessions.append(session)
            return session

        first = self.export([source], output, format="opju", pdf=True, plot=plot, session_factory=factory)
        opju = Path(first["tasks"][0]["output"]["path"])
        pdf = Path(first["tasks"][0]["pdfs"][0]["path"])
        original = opju.read_bytes()
        pdf_bytes = pdf.read_bytes()
        opju.write_bytes(b"OPJU-REPLACED")
        replaced = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=factory,
        )
        self.assertEqual(replaced["tasks"][0]["status"], "blocked")
        self.assertEqual(opju.read_bytes(), b"OPJU-REPLACED")
        self.assertEqual(pdf.read_bytes(), pdf_bytes)
        self.assertFalse(sessions[-1].recovered)

        opju.write_bytes(original)
        retitled = self.export(
            [source], output, format="opju", pdf=True, plot={**plot, "title": "Other"},
            resume=True, session_factory=factory,
        )
        self.assertEqual(retitled["tasks"][0]["status"], "blocked")
        self.assertEqual(opju.read_bytes(), original)
        self.assertFalse(sessions[-1].recovered)

    def test_recovery_start_failure_keeps_the_success_receipt(self):
        source = self.write_csv(self.root / "dead.csv")
        output = self.root / "dead-out"
        plot = {"kind": "line", "x": "x", "y": ["y"], "title": "Dead"}
        first = self.export(
            [source], output, format="opju", pdf=True, plot=plot,
            session_factory=lambda: FakeSession(pdf_error=True),
        )
        opju = Path(first["tasks"][0]["output"]["path"])
        original = opju.read_bytes()

        class Dead(FakeSession):
            def start(self):
                raise DataImportError("Could not prove which Origin process belongs to this COM instance")

            def export_saved_pdfs(self, project, graphs):
                raise AssertionError("PDF export ran without a proved session")

        failed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=Dead,
        )
        self.assertEqual(failed["tasks"][0]["status"], "succeeded")
        self.assertIsNone(failed["tasks"][0]["error"])
        self.assertEqual(failed["tasks"][0]["pdf_status"], "failed")
        self.assertIn("Could not prove", failed["tasks"][0]["pdf_errors"][0]["message"])
        self.assertEqual(opju.read_bytes(), original)
        stored = load_batch_record(failed["record_path"])["tasks"][first["tasks"][0]["task_id"]]
        self.assertEqual(stored["status"], "succeeded")
        self.assertEqual(stored["output"]["sha256"], hashlib.sha256(original).hexdigest())
        self.assertIn("Could not prove", stored["pdf_errors"][0]["message"])

    def test_modes_that_cannot_be_honored_are_rejected(self):
        source = self.write_csv(self.root / "mode.csv")
        with self.assertRaisesRegex(DataImportError, "cannot export PDFs"):
            build_batch_request([source], self.root / "u1", layout="unified", format="opju", pdf=True, plot={"kind": "none"})
        with self.assertRaisesRegex(DataImportError, "does not support resume"):
            build_batch_request([source], self.root / "u2", layout="unified", format="xlsx", plot={"kind": "none"}, resume=True)
        with self.assertRaisesRegex(DataImportError, "retry_task_ids"):
            build_batch_request(
                [source], self.root / "u3", format="xlsx", plot={"kind": "none"}, retry_task_ids=["abc"],
            )
        built = build_batch_request([source], self.root / "u4", layout="unified", format="xlsx", plot={"kind": "none"})
        with self.assertRaisesRegex(DataImportError, "does not support resume"):
            execute_batch(built, resume=True)
        per_file = build_batch_request(
            [source], self.root / "u5", format="xlsx", plot={"kind": "none"}, resume=True, retry_task_ids=["abc"],
        )
        per_file["resume"] = False
        with self.assertRaisesRegex(DataImportError, "retry_task_ids"):
            execute_batch(per_file)
        unified_pdf = build_batch_request(
            [source], self.root / "u6", layout="unified", format="opju", plot={"kind": "none"},
        )
        unified_pdf["pdf"] = True
        with self.assertRaisesRegex(DataImportError, "cannot export PDFs"):
            execute_batch(unified_pdf)

        pdf_cli = run_cli(
            "batch", "-i", source, "-o", self.root / "cli-pdf", "--layout", "unified",
            "--format", "opju", "--pdf", "--plot", "none",
        )
        self.assertEqual(pdf_cli.returncode, 4, pdf_cli.stdout + pdf_cli.stderr)
        self.assertIn("cannot export PDFs", pdf_cli.stdout)
        resume_cli = run_cli(
            "batch", "-i", source, "-o", self.root / "cli-resume", "--layout", "unified",
            "--format", "xlsx", "--plot", "none", "--resume",
        )
        self.assertEqual(resume_cli.returncode, 4, resume_cli.stdout + resume_cli.stderr)
        self.assertIn("does not support resume", resume_cli.stdout)

        record = self.root / "retry-record.json"
        from origin_bridge.batch import save_batch_record

        save_batch_record(record, {"tasks": {
            "bad": {"task_id": "bad", "status": "failed", "source": "a"},
            "pdf": {"task_id": "pdf", "status": "succeeded", "pdf_status": "failed", "source": "b"},
        }})
        args = build_parser().parse_args([
            "batch", "-i", str(source), "-o", str(self.root / "cli-retry"),
            "--format", "xlsx", "--plot", "none", "--retry-failed", "--record", str(record),
        ])
        captured = []

        def capture(request, **kwargs):
            del kwargs
            captured.append(request)
            return {"ok": True, "tasks": []}

        with mock.patch("origin_bridge.batch.execute_batch", side_effect=capture):
            from origin_bridge.cli import _run_batch

            payload = _run_batch(args)
        self.assertTrue(payload["ok"])
        self.assertEqual(captured[0]["resume"], True)
        self.assertEqual(captured[0]["retry_task_ids"], ["bad"])

    def test_other_thread_gets_already_active_without_waiting(self):
        started = threading.Event()
        release = threading.Event()
        errors = []

        def hold():
            with origin_export_lock():
                started.set()
                self.assertTrue(release.wait(5))

        def other():
            began = time.perf_counter()
            try:
                with origin_export_lock():
                    errors.append("acquired")
            except Exception as exc:
                errors.append((time.perf_counter() - began, str(exc)))

        holder = threading.Thread(target=hold)
        holder.start()
        self.assertTrue(started.wait(5))
        competitor = threading.Thread(target=other)
        competitor.start()
        competitor.join(2)
        release.set()
        holder.join(5)
        self.assertFalse(competitor.is_alive())
        self.assertEqual(len(errors), 1)
        elapsed, message = errors[0]
        self.assertLess(elapsed, 1.0)
        self.assertIn("already active", message)
        try:
            with origin_export_lock():
                raise RuntimeError("boom")
        except RuntimeError:
            pass
        with origin_export_lock():
            nested = False
            with origin_export_lock():
                nested = True
        self.assertTrue(nested)

    def test_pdf_checker_rejects_objstm_and_accepts_an_uncompressed_page(self):
        plain = self.root / "plain.pdf"
        plain.write_bytes(PDF + b"\n/Filter /FlateDecode\n")
        self.assertGreaterEqual(validate_pdf_file(plain)["pages"], 1)
        objstm = self.root / "objstm.pdf"
        objstm.write_bytes(
            b"%PDF-1.4\n1 0 obj << /Type /ObjStm /N 1 >> stream\nzzzz\nendstream\nendobj\nstartxref\n0\n%%EOF\n"
        )
        with self.assertRaisesRegex(DataImportError, "ObjStm"):
            validate_pdf_file(objstm)
        compressed = self.root / "compressed.pdf"
        compressed.write_bytes(b"%PDF-1.4\n1 0 obj << /Filter /FlateDecode >> stream\nzzzz\nendstream\nendobj\nstartxref\n0\n%%EOF\n")
        with self.assertRaisesRegex(DataImportError, "uncompressed /Type /Page"):
            validate_pdf_file(compressed)

    def test_keep_open_and_recycle_report_every_proved_process(self):
        folder = self.root / "life"
        paths = [self.write_csv(folder / f"{name}.csv") for name in ("a", "b")]
        opened: list[FakeSession] = []

        def factory():
            session = FakeSession()
            session.owned_pids = [81000 + len(opened)]
            session.owned_identities = [{
                "pid": 81000 + len(opened),
                "creation_filetime": 5000 + len(opened),
                "image": "Origin64.exe",
            }]
            opened.append(session)
            return session

        held = self.export(
            [paths[0]], self.root / "held-out", format="opju", plot={"kind": "none"},
            keep_open=True, session_factory=factory,
        )
        self.assertTrue(held["ok"], held)
        self.assertEqual(held["origin"]["stops"], 0)
        self.assertEqual(held["origin"]["owned_pids"], [81000])
        self.assertEqual(held["origin"]["owned_processes"][0]["creation_filetime"], 5000)
        self.assertTrue(opened[0].healthy())
        opened[0].close()

        recycled: list[FakeSession] = []

        def recycled_factory():
            session = FakeSession()
            session.owned_pids = [82000 + len(recycled)]
            session.owned_identities = [{
                "pid": 82000 + len(recycled),
                "creation_filetime": 6000 + len(recycled),
                "image": "Origin64.exe",
            }]
            recycled.append(session)
            return session

        result = self.export(
            [folder], self.root / "recycle-out", format="opju", plot={"kind": "none"},
            recycle_every=1, keep_open=True, session_factory=recycled_factory,
        )
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["origin"]["starts"], 2)
        self.assertEqual(result["origin"]["stops"], 2)
        self.assertGreaterEqual(result["origin"]["recycled"], 1)
        self.assertEqual(result["origin"]["owned_pids"], [82000, 82001])
        self.assertEqual(
            [item["creation_filetime"] for item in result["origin"]["owned_processes"]],
            [6000, 6001],
        )

    def test_unknown_and_reused_pids_are_not_terminated(self):
        if os.name != "nt":
            self.skipTest("Windows process handles")
        session = OriginSession(timeout_s=0)
        session._terminate_handle = lambda handle: (_ for _ in ()).throw(AssertionError("terminated"))
        with mock.patch("origin_bridge.session.origin_pids", return_value={424242}), \
             mock.patch("origin_bridge.session.terminate_pid") as killer:
            self.assertEqual(session.terminate_owned(), [])
            session._deadline = time.monotonic() - 1
            session._enforce_deadline()
            killer.assert_not_called()

        session.started = True
        session._alive = True
        session._live = True
        session.op = object()
        session._owned = [OwnedProcess(42, 1000, 7, r"C:\Origin64.exe")]
        killed = []
        session._terminate_handle = lambda handle: killed.append(handle) or True
        session._handle_status = lambda handle: {"alive": True, "creation": 1000}
        session._pid_creation = lambda pid: 2000
        self.assertFalse(session.healthy())
        self.assertEqual(session.terminate_owned(), [])
        session._enforce_deadline()
        self.assertEqual(killed, [])

        session._pid_creation = lambda pid: 1000
        self.assertTrue(session.healthy())
        with mock.patch("origin_bridge.session.terminate_pid") as killer, \
             mock.patch("origin_bridge.session.origin_pids", return_value={42, 99}):
            session._deadline = time.monotonic() - 1
            session._enforce_deadline()
            killer.assert_not_called()
        self.assertEqual(killed, [7])

        closed = []
        session._close_handle = lambda handle: closed.append(handle)
        session._pid_creation = lambda pid: None
        session.op = None
        session._live = False
        session.close()
        session.close()
        self.assertEqual(closed, [7])

        self.assertIsNone(claim_sentinel_holder([], is_origin_image=lambda pid, image: True))
        self.assertIsNone(claim_sentinel_holder(
            [
                {"pid": 10, "creation_filetime": 1, "image": "Origin64.exe"},
                {"pid": 11, "creation_filetime": 2, "image": "Origin64.exe"},
            ],
            is_origin_image=lambda pid, image: True,
        ))
        claimed = claim_sentinel_holder(
            [{"pid": os.getpid(), "creation_filetime": 5, "image": "friendly name"}],
            is_origin_image=lambda pid, image: image == "friendly name",
        )
        self.assertEqual(claimed["pid"], os.getpid())
        self.assertIsNone(claim_sentinel_holder(
            [{"pid": os.getpid(), "creation_filetime": pid_creation_filetime(os.getpid()) or 1, "image": "Origin"}],
            is_origin_image=_origin_image,
        ))

    def test_restart_manager_sees_this_process_holding_a_file(self):
        if os.name != "nt":
            self.skipTest("Windows Restart Manager")
        directory = Path(tempfile.mkdtemp(prefix="ori_rm_", dir=self.root))
        path = directory / "held.bin"
        path.write_bytes(b"held-by-this-process")
        stream = path.open("rb")
        try:
            holders = sentinel_holders(path)
        finally:
            stream.close()
        ours = [item for item in holders if item["pid"] == os.getpid()]
        self.assertTrue(ours, holders)
        self.assertEqual(ours[0]["creation_filetime"], pid_creation_filetime(os.getpid()))
        self.assertIsNone(claim_sentinel_holder(holders, is_origin_image=_origin_image))

    def test_unproved_launch_does_not_terminate_a_candidate(self):
        class LiveOp:
            def __init__(self):
                self.shown = None
                self.exited = False
                self.value = ""

            def set_show(self, value):
                self.shown = value

            def set_lt_str(self, name, value):
                del name
                self.value = value
                return True

            def get_lt_str(self, name):
                del name
                return self.value

            def lt_exec(self, script):
                del script
                return True

            def lt_int(self, expr):
                del expr
                return 1

            def exit(self):
                self.exited = True

        LiveOp.__module__ = "originpro"
        op = LiveOp()
        closed = []
        support = SimpleNamespace(
            OriginExportError=RuntimeError,
            origin_process_running=lambda: False,
            _load_originpro=lambda: op,
            close_origin_app=lambda app, started: closed.append((app, started)) or app.exit(),
            _cleanup_origin_temp_dirs=lambda directories: None,
        )
        holders = [
            {"pid": 111, "creation_filetime": 10, "image": "Origin64.exe"},
            {"pid": 222, "creation_filetime": 20, "image": "Origin64.exe"},
        ]
        session = OriginSession(timeout_s=0)
        with mock.patch("origin_bridge.exporter._origin_support", return_value=support), \
             mock.patch("origin_bridge.session.origin_pids", return_value=set()), \
             mock.patch("origin_bridge.session.sentinel_holders", return_value=holders), \
             mock.patch("origin_bridge.session.terminate_pid") as killer, \
             mock.patch("origin_bridge.session.terminate_process_handle") as handle_killer:
            with self.assertRaisesRegex(DataImportError, "no candidate process was terminated"):
                session.start()
            killer.assert_not_called()
            handle_killer.assert_not_called()
        self.assertIs(op.shown, False)
        self.assertEqual(closed, [(op, True)])
        self.assertTrue(op.exited)
        self.assertEqual(session.owned_pids, set())

    def test_recover_fail_indexes_keeps_the_project_and_a_later_resume_fills(self):
        source = self.workbook(self.root / "retry-one.xlsx")
        output = self.root / "retry-one-out"
        plot = {"kind": "line", "x": "x", "y": ["y"]}
        first = self.export(
            [source], output, format="opju", pdf=True, plot=plot,
            session_factory=lambda: FakeSession(pdf_fail_indexes={2}),
        )
        task = first["tasks"][0]
        self.assertEqual(task["pdf_status"], "failed")
        opju = Path(task["output"]["path"])
        original = opju.read_bytes()
        kept = Path(task["pdfs"][0]["path"])
        kept_bytes = kept.read_bytes()
        sessions: list[FakeSession] = []

        def failing_recover():
            session = FakeSession()
            session.recover_fail_indexes.add(2)
            sessions.append(session)
            return session

        failed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=failing_recover,
        )
        failed_task = failed["tasks"][0]
        self.assertEqual(failed_task["status"], "succeeded")
        self.assertIsNone(failed_task["error"])
        self.assertEqual(failed_task["pdf_status"], "failed")
        self.assertTrue(failed_task["pdf_errors"])
        self.assertIn("recover failed", failed_task["pdf_errors"][0]["message"])
        self.assertEqual([item["index"] for item in failed_task["pdfs"]], [1])
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertEqual([item["index"] for item in sessions[0].recovered[0]], [2])
        missing = Path(failed_task["pdf_errors"][0]["path"])
        self.assertFalse(missing.is_file())

        filled = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: FakeSession(),
        )
        self.assertEqual(filled["tasks"][0]["pdf_status"], "ok", filled)
        self.assertEqual(sorted(item["index"] for item in filled["tasks"][0]["pdfs"]), [1, 2])
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertTrue(missing.is_file())

    def test_record_without_graph_identities_uses_resolved_plot(self):
        source = self.write_csv(self.root / "legacy.csv", "x,y\n0,1\n1,4\n")
        output = self.root / "legacy-out"
        plot = {"kind": "line", "x": "x", "y": ["y"], "title": "Legacy"}
        first = self.export(
            [source], output, format="opju", pdf=True, plot=plot, session_factory=lambda: FakeSession(),
        )
        task = first["tasks"][0]
        opju = Path(task["output"]["path"])
        original = opju.read_bytes()
        pdf = Path(task["pdfs"][0]["path"])
        pdf.unlink()
        record_path = Path(first["record_path"])
        stored = load_batch_record(record_path)
        entry = stored["tasks"][task["task_id"]]
        self.assertTrue(entry.get("graph_identities"))
        resolved = [item for item in entry["resolved_plot"] if item.get("kind") != "none"]
        self.assertEqual(len(resolved), 1)
        entry.pop("graph_identities")
        save_batch_record(record_path, stored)
        sessions: list[FakeSession] = []

        def factory():
            session = FakeSession()
            sessions.append(session)
            return session

        filled = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True, session_factory=factory,
        )
        self.assertEqual(filled["tasks"][0]["pdf_status"], "ok", filled)
        self.assertEqual(opju.read_bytes(), original)
        requested = sessions[-1].recovered[0][0]
        self.assertEqual(requested["table"], resolved[0]["table"])
        self.assertEqual(requested["graph"], resolved[0]["title"] or resolved[0]["table"])
        self.assertEqual(requested["short_name"], "")
        self.assertGreaterEqual(validate_pdf_file(pdf)["pages"], 1)

    def test_snapshot_copy_leaves_a_mismatched_source_in_place(self):
        from origin_bridge import batch as batch_mod

        stage = getattr(batch_mod, "_stage_recovery_snapshot", None)
        if stage is None:
            self.fail("recovery does not stage a verified snapshot")
        src = self.root / "source.opju"
        payload = b"ORIGINAL-OPJU-BYTES"
        src.write_bytes(payload)
        expected = hashlib.sha256(payload).hexdigest()
        src.write_bytes(b"FOREIGN_DURING_COPY")
        dest = self.root / "snap" / "snapshot.opju"
        with self.assertRaises(DataImportError) as caught:
            stage(src, dest, expected)
        self.assertEqual(src.read_bytes(), b"FOREIGN_DURING_COPY")
        self.assertNotIn("restored", str(caught.exception).lower())
        self.assertIn("left in place", str(caught.exception))

    def test_export_saved_pdfs_rejects_open_provenance_and_ambiguity_without_restoring(self):
        project = self.root / "direct.opju"
        original = b"OPJU-ORIGINAL-BYTES-0123456789"
        project.write_bytes(original)
        source = self.write_csv(self.root / "direct.csv")
        source_name = str(source.resolve())
        dest = self.root / "direct.pdf"
        request = [{
            "table": "direct",
            "graph": "direct",
            "short_name": "G1",
            "source_name": source_name,
            "dest": str(dest),
        }]

        class Closed:
            def open(self, path, *_args):
                return False

            def pages(self, _kind):
                raise AssertionError("pages were read after Origin refused to open")

        session = _armed_session(Closed())
        with self.assertRaisesRegex(DataImportError, "did not open"):
            session.export_saved_pdfs(project, request)
        self.assertEqual(project.read_bytes(), original)
        self.assertFalse(dest.exists())

        class MutatingClosed:
            def open(self, path, *_args):
                Path(path).write_bytes(b"FOREIGN_OPEN_FAIL")
                return False

            def pages(self, _kind):
                raise AssertionError("pages were read after Origin refused to open")

        project.write_bytes(original)
        session = _armed_session(MutatingClosed())
        with self.assertRaises(DataImportError) as caught:
            session.export_saved_pdfs(project, request)
        self.assertNotIn("restored", str(caught.exception).lower())
        self.assertIn("left in place", str(caught.exception))
        self.assertEqual(project.read_bytes(), b"FOREIGN_OPEN_FAIL")
        self.assertFalse(dest.exists())

        project.write_bytes(original)
        decoy = _Graph("G1", "direct")
        com = _Com(
            books=[_Book([
                _Sheet("direct", [[]]),
                _Sheet("Import Provenance", [["unrelated-source"]]),
            ])],
            graphs=[decoy],
        )
        session = _armed_session(com)
        with self.assertRaisesRegex(DataImportError, "provenance"):
            session.export_saved_pdfs(project, request)
        self.assertEqual(decoy.saves, [])
        self.assertEqual(project.read_bytes(), original)
        self.assertFalse(dest.exists())

        project.write_bytes(original)
        wrong = _Graph("G9", "not-the-recorded-graph")
        com = _Com(
            books=[_Book([
                _Sheet("direct", [[]]),
                _Sheet("Import Provenance", [[source_name]]),
            ])],
            graphs=[wrong],
        )
        session = _armed_session(com)
        refused = session.export_saved_pdfs(project, request)
        self.assertEqual(len(refused), 1)
        self.assertFalse(refused[0]["ok"])
        self.assertIn("unique", refused[0]["error"])
        self.assertEqual(wrong.saves, [])
        self.assertEqual(project.read_bytes(), original)
        self.assertFalse(dest.exists())

        project.write_bytes(original)
        first = _Graph("G1", "direct")
        second = _Graph("G1", "direct")
        com = _Com(
            books=[_Book([
                _Sheet("direct", [[]]),
                _Sheet("direct", [[]]),
                _Sheet("Import Provenance", [[source_name]]),
            ])],
            graphs=[first, second],
        )
        session = _armed_session(com)
        ambiguous = session.export_saved_pdfs(project, request)
        self.assertFalse(ambiguous[0]["ok"])
        self.assertIn("unique", ambiguous[0]["error"])
        self.assertEqual(first.saves, [])
        self.assertEqual(second.saves, [])
        self.assertEqual(project.read_bytes(), original)
        self.assertFalse(dest.exists())

    def test_foreign_update_during_recovery_is_left_in_place(self):
        source, output, plot, opju, original, kept, kept_bytes, record, receipt, missing = self._partial_pdf(self.root / "foreign")
        seen = []

        def on_open(path: Path) -> None:
            seen.append(Path(path).read_bytes())
            opju.write_bytes(b"FOREIGN_UPDATE")

        com = _com_for_record(load_batch_record(record), source, on_open=on_open)
        resumed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: ProbeSession(com),
        )
        task = resumed["tasks"][0]
        message = task["error"]["message"] if task["error"] else ""
        self.assertEqual(task["status"], "blocked", resumed)
        self.assertEqual(task["pdf_status"], "failed", resumed)
        self.assertIn("left in place", message)
        self.assertNotIn("restored", message.lower())
        self.assertEqual(opju.read_bytes(), b"FOREIGN_UPDATE")
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertFalse(missing.is_file())
        self.assertEqual(record.read_bytes(), receipt)
        self.assertTrue(com.opened)
        self.assertTrue(all(not _same_path(path, opju) for path in com.opened))
        self.assertTrue(all(path.name == "snapshot.opju" for path in com.opened))
        self.assertEqual(seen, [original])
        self.assertFalse(list(opju.parent.glob("*.restore-tmp")))

    def test_project_replaced_before_snapshot_open_is_not_opened(self):
        source, output, plot, opju, _original, kept, kept_bytes, record, receipt, missing = self._partial_pdf(self.root / "before-open")

        def on_start() -> None:
            opju.write_bytes(b"FOREIGN_BEFORE_OPEN")

        com = _com_for_record(load_batch_record(record), source)
        resumed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: ProbeSession(com, on_start=on_start),
        )
        task = resumed["tasks"][0]
        message = task["error"]["message"] if task["error"] else ""
        self.assertEqual(task["status"], "blocked", resumed)
        self.assertEqual(task["pdf_status"], "failed")
        self.assertIn("left in place", message)
        self.assertNotIn("restored", message.lower())
        self.assertEqual(opju.read_bytes(), b"FOREIGN_BEFORE_OPEN")
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertFalse(missing.is_file())
        self.assertEqual(record.read_bytes(), receipt)
        self.assertEqual(com.opened, [])

    def test_project_replaced_after_snapshot_copy_is_not_opened(self):
        from origin_bridge import batch as batch_mod

        stage = getattr(batch_mod, "_stage_recovery_snapshot", None)
        if stage is None:
            self.fail("recovery does not stage a verified snapshot")
        source, output, plot, opju, _original, kept, kept_bytes, record, receipt, missing = self._partial_pdf(self.root / "during-copy")

        def wrapper(src, dest, expected):
            stage(src, dest, expected)
            Path(src).write_bytes(b"FOREIGN_DURING_COPY")

        com = _com_for_record(load_batch_record(record), source)
        with mock.patch.object(batch_mod, "_stage_recovery_snapshot", wrapper):
            resumed = self.export(
                [source], output, format="opju", pdf=True, plot=plot, resume=True,
                session_factory=lambda: ProbeSession(com),
            )
        task = resumed["tasks"][0]
        message = task["error"]["message"] if task["error"] else ""
        self.assertEqual(task["status"], "blocked", resumed)
        self.assertEqual(task["pdf_status"], "failed")
        self.assertIn("left in place", message)
        self.assertNotIn("restored", message.lower())
        self.assertEqual(opju.read_bytes(), b"FOREIGN_DURING_COPY")
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertFalse(missing.is_file())
        self.assertEqual(record.read_bytes(), receipt)
        self.assertEqual(com.opened, [])

    def test_project_replaced_before_install_is_not_installed(self):
        source, output, plot, opju, original, kept, kept_bytes, record, receipt, missing = self._partial_pdf(self.root / "before-install")

        def on_save(_dest) -> None:
            opju.write_bytes(b"FOREIGN_BEFORE_INSTALL")

        com = _com_for_record(load_batch_record(record), source, on_save=on_save)
        opened_bytes: list[bytes] = []
        original_open = com.open

        def open_and_remember(path, *args):
            opened_bytes.append(Path(path).read_bytes())
            return original_open(path, *args)

        com.open = open_and_remember
        resumed = self.export(
            [source], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: ProbeSession(com),
        )
        task = resumed["tasks"][0]
        message = task["error"]["message"] if task["error"] else ""
        self.assertEqual(task["status"], "blocked", resumed)
        self.assertEqual(task["pdf_status"], "failed")
        self.assertIn("left in place", message)
        self.assertNotIn("restored", message.lower())
        self.assertEqual(opju.read_bytes(), b"FOREIGN_BEFORE_INSTALL")
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertFalse(missing.is_file())
        self.assertEqual(record.read_bytes(), receipt)
        self.assertTrue(com.opened)
        self.assertTrue(all(not _same_path(path, opju) for path in com.opened))
        self.assertEqual(opened_bytes, [original])

    def test_com_error_dirties_only_the_snapshot_and_the_batch_continues(self):
        source, output, plot, opju, original, kept, kept_bytes, record, _receipt, missing = self._partial_pdf(self.root / "snapshot-error")
        follower = self.write_csv(self.root / "snapshot-error" / "next.csv", "x,y\n0,5\n1,6\n")
        com = _com_for_record(load_batch_record(record), source)

        def on_open(path: Path) -> None:
            if _same_path(path, opju):
                raise AssertionError("recovery opened the production project")
            path.write_bytes(path.read_bytes() + b"\nCOM")
            raise RuntimeError("com broke snapshot")

        com.on_open = on_open
        resumed = self.export(
            [source, follower], output, format="opju", pdf=True, plot=plot, resume=True,
            session_factory=lambda: ProbeSession(com),
        )
        by_name = {Path(item["source"]).name: item for item in resumed["tasks"]}
        failed = by_name[source.name]
        message = " ".join(
            [failed["pdf_errors"][0]["message"] if failed["pdf_errors"] else ""]
            + [failed["error"]["message"] if failed["error"] else ""]
        )
        self.assertEqual(failed["status"], "succeeded", resumed)
        self.assertEqual(failed["pdf_status"], "failed")
        self.assertNotIn("restored", message.lower())
        self.assertTrue("not written back" in message or "com broke snapshot" in message)
        self.assertEqual(opju.read_bytes(), original)
        self.assertEqual(kept.read_bytes(), kept_bytes)
        self.assertFalse(missing.is_file())
        self.assertEqual(load_batch_record(record)["tasks"][failed["task_id"]]["output"]["sha256"], hashlib.sha256(original).hexdigest())
        self.assertTrue(com.opened)
        self.assertTrue(all(not _same_path(path, opju) for path in com.opened))
        self.assertEqual(by_name["next.csv"]["status"], "succeeded", resumed)
        self.assertTrue(Path(by_name["next.csv"]["output"]["path"]).is_file())

    def _partial_pdf(self, folder: Path):
        folder.mkdir(parents=True, exist_ok=True)
        source = self.workbook(folder / "book.xlsx")
        output = folder / "out"
        plot = {"kind": "line", "x": "x", "y": ["y"]}
        first = self.export(
            [source], output, format="opju", pdf=True, plot=plot,
            session_factory=lambda: FakeSession(pdf_fail_indexes={2}),
        )
        task = first["tasks"][0]
        self.assertEqual(task["status"], "succeeded", first)
        self.assertEqual(task["pdf_status"], "failed", first)
        self.assertEqual(len(task["pdfs"]), 1)
        opju = Path(task["output"]["path"])
        kept = Path(task["pdfs"][0]["path"])
        record = Path(first["record_path"])
        stored = load_batch_record(record)
        missing_identity = next(
            item for item in stored["tasks"][task["task_id"]]["graph_identities"]
            if item["index"] != task["pdfs"][0]["index"]
        )
        from origin_bridge.batch import _pdf_target

        missing = _pdf_target(opju, int(missing_identity["index"]), str(missing_identity["table"]))
        self.assertFalse(missing.is_file())
        return source, output, plot, opju, opju.read_bytes(), kept, kept.read_bytes(), record, record.read_bytes(), missing


def _same_path(left, right) -> bool:
    return os.path.normcase(str(Path(left).resolve())) == os.path.normcase(str(Path(right).resolve()))


def _armed_session(op) -> OriginSession:
    session = OriginSession(timeout_s=0)
    session.op = op
    session.started = True
    session._alive = True
    session._live = False
    session._closed = False
    return session


class _Sheet:
    def __init__(self, lname: str, columns: list[list]):
        self.lname = lname
        self.cols = len(columns)
        self._columns = [list(column) for column in columns]

    def to_list(self, index: int):
        return list(self._columns[index])


class _Book:
    def __init__(self, sheets: list[_Sheet]):
        self._sheets = list(sheets)

    def __iter__(self):
        return iter(self._sheets)


class _Graph:
    def __init__(self, short: str, long_name: str, on_save=None):
        self.name = short
        self.lname = long_name
        self.obj = None
        self.saves: list[str] = []
        self.on_save = on_save

    def save_fig(self, dest, replace=False):
        del replace
        self.saves.append(str(dest))
        if self.on_save:
            self.on_save(dest)
        path = Path(dest)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(PDF)
        return str(path)


class _Com:
    def __init__(self, *, books, graphs, on_open=None, on_save=None):
        self.books = books
        self.graphs = graphs
        self.on_open = on_open
        self.on_save = on_save
        self.opened: list[Path] = []

    def open(self, path, *_args):
        opened = Path(path)
        self.opened.append(opened)
        if self.on_open:
            self.on_open(opened)
        return True

    def pages(self, kind):
        if kind == "w":
            return self.books
        if kind == "g":
            return self.graphs
        return []


def _com_for_record(record: dict, source: Path, *, on_open=None, on_save=None) -> _Com:
    entry = next(iter(record["tasks"].values()))
    sheets = [_Sheet(str(item["table"]), [[]]) for item in entry["graph_identities"]]
    sheets.append(_Sheet("Import Provenance", [[str(source.resolve())]]))
    graphs = [
        _Graph(str(item.get("short_name") or ""), str(item["graph"]), on_save=on_save)
        for item in entry["graph_identities"]
    ]
    return _Com(books=[_Book(sheets)], graphs=graphs, on_open=on_open, on_save=on_save)


class ProbeSession(OriginSession):
    """Run the real PDF recovery method against a COM stand-in."""

    def __init__(self, com: _Com, *, on_start=None):
        super().__init__(timeout_s=0)
        self._com = com
        self._on_start = on_start
        self.starts = 0
        self.stops = 0
        self._on = False

    def start(self) -> None:
        if self._on_start:
            self._on_start()
        self.op = self._com
        self.started = True
        self._alive = True
        self._live = False
        self._closed = False
        self._on = True
        self.starts += 1

    def close(self) -> None:
        if self._on:
            self.stops += 1
        self._on = False
        self.started = False
        self._alive = False
        self._closed = True

    def healthy(self) -> bool:
        return self._on and self.op is not None

    def write_project(self, stage_dir, tables, prepared, converted, *, pdf=False, conversion_notes=None):
        fake = FakeSession()
        fake._on = True
        return fake.write_project(
            stage_dir, tables, prepared, converted, pdf=pdf, conversion_notes=conversion_notes,
        )


if __name__ == "__main__":
    unittest.main()
