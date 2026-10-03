"""Write validated, typed tabular imports to XLSX or native Origin projects."""

from __future__ import annotations

import getpass
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import DataColumn, DataImportError, PlannedTable, PreparedImport
from .dates import parse_iso_datetime

_XLSX_MAX_ROWS = 1_048_576
_XLSX_MAX_COLUMNS = 16_384
_XLSX_MAX_CELL_TEXT = 32_767
_PROVENANCE_SHEET = "Provenance"
_PROVENANCE_HEADERS = (
    "record", "table_name", "source_path", "source_sheet", "source_sha256",
    "row_count", "column_index", "column_name", "original_name", "kind",
    "unit", "missing_count", "read_options_json", "plot_json_or_warnings_json",
)
_NUMBER_RE = re.compile(r"^[+-]?(?:(?:\d+(?:\.\d*)?)|(?:\.\d+))(?:[eEdD][+-]?\d+)?$")
_PLOT_TYPES = {"line": "l", "scatter": "s", "line_symbol": "y", "column": "c"}
_DATE_FMT = "yyyy-mm-dd hh:mm:ss.000"


def execute_import(prepared: PreparedImport) -> dict[str, Any]:
    """Export a prepared import and return a JSON-serializable receipt."""
    tables, output, converted, conversion_warnings, conversion_notes = _validate_prepared(prepared)
    _verify_sources(tables)
    output.parent.mkdir(parents=True, exist_ok=True)

    if output.exists() and not prepared.overwrite:
        raise FileExistsError(f"Output already exists and overwrite is disabled: {output}")

    stage_dir = Path(tempfile.mkdtemp(prefix=f".{output.stem}_origin_stage_", dir=output.parent))
    runtime_warnings: list[str] = []
    try:
        with origin_export_lock():
            if prepared.format == "xlsx":
                staged = stage_dir / output.name
                _write_xlsx(staged, tables, converted, conversion_notes)
                if not staged.is_file() or staged.stat().st_size == 0:
                    raise DataImportError(f"Exporter did not create a non-empty output: {staged}")
                _install_staged_file(staged, output, overwrite=prepared.overwrite)
            else:
                staged = _write_opju(stage_dir, tables, prepared, converted, conversion_notes)
                if not staged.is_file() or staged.stat().st_size == 0:
                    raise DataImportError(f"Exporter did not create a non-empty output: {staged}")
                _install_staged_file(staged, output, overwrite=prepared.overwrite)
                if prepared.keep_open:
                    warning = _show_installed_origin(output)
                    if warning:
                        runtime_warnings.append(warning)
        size = output.stat().st_size
        if size <= 0:
            raise DataImportError(f"Installed output is empty: {output}")
        receipt = {
            "ok": True,
            "output": {
                "path": str(output),
                "format": prepared.format,
                "size_bytes": size,
                "sha256": _sha256_file(output),
            },
            "tables": [
                _table_receipt(item, sheet_name, prepared.format)
                for item, sheet_name in zip(tables, _output_sheet_names(tables, prepared.format))
            ],
            "warnings": _all_warnings(prepared, tables, conversion_warnings) + runtime_warnings,
        }
        json.dumps(receipt, ensure_ascii=False)
        return receipt
    finally:
        shutil.rmtree(stage_dir, ignore_errors=True)


def validate_prepared_import(prepared: PreparedImport) -> list[str]:
    """Validate the prepared export without writing files or launching Origin."""
    tables, _output, _converted, conversion_warnings, _notes = _validate_prepared(prepared)
    return _all_warnings(prepared, tables, conversion_warnings)


def _validate_prepared(
    prepared: PreparedImport,
) -> tuple[tuple[PlannedTable, ...], Path, list[list[list[Any]]], list[str], list[list[str]]]:
    if prepared.format not in {"xlsx", "opju"}:
        raise DataImportError("format must be 'xlsx' or 'opju'")
    if not prepared.tables:
        raise DataImportError("Import contains no tables")
    output = Path(prepared.output).expanduser().resolve(strict=False)
    if output.suffix.casefold() != f".{prepared.format}":
        raise DataImportError(f"Output path must end in .{prepared.format}: {output}")

    names: set[str] = set()
    source_hashes: dict[str, str] = {}
    converted: list[list[list[Any]]] = []
    conversion_warnings: list[str] = []
    notes_by_table: list[list[str]] = []
    for planned in prepared.tables:
        table = planned.table
        if not planned.name.strip():
            raise DataImportError("Every table needs a non-empty output name")
        if planned.name.casefold() in names:
            raise DataImportError(f"Duplicate table name: {planned.name}")
        names.add(planned.name.casefold())
        if not table.columns:
            raise DataImportError(f"Table has no columns: {planned.name}")
        if prepared.format == "xlsx" and len(table.columns) > _XLSX_MAX_COLUMNS:
            raise DataImportError(f"Table exceeds the XLSX column limit: {planned.name}")
        if prepared.format == "xlsx" and table.n_rows + 1 > _XLSX_MAX_ROWS:
            raise DataImportError(f"Table exceeds the XLSX row limit: {planned.name}")
        n_rows = len(table.columns[0].values)
        column_values: list[list[Any]] = []
        column_notes: list[str] = []
        for column in table.columns:
            if column.kind not in {"number", "text", "datetime"}:
                raise DataImportError(f"Unsupported column kind {column.kind!r}: {planned.name}/{column.name}")
            if len(column.values) != n_rows:
                raise DataImportError(f"Columns have different row counts: {planned.name}")
            # One pass keeps the integrity checks and the warning facts together.
            values, notes = _convert_column(column, planned.name, prepared.format)
            column_values.append(values)
            column_notes.extend(notes)
        converted.append(column_values)
        notes_by_table.append(column_notes)
        conversion_warnings.extend(column_notes)

        source_path = Path(table.source).expanduser().resolve(strict=False)
        if _path_key(source_path) == _path_key(output):
            raise DataImportError(f"Output path collides with an input file: {source_path}")
        if not re.fullmatch(r"[0-9a-fA-F]{64}", table.source_hash):
            raise DataImportError(f"Invalid SHA-256 for source file: {source_path}")
        key = _path_key(source_path)
        old_hash = source_hashes.setdefault(key, table.source_hash.casefold())
        if old_hash != table.source_hash.casefold():
            raise DataImportError(f"Tables disagree about the source hash: {source_path}")
        _validate_plot(planned)

    if prepared.format == "xlsx":
        for row_index, values in enumerate(_provenance_rows(tuple(prepared.tables), "xlsx", notes_by_table), start=2):
            for column_index, value in enumerate(values, start=1):
                if len(value) > _XLSX_MAX_CELL_TEXT:
                    raise DataImportError(
                        f"XLSX cell text exceeds {_XLSX_MAX_CELL_TEXT} characters in provenance "
                        f"row {row_index}, column {column_index}"
                    )
        for planned in prepared.tables:
            for column in planned.table.columns:
                for row_index, value in enumerate(column.values, start=1):
                    if value is not None and column.kind == "text" and len(value) > _XLSX_MAX_CELL_TEXT:
                        raise DataImportError(
                            f"Text exceeds the XLSX cell limit in {planned.name}/{column.name}, row {row_index}"
                        )

    if output.exists() and not prepared.overwrite:
        raise FileExistsError(f"Output already exists and overwrite is disabled: {output}")
    return tuple(prepared.tables), output, converted, conversion_warnings, notes_by_table


def _validate_plot(planned: PlannedTable) -> None:
    plot = planned.plot
    if plot.kind not in {"none", *_PLOT_TYPES}:
        raise DataImportError(f"Unsupported plot kind {plot.kind!r}: {planned.name}")
    if plot.kind == "none":
        return
    table = planned.table
    if not plot.y:
        raise DataImportError(f"A {plot.kind} plot needs at least one Y column: {planned.name}")
    if table.n_rows == 0:
        raise DataImportError(f"Cannot plot an empty table: {planned.name}")
    if len(set(plot.y)) != len(plot.y):
        raise DataImportError(f"Plot contains a repeated Y column: {planned.name}")
    if plot.x is not None and not 0 <= plot.x < len(table.columns):
        raise DataImportError(f"X column index is outside the table: {planned.name}")
    for y_index in plot.y:
        if not 0 <= y_index < len(table.columns):
            raise DataImportError(f"Y column index is outside the table: {planned.name}")
        if table.columns[y_index].kind != "number":
            raise DataImportError(f"Y columns must be numeric: {planned.name}/{table.columns[y_index].name}")
        if y_index == plot.x:
            raise DataImportError(f"A column cannot be both X and Y: {planned.name}")
    error_pairs: set[int] = set()
    for y_index, error_index in plot.y_error:
        if y_index not in plot.y:
            raise DataImportError(f"Y error pair refers to an unplotted Y column: {planned.name}")
        if y_index in error_pairs:
            raise DataImportError(f"A Y column has more than one error column: {planned.name}")
        error_pairs.add(y_index)
        if not 0 <= error_index < len(table.columns):
            raise DataImportError(f"Y error column index is outside the table: {planned.name}")
        if table.columns[error_index].kind != "number":
            raise DataImportError(f"Y error columns must be numeric: {planned.name}")


def _coerce_column(column: DataColumn, table_name: str) -> list[Any]:
    """Convert one column. Prefer the list returned with warnings by ``_convert_column``."""
    values, _notes = _convert_column(column, table_name, "opju")
    return values


def _convert_column(
    column: DataColumn,
    table_name: str,
    output_format: str,
) -> tuple[list[Any], list[str]]:
    """Validate every cell once and collect the native-format warnings for that pass.

    The returned lists are the values writers and the Origin read-back use. Callers
    discard them with the export; this is not a cross-task cache.
    """
    converted: list[Any] = []
    empty_strings = 0
    rounded: list[tuple[str, float]] = []
    aware = 0
    naive = 0
    xlsx_submillisecond = 0
    origin_submicrosecond = 0
    for row_index, value in enumerate(column.values, start=1):
        if value is None:
            converted.append("" if column.kind == "text" else None)
            continue
        if not isinstance(value, str):
            raise DataImportError(
                f"Column value must be text or null: {table_name}/{column.name}, row {row_index}"
            )
        if column.kind == "text":
            converted.append(value)
            if value == "":
                empty_strings += 1
            continue
        if column.kind == "number":
            exact, number = _numeric_parts(value, table_name, column.name, row_index)
            converted.append(number)
            if exact.is_finite() and math.isfinite(number) and Decimal(str(number)) != exact:
                rounded.append((value, number))
            continue
        parsed = _parse_datetime(value, table_name, column.name, row_index)
        if parsed.utcoffset() is not None:
            from datetime import timezone

            aware += 1
            parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
        else:
            naive += 1
        converted.append(parsed)
        fraction = re.search(r"[Tt ]\d{2}:\d{2}:\d{2}\.(\d+)", value)
        if fraction:
            digits = fraction.group(1)
            if len(digits) > 3 and any(digit != "0" for digit in digits[3:]):
                xlsx_submillisecond += 1
            if len(digits) > 6 and any(digit != "0" for digit in digits[6:]):
                origin_submicrosecond += 1

    warnings: list[str] = []
    if column.kind == "text" and output_format == "xlsx" and empty_strings:
        warnings.append(
            f"{table_name}/{column.name}: {empty_strings} empty string(s) are written as blank XLSX cells "
            "and may read back as missing values"
        )
    elif column.kind == "number" and rounded:
        raw, number = rounded[0]
        warnings.append(
            f"{table_name}/{column.name}: {len(rounded)} numeric value(s) were rounded "
            f"to binary64; example {raw!r} -> {number!r}"
        )
    elif column.kind == "datetime":
        if aware:
            if naive:
                warnings.append(
                    f"{table_name}/{column.name}: mixed naive and timezone-aware timestamps; "
                    f"{aware} aware value(s) converted to UTC and {naive} naive value(s) retained"
                )
            else:
                warnings.append(
                    f"{table_name}/{column.name}: {aware} timezone-aware timestamp(s) converted to UTC; "
                    "the Origin date column does not store the original timezone"
                )
        if output_format == "xlsx" and xlsx_submillisecond:
            warnings.append(
                f"{table_name}/{column.name}: {xlsx_submillisecond} timestamp(s) have precision finer than "
                "milliseconds; XLSX native date cells store milliseconds"
            )
        if output_format == "opju" and origin_submicrosecond:
            warnings.append(
                f"{table_name}/{column.name}: {origin_submicrosecond} timestamp(s) have precision finer than "
                "microseconds; Origin date values are limited by datetime parsing and serial-number precision"
            )
    return converted, warnings


def _numeric_parts(raw: str, table_name: str, column_name: str, row_index: int) -> tuple[Decimal, float]:
    token = raw.strip().replace("d", "e").replace("D", "E")
    try:
        if not _NUMBER_RE.fullmatch(token):
            raise InvalidOperation
        exact = Decimal(token)
        value = float(exact)
    except (InvalidOperation, OverflowError, ValueError) as exc:
        raise DataImportError(
            f"Invalid numeric value in {table_name}/{column_name}, row {row_index}: {raw!r}"
        ) from exc
    if not exact.is_finite() or not math.isfinite(value):
        raise DataImportError(
            f"Non-finite numeric value in {table_name}/{column_name}, row {row_index}: {raw!r}"
        )
    if exact != 0 and value == 0:
        raise DataImportError(
            f"Numeric value underflows to zero in {table_name}/{column_name}, row {row_index}: {raw!r}"
        )
    return exact, value


def _numeric_value(raw: str, table_name: str, column_name: str, row_index: int) -> float:
    return _numeric_parts(raw, table_name, column_name, row_index)[1]


def _datetime_value(raw: str, table_name: str, column_name: str, row_index: int) -> datetime:
    value = _parse_datetime(raw, table_name, column_name, row_index)
    if value.utcoffset() is not None:
        from datetime import timezone
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _parse_datetime(raw: str, table_name: str, column_name: str, row_index: int) -> datetime:
    try:
        value = parse_iso_datetime(raw)
    except ValueError as exc:
        raise DataImportError(
            f"Invalid ISO date/time in {table_name}/{column_name}, row {row_index}: {raw!r}"
        ) from exc
    return value


def _verify_sources(tables: tuple[PlannedTable, ...]) -> None:
    checked: dict[str, str] = {}
    for planned in tables:
        path = Path(planned.table.source).expanduser().resolve(strict=False)
        key = _path_key(path)
        expected = planned.table.source_hash.casefold()
        actual = checked.get(key)
        if actual is None:
            if not path.is_file():
                raise DataImportError(f"Source file no longer exists: {path}")
            actual = _sha256_file(path)
            checked[key] = actual
        if actual != expected:
            raise DataImportError(f"Source changed after planning; prepare the import again: {path}")


def _write_xlsx(
    path: Path,
    tables: tuple[PlannedTable, ...],
    converted: list[list[list[Any]]],
    conversion_notes: list[list[str]] | None = None,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    names = _output_sheet_names(tables, "xlsx")
    for table_index, (planned, sheet_name) in enumerate(zip(tables, names)):
        sheet = workbook.create_sheet(sheet_name)
        column_values = converted[table_index]
        for column_index, column in enumerate(planned.table.columns, start=1):
            header = sheet.cell(1, column_index, column.name)
            header.data_type = "s"
            header.font = Font(bold=True)
            header.number_format = "@"
            sheet.column_dimensions[get_column_letter(column_index)].width = min(42, max(12, len(column.name) + 2))
            for row_index, value in enumerate(column_values[column_index - 1], start=2):
                cell = sheet.cell(row_index, column_index)
                if value is None:
                    continue
                if column.kind == "text":
                    cell.value = value
                    cell.data_type = "s"
                    cell.number_format = "@"
                elif column.kind == "datetime":
                    cell.value = value
                    cell.number_format = _DATE_FMT
                else:
                    cell.value = value
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 2:
            sheet.auto_filter.ref = sheet.dimensions

    provenance = workbook.create_sheet(_PROVENANCE_SHEET)
    for col_index, name in enumerate(_PROVENANCE_HEADERS, start=1):
        cell = provenance.cell(1, col_index, name)
        cell.data_type = "s"
        cell.font = Font(bold=True)
    for row_index, values in enumerate(_provenance_rows(tables, "xlsx", conversion_notes), start=2):
        for col_index, value in enumerate(values, start=1):
            cell = provenance.cell(row_index, col_index, value)
            cell.data_type = "s"
            cell.number_format = "@"
    provenance.freeze_panes = "A2"
    provenance.auto_filter.ref = provenance.dimensions
    provenance.column_dimensions["C"].width = 72
    provenance.column_dimensions["M"].width = 56
    provenance.column_dimensions["N"].width = 56

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    if not zipfile.is_zipfile(path):
        raise DataImportError(f"XLSX validation failed: {path}")
    checked = load_workbook(path, read_only=True, data_only=False)
    try:
        if set(checked.sheetnames) != set(names) | {_PROVENANCE_SHEET}:
            raise DataImportError(f"XLSX workbook is missing a table or provenance sheet: {path}")
        for planned, name in zip(tables, names):
            sheet = checked[name]
            if sheet.max_column != len(planned.table.columns) or sheet.max_row != planned.table.n_rows + 1:
                raise DataImportError(f"XLSX table dimensions did not round-trip: {planned.name}")
            headers = tuple(sheet.cell(1, index).value for index in range(1, sheet.max_column + 1))
            if headers != tuple(column.name for column in planned.table.columns):
                raise DataImportError(f"XLSX column names did not round-trip: {planned.name}")
    finally:
        checked.close()


def _provenance_rows(
    tables: tuple[PlannedTable, ...],
    output_format: str,
    conversion_notes: list[list[str]] | None = None,
) -> list[tuple[str, ...]]:
    rows: list[tuple[str, ...]] = []
    output_names = _output_sheet_names(tables, output_format)
    for index, (planned, output_name) in enumerate(zip(tables, output_names)):
        table = planned.table
        if conversion_notes is None:
            table_notes = _table_conversion_warnings(planned, output_format)
        else:
            table_notes = list(conversion_notes[index])
        source = str(Path(table.source).expanduser().resolve(strict=False))
        read_options = _json_text(table.read_options)
        notes = _json_text({
            "plot": _plot_dict(planned),
            "warnings": list(table.warnings) + table_notes,
            "output_sheet": output_name,
            "datetime_policy": "timezone-aware values converted to UTC; naive values retained",
            "numeric_policy": "Origin/XLSX number cells use finite binary64 values",
        })
        rows.append((
            "table", planned.name, source, table.source_sheet or "", table.source_hash.casefold(),
            str(table.n_rows), "", "", "", "", "", "", read_options, notes,
        ))
        for index, column in enumerate(table.columns):
            rows.append((
                "column", planned.name, source, table.source_sheet or "", table.source_hash.casefold(),
                str(table.n_rows), str(index), column.name, column.original_name, column.kind,
                column.unit, str(sum(value is None for value in column.values)), read_options, notes,
            ))
    return rows


def _write_opju(
    stage_dir: Path,
    tables: tuple[PlannedTable, ...],
    prepared: PreparedImport,
    converted: list[list[list[Any]]],
    conversion_notes: list[list[str]] | None = None,
) -> Path:
    """Start one Origin session, write one project, and close that session."""
    from .session import OriginSession, OriginSessionLost

    sto = _origin_support()
    session = OriginSession(timeout_s=0)
    try:
        session.start()
        result = session.write_project(
            stage_dir, tables, prepared, converted, pdf=False, conversion_notes=conversion_notes,
        )
        saved = result["opju"]
        if not saved.is_file() or saved.stat().st_size == 0:
            raise sto.OriginExportError("Origin did not create a usable .opju file")
        return saved
    except sto.OriginExportError:
        raise
    except OriginSessionLost as exc:
        raise sto.OriginExportError(str(exc)) from exc
    except Exception as exc:
        raise sto.OriginExportError(f"Origin project export failed: {exc}") from exc
    finally:
        session.close()


def _populate_origin_project(
    op,
    tables: tuple[PlannedTable, ...],
    converted: list[list[list[Any]]],
    conversion_notes: list[list[str]] | None = None,
):
    """Replace the current Origin project with these tables. ``op.new`` returns None."""
    sto = _origin_support()
    op.new(False)
    books = list(op.pages("w"))
    if not books:
        raise sto.OriginExportError("Origin did not create a workbook")
    book = books[0]
    book.lname = "Imported Data"
    for extra in books[1:]:
        extra.destroy()
    for key in ("g", "m"):
        try:
            leftovers = list(op.pages(key))
        except Exception:
            leftovers = []
        for page in leftovers:
            destroy = getattr(page, "destroy", None)
            if callable(destroy):
                destroy()
    if list(op.pages("g")):
        raise sto.OriginExportError("Origin still had graph pages after starting a new project")

    short_names = _output_sheet_names(tables, "opju") + [_origin_provenance_sheet_name(tables)]
    data_sheets = []
    graph_pages = []
    for index, (planned, short_name) in enumerate(zip(tables, short_names[:-1])):
        sheet = book[0] if index == 0 else book.add_sheet(short_name)
        sheet.name = short_name
        sheet.lname = planned.name
        _fill_origin_sheet(op, sheet, planned, converted[index])
        data_sheets.append(sheet)
        if planned.plot.kind != "none":
            graph_pages.append(_plot_origin_table(op, sheet, planned))

    provenance = book.add_sheet(short_names[-1])
    provenance.name = short_names[-1]
    provenance.lname = "Import Provenance"
    _fill_origin_provenance(provenance, tables, conversion_notes)
    _verify_origin_project(op, data_sheets, graph_pages, tables, converted)
    return data_sheets, graph_pages


def _show_installed_origin(path: Path) -> str | None:
    sto = _origin_support()
    if sto.origin_process_running():
        return "The .opju was saved, but Origin was started elsewhere before it could be reopened."
    op = sto._load_originpro()
    if sto.origin_process_running():
        return "The .opju was saved, but Origin became busy before it could be reopened."
    try:
        result = op.open(str(path))
        if result is False:
            raise RuntimeError("Origin returned false while opening the saved project")
        op.set_show(True)
        return None
    except Exception as exc:
        sto.close_origin_app(op, started=True)
        return f"The .opju was saved, but Origin could not remain open: {exc}"


def _fill_origin_sheet(op, sheet, planned: PlannedTable, column_values: list[list[Any]]) -> None:
    table = planned.table
    sheet.cols = len(table.columns)
    formats = _origin_format_codes()
    dso = _origin_dso(op) if any(column.kind == "datetime" for column in table.columns) else None
    for index, column in enumerate(table.columns):
        values = list(column_values[index])
        if column.kind in {"number", "datetime"}:
            if column.kind == "datetime":
                values = [float("nan") if value is None else _origin_date_number(value, dso) for value in values]
            else:
                values = [float("nan") if value is None else value for value in values]
        sheet.obj[index].SetDataFormat(formats[column.kind])
        original_name = column.original_name.strip()
        comments = f"Source column: {original_name}" if original_name and original_name != column.name else ""
        sheet.from_list(index, values, lname=column.name, units=column.unit, comments=comments)
        if column.kind == "datetime":
            sheet.as_date(index, _DATE_FMT)

    sheet.cols_axis()
    if planned.plot.kind != "none":
        plot = planned.plot
        if plot.x is not None:
            sheet.cols_axis("X", plot.x, plot.x)
            x_column = table.columns[plot.x]
            if x_column.kind == "text":
                categories = list(dict.fromkeys(value for value in x_column.values if value is not None))
                if categories:
                    sheet.obj[plot.x].SetCategMapSortCategories(_origin_custom_category_sort(), categories)
        for y_index in plot.y:
            sheet.cols_axis("Y", y_index, y_index)
        for _y_index, error_index in plot.y_error:
            sheet.cols_axis("E", error_index, error_index)


def _plot_origin_table(op, sheet, planned: PlannedTable):
    spec = planned.plot
    title = spec.title.strip() or planned.name
    graph = op.new_graph(lname=title, hidden=False)
    if graph is None:
        raise _origin_error_type()(f"Origin could not create graph: {title}")
    graph.lname = title
    layer = graph[0]
    error_columns = dict(spec.y_error)
    x_column = "#" if spec.x is None else spec.x
    for y_index in spec.y:
        kwargs: dict[str, Any] = {"coly": y_index, "colx": x_column, "type": _PLOT_TYPES[spec.kind]}
        if y_index in error_columns:
            kwargs["colyerr"] = error_columns[y_index]
        if layer.add_plot(sheet, **kwargs) is None:
            raise _origin_error_type()(f"Origin could not plot {planned.name}/{sheet.get_label(y_index, 'L')}")
    plots = list(layer.obj.DataPlots)
    expected_plot_objects = len(spec.y) + len(spec.y_error)
    if len(plots) != expected_plot_objects:
        raise _origin_error_type()(f"Expected {expected_plot_objects} plot objects in {title}; Origin made {len(plots)}")
    x_title = spec.x_label.strip()
    if not x_title:
        x_title = _axis_label(planned.table.columns[spec.x]) if spec.x is not None else "Row"
    y_title = spec.y_label.strip() or _default_y_axis_label(planned)
    if x_title:
        layer.axis("x").title = x_title
    if y_title:
        layer.axis("y").title = y_title
    layer.rescale()
    return graph


def _fill_origin_provenance(
    sheet,
    tables: tuple[PlannedTable, ...],
    conversion_notes: list[list[str]] | None = None,
) -> None:
    rows = [tuple(_PROVENANCE_HEADERS), *_provenance_rows(tables, "opju", conversion_notes)]
    columns = list(zip(*rows))
    formats = _origin_format_codes()
    sheet.cols = len(_PROVENANCE_HEADERS)
    for index, (header, values) in enumerate(zip(_PROVENANCE_HEADERS, columns)):
        sheet.obj[index].SetDataFormat(formats["text"])
        sheet.from_list(index, list(values), lname=header)
    sheet.cols_axis()


def _verify_origin_project(
    op,
    sheets,
    graphs,
    tables: tuple[PlannedTable, ...],
    converted: list[list[list[Any]]],
) -> None:
    if len(sheets) != len(tables):
        raise _origin_error_type()(f"Origin workbook has {len(sheets)} data sheets; expected {len(tables)}")
    expected_graphs = sum(item.plot.kind != "none" for item in tables)
    if len(graphs) != expected_graphs:
        raise _origin_error_type()(f"Origin project has {len(graphs)} graphs; expected {expected_graphs}")
    dso = _origin_dso(op) if any(col.kind == "datetime" for item in tables for col in item.table.columns) else 0
    graph_index = 0
    for table_index, (sheet, planned) in enumerate(zip(sheets, tables)):
        table = planned.table
        if sheet.cols != len(table.columns):
            raise _origin_error_type()(f"Wrong column count after Origin import: {planned.name}")
        if table.n_rows and sheet.rows != table.n_rows:
            raise _origin_error_type()(f"Wrong row count after Origin import: {planned.name}")
        for column_index, column in enumerate(table.columns):
            actual = sheet.to_list(column_index)
            expected = list(converted[table_index][column_index])
            if column.kind == "datetime":
                expected = [float("nan") if value is None else _origin_date_number(value, dso) for value in expected]
            elif column.kind == "number":
                expected = [float("nan") if value is None else value for value in expected]
            if len(actual) != len(expected):
                raise _origin_error_type()(f"Wrong data length after Origin import: {planned.name}/{column.name}")
            for row_index, (found, value) in enumerate(zip(actual, expected), start=1):
                if column.kind == "text":
                    if found not in (value, None if value == "" else value):
                        raise _origin_error_type()(f"Text changed in Origin: {planned.name}/{column.name}, row {row_index}")
                elif value is None or (isinstance(value, float) and math.isnan(value)):
                    if not isinstance(found, (int, float)) or not math.isnan(float(found)):
                        raise _origin_error_type()(f"Missing numeric cell changed in Origin: {planned.name}/{column.name}, row {row_index}")
                elif not isinstance(found, (int, float)) or not math.isclose(float(found), float(value), rel_tol=0, abs_tol=1e-9):
                    raise _origin_error_type()(f"Numeric value changed in Origin: {planned.name}/{column.name}, row {row_index}")
            if sheet.get_label(column_index, "L") != column.name or sheet.get_label(column_index, "U") != column.unit:
                raise _origin_error_type()(f"Column labels changed in Origin: {planned.name}/{column.name}")
        if planned.plot.kind != "none":
            expected_plots = len(planned.plot.y) + len(planned.plot.y_error)
            if len(list(graphs[graph_index][0].obj.DataPlots)) != expected_plots:
                raise _origin_error_type()(f"Wrong graph plot count after Origin import: {planned.name}")
            graph_index += 1


def _origin_date_number(value: datetime, dso: float) -> float:
    """Match originpro's pandas-to-Origin date conversion without pandas."""
    seconds = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    julian_date = value.toordinal() + 1_721_424.5 + seconds / 86_400
    return julian_date + dso - 2_415_018.5


def _origin_dso(op) -> float:
    try:
        return float(op.sysvar["DSO"])
    except Exception as exc:
        raise _origin_error_type()(f"Origin did not expose the @DSO date offset: {exc}") from exc


def _origin_format_codes() -> dict[str, Any]:
    try:
        from originpro.config import po
        return {"number": po.DF_DOUBLE, "text": po.DF_TEXT, "datetime": po.DF_DATE}
    except Exception as exc:
        raise _origin_error_type()(f"OriginPro column-format constants are unavailable: {exc}") from exc


def _origin_custom_category_sort():
    try:
        from originpro.config import po
        return po.CM_SORTING_CUSTOM
    except Exception as exc:
        raise _origin_error_type()(f"OriginPro categorical sort constant is unavailable: {exc}") from exc


def _origin_support():
    # Reuse the existing process guard, launcher, close policy, and save helper.
    import spectra_to_origin
    return spectra_to_origin


def _origin_error_type():
    return _origin_support().OriginExportError


def _axis_label(column: DataColumn) -> str:
    return f"{column.name} ({column.unit})" if column.unit else column.name


def _default_y_axis_label(planned: PlannedTable) -> str:
    if len(planned.plot.y) == 1:
        return _axis_label(planned.table.columns[planned.plot.y[0]])
    units = {planned.table.columns[index].unit for index in planned.plot.y}
    return next(iter(units)) if len(units) == 1 else ""


def _unique_excel_names(raw_names: list[str], reserved: set[str]) -> list[str]:
    used = set(reserved)
    result: list[str] = []
    for raw in raw_names:
        base = re.sub(r"[\[\]:*?/\\]", "_", raw).strip().strip("'") or "Table"
        base = base[:31]
        candidate = base
        serial = 2
        while candidate.casefold() in used:
            suffix = f"_{serial}"
            candidate = f"{base[:31 - len(suffix)]}{suffix}"
            serial += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def _output_sheet_names(tables: tuple[PlannedTable, ...], output_format: str) -> list[str]:
    """Return physical worksheet names in planned-table order."""
    raw_names = [item.name for item in tables]
    if output_format == "xlsx":
        return _unique_excel_names(raw_names, {_PROVENANCE_SHEET.casefold()})
    if output_format == "opju":
        return _unique_origin_names(raw_names + [_PROVENANCE_SHEET])[:-1]
    raise DataImportError(f"Unsupported output format: {output_format}")


def _origin_provenance_sheet_name(tables: tuple[PlannedTable, ...]) -> str:
    return _unique_origin_names([item.name for item in tables] + [_PROVENANCE_SHEET])[-1]


def _unique_origin_names(raw_names: list[str], limit: int = 13) -> list[str]:
    used: set[str] = set()
    result: list[str] = []
    for raw in raw_names:
        base = re.sub(r"[^0-9A-Za-z]+", "", raw.strip()) or "Table"
        if base[0].isdigit():
            base = "T" + base
        base = base[:limit]
        candidate = base
        serial = 2
        while candidate.casefold() in used:
            suffix = str(serial)
            candidate = f"{base[:limit - len(suffix)]}{suffix}"
            serial += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def _table_receipt(planned: PlannedTable, output_sheet: str, output_format: str) -> dict[str, Any]:
    table = planned.table
    graph = None
    if output_format == "opju" and planned.plot.kind != "none":
        graph = {
            "type": planned.plot.kind,
            "title": planned.plot.title or planned.name,
            "plot_count": len(planned.plot.y),
            "error_bar_count": len(planned.plot.y_error),
        }
    return {
        "name": planned.name,
        "output_sheet": output_sheet,
        "source": str(Path(table.source).expanduser().resolve(strict=False)),
        "source_sheet": table.source_sheet,
        "rows": table.n_rows,
        "columns": [
            {"index": i, "name": col.name, "kind": col.kind, "unit": col.unit,
             "missing": sum(value is None for value in col.values)}
            for i, col in enumerate(table.columns)
        ],
        "graph": graph,
    }


def _all_warnings(
    prepared: PreparedImport,
    tables: tuple[PlannedTable, ...],
    conversion_warnings: list[str] | None = None,
) -> list[str]:
    warnings = list(prepared.warnings)
    for planned in tables:
        warnings.extend(planned.table.warnings)
    if conversion_warnings is None:
        for planned in tables:
            warnings.extend(_table_conversion_warnings(planned, prepared.format))
    else:
        warnings.extend(conversion_warnings)
    return list(dict.fromkeys(str(item) for item in warnings if str(item).strip()))


def _table_conversion_warnings(planned: PlannedTable, output_format: str) -> list[str]:
    warnings: list[str] = []
    for column in planned.table.columns:
        _values, notes = _convert_column(column, planned.name, output_format)
        warnings.extend(notes)
    return warnings


def _plot_dict(planned: PlannedTable) -> dict[str, Any]:
    plot = planned.plot
    return {
        "kind": plot.kind, "x": plot.x, "y": list(plot.y),
        "y_error": [list(pair) for pair in plot.y_error],
        "title": plot.title, "x_label": plot.x_label, "y_label": plot.y_label,
    }


def _json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _path_key(path: Path) -> str:
    return os.path.normcase(str(Path(path).expanduser().resolve(strict=False)))


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class _OriginExportLock:
    """OS-backed per-user lock; process death releases the lock automatically."""

    def __init__(self) -> None:
        user_key = hashlib.sha256(getpass.getuser().encode("utf-8")).hexdigest()[:16]
        self.path = Path(tempfile.gettempdir()) / f"origin_bridge_{user_key}.lock"
        self.stream = None

    def __enter__(self):
        self.stream = self.path.open("a+b")
        self.stream.seek(0, os.SEEK_END)
        if self.stream.tell() == 0:
            self.stream.write(b"\0")
            self.stream.flush()
        self.stream.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(self.stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.stream.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except (OSError, BlockingIOError) as exc:
            self.stream.close()
            self.stream = None
            raise _origin_error_type()("Another generic Origin export is already active") from exc
        return self

    def __exit__(self, exc_type, exc, traceback) -> None:
        if self.stream is None:
            return
        try:
            self.stream.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(self.stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.stream.fileno(), fcntl.LOCK_UN)
        finally:
            self.stream.close()
            self.stream = None


def _origin_export_lock() -> _OriginExportLock:
    return _OriginExportLock()


_export_gate = threading.Lock()
_export_local = threading.local()


class origin_export_lock:
    """Reentrant per-thread, exclusive across threads and processes.

    The file lock is the one old entry points already take. A second acquire on
    the same thread nests. Another thread or process gets the existing
    "already active" error instead of attaching to the same Origin session.
    """

    def __enter__(self):
        depth = getattr(_export_local, "depth", 0)
        if depth == 0:
            if not _export_gate.acquire(blocking=False):
                raise _origin_error_type()("Another generic Origin export is already active")
            lock = _origin_export_lock()
            try:
                lock.__enter__()
            except BaseException:
                _export_gate.release()
                raise
            _export_local.lock = lock
        _export_local.depth = depth + 1
        return self

    def __exit__(self, exc_type, exc, traceback) -> None:
        _export_local.depth -= 1
        if _export_local.depth == 0:
            try:
                _export_local.lock.__exit__(exc_type, exc, traceback)
            finally:
                _export_local.lock = None
                _export_gate.release()


def _install_staged_file(staged: Path, destination: Path, *, overwrite: bool) -> None:
    """Atomically install a same-volume staged file without a no-clobber race."""
    try:
        if overwrite:
            os.replace(staged, destination)
        elif os.name == "nt":
            # On Windows, os.rename fails atomically if destination already exists.
            os.rename(staged, destination)
        else:
            # A same-filesystem hard link atomically fails if destination exists.
            os.link(staged, destination)
            staged.unlink()
    except FileExistsError:
        raise FileExistsError(f"Output appeared during export and was preserved: {destination}") from None
    except OSError as exc:
        raise DataImportError(f"Could not install output at {destination}: {exc}") from exc
