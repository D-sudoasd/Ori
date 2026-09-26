"""CLI contracts tested through actual files and the public entry point."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import spectra_to_origin as sto


class CliTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.source = self.root / "input"
        self.source.mkdir()
        (self.source / "sample1.txt").write_text("0 10\n1 11\n", encoding="utf-8")
        (self.source / "sample2.txt").write_text("0.0 12\n1.0 13\n", encoding="utf-8")

    def run_command(self, *flags):
        stdout, stderr = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            result = sto.main(list(flags))
        return result, stdout.getvalue(), stderr.getvalue()

    def test_check_does_not_require_output_or_start_origin(self):
        with mock.patch.object(sto, "_load_originpro", side_effect=AssertionError("Origin must not start")):
            code, output, error = self.run_command("--check", "-i", str(self.source))
        self.assertEqual(code, 0, error)
        self.assertIn("2 条谱线", output)
        self.assertIn("XYYY", output)
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), ["input"])

    def test_mixed_valid_and_missing_input_fails_before_writing(self):
        target = self.root / "out.xlsx"
        code, output, error = self.run_command("--cli", "--xlsx-only", "-i", str(self.source), str(self.root / "missing"), "-o", str(target))
        self.assertEqual(code, 2)
        self.assertIn("missing", error)
        self.assertEqual(output, "")
        self.assertFalse(target.exists())

    def test_custom_axes_reach_saved_workbook(self):
        target = self.root / "out.xlsx"
        code, output, error = self.run_command("--cli", "--xlsx-only", "-i", str(self.source), "-o", str(target), "--x-name", "Raman shift", "--x-unit", "cm^-1", "--y-name", "Counts", "--y-unit", "counts")
        self.assertEqual(code, 0, error)
        workbook = load_workbook(target)
        self.addCleanup(workbook.close)
        sheet = workbook.worksheets[0]
        self.assertEqual(sheet.cell(1, 1).value, "Raman shift")
        self.assertEqual(sheet.cell(2, 1).value, "cm^-1")
        self.assertEqual(sheet.cell(2, 2).value, "counts")
        self.assertEqual(sheet.cell(4, 2).value, 10)
        self.assertTrue(target.with_suffix(".csv").is_file())

    def test_invalid_groups_and_conflicting_group_options_are_usage_errors(self):
        for extra in [("--n-groups", "0"), ("--n-groups", "-1"), ("--n-groups", "oops"), ("--n-groups", "2", "--group-by-name")]:
            with self.subTest(extra=extra), contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit) as failure:
                    sto.build_parser().parse_args(list(extra))
                self.assertEqual(failure.exception.code, 2)

    def test_bad_data_reports_source_line_and_fails(self):
        bad = self.source / "broken.csv"
        bad.write_text("x,y\n0,10\n1,NaN\n", encoding="utf-8")
        code, output, error = self.run_command("--check", "-i", str(bad))
        self.assertEqual(code, 4)
        self.assertIn("broken.csv", error)
        self.assertIn("3", error)
        self.assertEqual(output, "")

    def test_xlsx_io_failure_has_no_traceback(self):
        with mock.patch.object(sto, "export_spectra", side_effect=PermissionError("workbook is open")):
            code, output, error = self.run_command("--cli", "--xlsx-only", "-i", str(self.source), "-o", str(self.root / "locked.xlsx"))
        self.assertEqual(code, 4)
        self.assertIn("workbook is open", error)
        self.assertNotIn("Traceback", error)


if __name__ == "__main__":
    unittest.main()
