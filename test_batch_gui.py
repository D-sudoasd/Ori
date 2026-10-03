"""Batch window seams: per-file discovery, spawn, cancel, resume, and column scope."""

from __future__ import annotations

import hashlib
import json
import multiprocessing
import os
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from origin_bridge.batch_config import config_path_for, find_config_for_record, load_batch_config, save_batch_config
from origin_bridge.gui import build_global_plot, widget_in_client


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _pump(app, timeout: float) -> None:
    deadline = time.monotonic() + timeout
    while app._busy and time.monotonic() < deadline:
        app.root.update()
        time.sleep(0.005)
    app.root.update()


def _batch_configs(folder: Path) -> list[Path]:
    return sorted(path for path in folder.glob("batch-config*.json") if path.is_file())


def _one_config(folder: Path) -> dict:
    matches = _batch_configs(folder)
    if len(matches) != 1:
        raise AssertionError(f"expected one batch config in {folder}, found {[path.name for path in matches]}")
    return load_batch_config(matches[0])


def _provenance_text(path: Path) -> str:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        chunks: list[str] = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if isinstance(value, str) and "y_error" in value:
                        chunks.append(value)
        return "\n".join(chunks)
    finally:
        book.close()


def _shutdown(app) -> None:
    process = getattr(app, "_import_process", None)
    if process is not None and getattr(process, "is_alive", lambda: False)():
        process.terminate()
        process.join(timeout=3)
    try:
        if app.root.winfo_exists():
            app.root.destroy()
    except Exception:
        pass


class BatchConfigTests(unittest.TestCase):
    def test_config_round_trip_does_not_touch_the_core_record(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            record = folder / "batch-record.json"
            record.write_text('{"schema_version":1,"kind":"ori-batch-record","tasks":{}}\n', encoding="utf-8")
            before = record.read_bytes()
            request = {
                "schema_version": 1,
                "inputs": [str(folder / "a.csv")],
                "output_dir": str(folder),
                "layout": "per_file",
                "format": "xlsx",
                "read_options": {"header": "auto", "skip_rows": 0},
                "plot": {"kind": "none", "x": "auto", "y": "auto", "y_error": "auto", "title": "auto", "x_label": "auto", "y_label": "auto"},
                "task_overrides": [],
                "overwrite": False,
                "pdf": False,
                "keep_open": False,
                "recycle_every": 0,
                "prefetch": 1,
                "max_resident_tasks": 2,
                "origin_timeout_s": 300,
                "origin_retries": 1,
                "output_name": "batch.xlsx",
                "record_path": str(record),
                "resume": False,
                "retry_task_ids": [],
            }
            path = save_batch_config(
                folder / "batch-config.json",
                name="实验甲",
                record_path=record,
                inputs=request["inputs"],
                request=request,
            )
            loaded = load_batch_config(path)
            self.assertEqual(loaded["name"], "实验甲")
            self.assertEqual(loaded["request"]["plot"]["kind"], "none")
            self.assertEqual(loaded["request"]["inputs"], request["inputs"])
            self.assertEqual(record.read_bytes(), before)
            first_config = path.read_bytes()
            save_batch_config(path, name="实验乙", record_path=record, inputs=request["inputs"], request=request)
            backup = Path(str(path) + ".bak")
            self.assertEqual(backup.read_bytes(), first_config)
            self.assertEqual(json.loads(backup.read_text(encoding="utf-8"))["name"], "实验甲")
            self.assertEqual(record.read_bytes(), before)

    def test_legacy_pair_stays_loadable_and_fresh_names_do_not_replace_it(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            legacy_record = folder / "batch-record.json"
            legacy_record.write_text('{"schema_version":1,"kind":"ori-batch-record","tasks":{"keep":1}}\n', encoding="utf-8")
            request = {
                "inputs": [str(folder / "a.csv")],
                "record_path": str(legacy_record.resolve()),
                "plot": {"kind": "none"},
            }
            legacy_config = save_batch_config(
                config_path_for(legacy_record),
                name="旧批次",
                record_path=legacy_record,
                inputs=request["inputs"],
                request=request,
            )
            self.assertEqual(legacy_config.name, "batch-config.json")
            self.assertEqual(find_config_for_record(legacy_record), legacy_config.resolve())
            before = {
                legacy_record: legacy_record.read_bytes(),
                legacy_config: legacy_config.read_bytes(),
            }
            fresh = folder / "batch-record-abc123.json"
            self.assertEqual(config_path_for(fresh).name, "batch-config-abc123.json")
            fresh_request = dict(request)
            fresh_request["record_path"] = str(fresh.resolve())
            save_batch_config(
                config_path_for(fresh),
                name="新批次",
                record_path=fresh,
                inputs=fresh_request["inputs"],
                request=fresh_request,
            )
            self.assertEqual(legacy_record.read_bytes(), before[legacy_record])
            self.assertEqual(legacy_config.read_bytes(), before[legacy_config])
            self.assertEqual(find_config_for_record(legacy_record).name, "batch-config.json")
            self.assertEqual(find_config_for_record(fresh).name, "batch-config-abc123.json")

    def test_inputs_must_match_the_request(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            record = str(folder / "batch-record.json")
            request = {"inputs": [str(folder / "a.csv")], "record_path": record, "plot": {"y": "auto"}}
            with self.assertRaises(ValueError):
                save_batch_config(folder / "batch-config.json", name="", record_path=record, inputs=[str(folder / "b.csv")], request=request)


class GlobalPlotTests(unittest.TestCase):
    def test_auto_scope_ignores_typed_columns(self):
        plot = build_global_plot(scope="自动（按每张表建议）", kind="line", explicit_x="9", explicit_y="other", explicit_error="err")
        self.assertEqual(plot["x"], "auto")
        self.assertEqual(plot["y"], "auto")
        self.assertEqual(plot["kind"], "line")

    def test_explicit_scope_keeps_names_and_indexes(self):
        plot = build_global_plot(scope="明确列", kind="auto", explicit_x="0", explicit_y="signal, other", explicit_error="")
        self.assertEqual(plot["kind"], "auto")
        self.assertEqual(plot["x"], 0)
        self.assertEqual(plot["y"], ["signal", "other"])
        self.assertEqual(plot["y_error"], {})

    def test_explicit_scope_requires_a_y_column(self):
        with self.assertRaises(ValueError):
            build_global_plot(scope="明确列", kind="line", explicit_x="x", explicit_y="", explicit_error="")

    def test_error_column_rejects_numeric_y_and_keeps_named_y(self):
        with self.assertRaisesRegex(ValueError, "误差列需要Y使用原列名"):
            build_global_plot(scope="明确列", kind="line", explicit_x="0", explicit_y="0", explicit_error="1")
        named_index = build_global_plot(scope="明确列", kind="line", explicit_x="time", explicit_y="signal", explicit_error="2")
        self.assertEqual(named_index["y"], ["signal"])
        self.assertEqual(named_index["y_error"], {"signal": 2})
        named_name = build_global_plot(scope="明确列", kind="line", explicit_x="time", explicit_y="signal", explicit_error="err")
        self.assertEqual(named_name["y_error"], {"signal": "err"})
        numeric = build_global_plot(scope="明确列", kind="line", explicit_x="0", explicit_y="1", explicit_error="")
        self.assertEqual(numeric["y"], [1])
        self.assertEqual(numeric["y_error"], {})


class _Pipe:
    def __init__(self, messages, raise_after=None):
        self.messages = list(messages)
        self.raise_after = raise_after
        self.recv_count = 0
        self.closed = False

    def poll(self, _timeout=0):
        if self.raise_after is not None and self.recv_count >= self.raise_after:
            raise OSError("broken pipe")
        return bool(self.messages)

    def recv(self):
        if self.raise_after is not None and self.recv_count >= self.raise_after:
            raise OSError("broken pipe")
        if not self.messages:
            raise EOFError
        self.recv_count += 1
        return self.messages.pop(0)

    def close(self):
        self.closed = True


class _Proc:
    def __init__(self):
        self.alive = True
        self.exitcode = None
        self.closed = False
        self.terminated = False

    def is_alive(self):
        if self.closed:
            raise ValueError("closed")
        return self.alive

    def join(self, timeout=None):
        if self.alive and timeout is None:
            raise AssertionError("joined a live process without a bound")
        return None

    def close(self):
        self.closed = True

    def terminate(self):
        self.terminated = True


class _Flag:
    def __init__(self):
        self.flag = False
        self.order = []

    def set(self):
        self.flag = True
        self.order.append("set")

    def is_set(self):
        return self.flag


@unittest.skipUnless(os.environ.get("SPECTRA_TEST_GUI") == "1", "set SPECTRA_TEST_GUI=1 for the Tk batch checks")
class RealTkBatchTests(unittest.TestCase):
    def _app(self):
        from origin_bridge.gui import GeneralDataApp

        try:
            app = GeneralDataApp()
        except Exception as exc:
            self.skipTest(f"Tk display is unavailable: {exc}")
        app.root.withdraw()
        return app

    def test_per_file_lists_1000_without_inspecting_and_keeps_the_heartbeat(self):
        from origin_bridge import gui
        from origin_bridge.readers import discover_files

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw) / "sources"
            folder.mkdir()
            for index in range(1000):
                (folder / f"sample{index}.csv").write_text("x,y\n1,2\n", encoding="utf-8")
            (folder / "notes.md").write_text("ignore", encoding="utf-8")
            app = self._app()
            samples = []
            after_ids = []

            def heartbeat():
                samples.append((time.monotonic(), len(app.table_tree.get_children())))
                if app._busy:
                    after_ids.append(app.root.after(10, heartbeat))

            try:
                after_ids.append(app.root.after(10, heartbeat))
                with mock.patch.object(gui.messagebox, "showwarning"), mock.patch.object(gui.messagebox, "showerror") as showerror:
                    app._add_paths([folder, folder / "notes.md"])
                    self.assertTrue(app._busy)
                    self.assertLess(len(app.table_tree.get_children()), 1000)
                    deadline = time.monotonic() + 60
                    while app._busy and time.monotonic() < deadline:
                        app.root.update()
                        time.sleep(0.005)
                    app.root.update()
                    showerror.assert_not_called()
                self.assertFalse(app._busy)
                self.assertEqual(app._summaries.inspect_calls, 0)
                names = [Path(app.table_tree.item(iid, "values")[0]).name for iid in app.table_tree.get_children()]
                self.assertEqual(len(names), 1000)
                self.assertEqual(set(names), {path.name for path in discover_files([folder])})
                self.assertTrue(all(row[4] == "未预览" for row in (app.table_tree.item(iid, "values") for iid in app.table_tree.get_children())))
                gaps = [samples[index][0] - samples[index - 1][0] for index in range(1, len(samples))]
                jumps = [samples[index][1] - samples[index - 1][1] for index in range(1, len(samples))]
                max_gap = max(gaps) if gaps else 999
                max_jump = max(jumps) if jumps else 1000
                partial = any(0 < count < 1000 for _stamp, count in samples)
                print(f"MEASURE per_file_list heartbeat_max_gap_s={max_gap:.4f} max_row_jump={max_jump} heartbeats={len(samples)}")
                self.assertGreaterEqual(len(samples), 5)
                self.assertTrue(partial)
                self.assertLess(max_jump, 80)
                self.assertLess(max_gap, 0.25)
                app.preview_selected()
                _pump(app, 20)
                self.assertEqual(app._summaries.inspect_calls, 1)
                self.assertEqual(len(app._paths), 1000)
                captured = {}
                out = Path(raw) / "out"
                with (
                    mock.patch.object(app, "_start_batch_process", lambda request: captured.setdefault("request", request)),
                    mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                ):
                    app.batch_kind_var.set("none")
                    app.export_button.invoke()
                    showerror.assert_not_called()
                self.assertEqual(len(captured["request"]["inputs"]), 1000)
                self.assertEqual(captured["request"]["plot"]["y"], "auto")
                self.assertEqual(app._summaries.inspect_calls, 1)
                record_path = Path(captured["request"]["record_path"])
                self.assertTrue(record_path.name.startswith("batch-record-"))
                self.assertNotEqual(record_path.name, "batch-record.json")
                self.assertTrue(config_path_for(record_path).is_file())
                self.assertEqual(config_path_for(record_path).name, "batch-config-" + record_path.name[len("batch-record-"):])
            finally:
                for after_id in after_ids:
                    try:
                        app.root.after_cancel(after_id)
                    except Exception:
                        pass
                _shutdown(app)

    def test_unified_switch_inspects_everything_already_listed(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            for name in ("a.csv", "b.csv", "c.csv"):
                (folder / name).write_text("x,y\n1,2\n", encoding="utf-8")
            app = self._app()
            try:
                app._add_paths([folder])
                _pump(app, 20)
                self.assertEqual(app._summaries.inspect_calls, 0)
                self.assertEqual(len(app._paths), 3)
                app.layout_var.set("unified")
                app._layout_changed()
                _pump(app, 20)
                self.assertEqual(app._summaries.inspect_calls, 3)
                self.assertEqual(len(app._tables), 3)
                self.assertFalse(app.pdf_var.get())
            finally:
                _shutdown(app)

    def test_progress_flood_cancel_and_exception_restore_the_window(self):
        from origin_bridge import gui

        app = self._app()
        try:
            messages = []
            for index in range(400):
                pdf_failed = index == 3
                messages.append({
                    "type": "progress",
                    "event": {
                        "type": "task",
                        "task_id": f"{index:04x}",
                        "index": index,
                        "count": 400,
                        "status": "blocked" if index == 4 else "succeeded",
                        "source": str(Path("C:/data") / f"file{index}.csv"),
                        "output": None,
                        "error": None if index != 4 else {"type": "FileExistsError", "message": "already present"},
                        "pdf_status": "failed" if pdf_failed or index == 4 else "not_applicable",
                        "pdf_errors": [
                            {"table": "file3", "graph": "GraphA", "short_name": "Graph1", "path": "C:/out/a.pdf", "message": "pdf export failed"}
                        ] if pdf_failed or index == 4 else [],
                    },
                })
            messages.append({
                "type": "result",
                "result": {
                    "ok": False,
                    "cancelled": False,
                    "counts": {"succeeded": 398, "skipped": 0, "failed": 0, "cancelled": 0, "blocked": 1, "pdf_failed": 2},
                    "tasks": [],
                },
            })
            pipe = _Pipe(messages)
            process = _Proc()
            cancel = _Flag()
            samples = []

            def heartbeat():
                samples.append((time.monotonic(), len(app.task_tree.get_children())))
                if not pipe.messages and process.alive:
                    process.alive = False
                    process.exitcode = 0
                if app._busy:
                    app.root.after(10, heartbeat)

            app._import_recv = pipe
            app._import_process = process
            app._import_cancel = cancel
            app._import_kind = "batch"
            app._import_result = None
            app._import_exit_waits = 0
            app._import_eof = False
            app._set_busy(True)
            self.assertEqual(str(app.stop_button.cget("state")), "normal")
            with mock.patch.object(gui.messagebox, "showinfo"), mock.patch.object(gui.messagebox, "showerror"):
                app.root.after(10, heartbeat)
                app._poll_import()
                app.stop_later_tasks()
                self.assertTrue(cancel.flag)
                self.assertEqual(str(app.stop_button.cget("state")), "normal")
                deadline = time.monotonic() + 30
                while app._busy and time.monotonic() < deadline:
                    app.root.update()
                    time.sleep(0.005)
                app.root.update()
            self.assertFalse(app._busy)
            self.assertEqual(str(app.export_button.cget("state")), "normal")
            self.assertEqual(str(app.stop_button.cget("state")), "disabled")
            self.assertIsNone(app._import_cancel)
            self.assertIsNone(app._import_process)
            self.assertIsNone(app._import_recv)
            self.assertIsNone(app._import_result)
            self.assertFalse(process.terminated)
            self.assertEqual(len(app.task_tree.get_children()), 400)
            values = [app.task_tree.item(iid, "values") for iid in app.task_tree.get_children()]
            pdf_row = next(row for row in values if row[0].endswith("file3.csv"))
            blocked_row = next(row for row in values if row[0].endswith("file4.csv"))
            self.assertIn("成功（PDF 失败）", pdf_row[1])
            self.assertIn("GraphA", pdf_row[2])
            self.assertIn("pdf export failed", pdf_row[2])
            self.assertIn("已阻塞（PDF 失败）", blocked_row[1])
            self.assertIn("already present", blocked_row[2])
            gaps = [samples[index][0] - samples[index - 1][0] for index in range(1, len(samples))]
            jumps = [samples[index][1] - samples[index - 1][1] for index in range(1, len(samples))]
            max_gap = max(gaps) if gaps else 999
            max_jump = max(jumps) if jumps else 400
            partial = any(0 < count < 400 for _stamp, count in samples)
            print(f"MEASURE batch_flood heartbeat_max_gap_s={max_gap:.4f} max_row_jump={max_jump} heartbeats={len(samples)}")
            self.assertGreaterEqual(len(samples), 8)
            self.assertTrue(partial)
            self.assertLess(max_jump, 40)
            self.assertLess(max_gap, 0.25)
            self.assertIn("PDF 失败", app.status_var.get())

            pipe = _Pipe([
                {"type": "progress", "event": {"type": "batch", "status": "started", "count": 1}},
                {"type": "progress", "event": {"type": "task", "task_id": "aa", "index": 0, "count": 1, "status": "failed", "source": "bad.csv", "error": {"type": "ValueError", "message": "boom"}, "pdf_status": "not_applicable", "pdf_errors": []}},
            ], raise_after=1)
            process = _Proc()
            process.alive = False
            process.exitcode = 7
            cancel = _Flag()
            app._import_recv = pipe
            app._import_process = process
            app._import_cancel = cancel
            app._import_kind = "batch"
            app._import_result = None
            app._import_exit_waits = 0
            app._import_eof = False
            app._set_busy(True)
            with mock.patch.object(gui.messagebox, "showerror") as showerror, mock.patch.object(gui.messagebox, "showinfo") as showinfo:
                while app._busy:
                    app._poll_import()
                showinfo.assert_not_called()
                showerror.assert_called()
            self.assertFalse(app._busy)
            self.assertIsNone(app._import_cancel)
            self.assertIsNone(app._import_process)
            self.assertEqual(str(app.export_button.cget("state")), "normal")
            self.assertFalse(process.terminated)
        finally:
            _shutdown(app)

    def test_destroy_requests_cancel_before_the_pipe_closes(self):
        app = self._app()
        order = []

        class _Event:
            def set(self):
                order.append("set")

        class _Recv:
            def close(self):
                order.append("close")

        process = _Proc()
        try:
            app._busy = True
            app._import_kind = "batch"
            app._import_cancel = _Event()
            app._import_recv = _Recv()
            app._import_process = process
            app.root.destroy()
            self.assertTrue(order)
            self.assertLess(order.index("set"), order.index("close"))
            self.assertFalse(process.terminated)
            self.assertTrue(process.is_alive())
        finally:
            process.alive = False
            _shutdown(app)

    def test_real_xlsx_batch_resume_retry_and_column_mismatch(self):
        from origin_bridge import gui
        from origin_bridge.batch import load_batch_record

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            good = folder / "good.csv"
            other = folder / "other.csv"
            bad = folder / "bad.csv"
            wrong = folder / "wrong.csv"
            good.write_text("x,signal\n0,1\n1,2\n", encoding="utf-8")
            other.write_text("x,signal\n3,4\n", encoding="utf-8")
            wrong.write_text("x,other\n5,6\n", encoding="utf-8")
            bad.write_bytes(b"")
            book = folder / "book.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "one"
            first.append(["x", "signal"])
            first.append([1, 2])
            second = workbook.create_sheet("two")
            second.append(["p", "q"])
            second.append([3, 4])
            workbook.save(book)
            workbook.close()
            out = folder / "batch-out"
            mismatch_out = folder / "mismatch-out"
            app = self._app()
            try:
                app.format_var.set("xlsx")
                app._format_changed()
                with mock.patch.object(gui.messagebox, "showwarning"), mock.patch.object(gui.messagebox, "showerror"), mock.patch.object(gui.messagebox, "showinfo"):
                    app._add_paths([good, other, bad])
                    _pump(app, 20)
                    self.assertEqual(app._summaries.inspect_calls, 0)
                    self.assertEqual(len(app._paths), 3)
                    app.batch_name_var.set("实验甲")
                    app.batch_kind_var.set("none")
                    app.batch_overwrite_var.set(False)
                    with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)):
                        app.export_button.invoke()
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                self.assertIsNone(app._import_process)
                self.assertEqual(str(app.export_button.cget("state")), "normal")
                config = _one_config(out)
                record_path = Path(config["record_path"])
                self.assertTrue(record_path.name.startswith("batch-record-"))
                self.assertEqual(config["name"], "实验甲")
                self.assertEqual(len(config["request"]["inputs"]), 3)
                self.assertFalse(config["request"]["overwrite"])
                self.assertEqual(config["request"]["plot"]["kind"], "none")
                record = load_batch_record(record_path)
                statuses = {Path(task["source"]).name: task["status"] for task in record["tasks"].values()}
                self.assertEqual(statuses["good.csv"], "succeeded")
                self.assertEqual(statuses["other.csv"], "succeeded")
                self.assertEqual(statuses["bad.csv"], "failed")
                good_output = next(Path(task["output"]["path"]) for task in record["tasks"].values() if Path(task["source"]).name == "good.csv")
                good_hash = _sha(good_output)
                record_before = record_path.read_bytes()
                saved_plot = json.loads(json.dumps(config["request"]["plot"]))
                _shutdown(app)

                app = self._app()
                app.format_var.set("xlsx")
                app._format_changed()
                app.plot_scope_var.set("明确列")
                app.explicit_y_var.set("NOPE")
                app.batch_kind_var.set("line")
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", return_value=str(record_path)),
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(gui.messagebox, "showwarning"),
                ):
                    app.continue_button.invoke()
                    self.assertFalse(showerror.called)
                    resumed = load_batch_config(config_path_for(record_path))
                    self.assertTrue(resumed["request"]["resume"])
                    self.assertFalse(resumed["request"]["overwrite"])
                    self.assertEqual(resumed["request"]["retry_task_ids"], [])
                    self.assertEqual(len(resumed["request"]["inputs"]), 3)
                    self.assertEqual(resumed["request"]["plot"], saved_plot)
                    self.assertNotIn("NOPE", json.dumps(resumed["request"]["plot"]))
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                self.assertEqual(_sha(good_output), good_hash)
                self.assertEqual(app.batch_name_var.get(), "实验甲")
                self.assertEqual(len(app._paths), 3)
                resumed_rows = [app.task_tree.item(iid, "values") for iid in app.task_tree.get_children()]
                self.assertTrue(any(str(row[0]).endswith("good.csv") and row[1] == "已跳过" for row in resumed_rows))
                self.assertTrue(any(str(row[0]).endswith("bad.csv") and row[1] == "失败" for row in resumed_rows))
                record = load_batch_record(record_path)
                statuses = {Path(task["source"]).name: task["status"] for task in record["tasks"].values()}
                self.assertEqual(statuses["good.csv"], "succeeded")
                self.assertEqual(statuses["bad.csv"], "failed")
                failed_id = next(task_id for task_id, task in record["tasks"].items() if Path(task["source"]).name == "bad.csv")
                succeeded_id = next(task_id for task_id, task in record["tasks"].items() if Path(task["source"]).name == "good.csv")
                bad.write_text("x,signal\n8,9\n", encoding="utf-8")
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", return_value=str(record_path)),
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showerror"),
                    mock.patch.object(gui.messagebox, "showwarning"),
                ):
                    app.retry_button.invoke()
                    retried = load_batch_config(config_path_for(record_path))
                    self.assertTrue(retried["request"]["resume"])
                    self.assertFalse(retried["request"]["overwrite"])
                    self.assertIn(failed_id, retried["request"]["retry_task_ids"])
                    self.assertNotIn(succeeded_id, retried["request"]["retry_task_ids"])
                    self.assertEqual(len(retried["request"]["inputs"]), 3)
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                record = load_batch_record(record_path)
                statuses = {Path(task["source"]).name: task["status"] for task in record["tasks"].values()}
                self.assertEqual(statuses["bad.csv"], "succeeded")
                self.assertEqual(statuses["good.csv"], "succeeded")
                self.assertEqual(_sha(good_output), good_hash)
                retried_rows = [app.task_tree.item(iid, "values") for iid in app.task_tree.get_children()]
                self.assertTrue(any(str(row[0]).endswith("good.csv") and row[1] == "已跳过" for row in retried_rows))
                self.assertNotEqual(record_path.read_bytes(), record_before)

                _shutdown(app)
                app = self._app()
                app.format_var.set("xlsx")
                app._format_changed()
                with mock.patch.object(gui.messagebox, "showwarning"), mock.patch.object(gui.messagebox, "showerror"), mock.patch.object(gui.messagebox, "showinfo"):
                    app._add_paths([good, wrong, book])
                    _pump(app, 20)
                    self.assertEqual(app._summaries.inspect_calls, 0)
                    for iid in app.table_tree.get_children():
                        if Path(app.table_tree.item(iid, "values")[0]).name == "good.csv":
                            app.table_tree.selection_set(iid)
                            app.table_tree.focus(iid)
                            break
                    app.preview_selected()
                    _pump(app, 20)
                    self.assertEqual(app._summaries.inspect_calls, 1)
                    app.plot_scope_var.set("明确列")
                    app.batch_kind_var.set("line")
                    app.explicit_x_var.set("x")
                    app.explicit_y_var.set("signal")
                    with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(mismatch_out)):
                        app.export_button.invoke()
                    launched = _one_config(mismatch_out)
                    self.assertEqual(launched["request"]["plot"]["y"], ["signal"])
                    self.assertEqual(launched["request"]["plot"]["x"], "x")
                    self.assertEqual(len(launched["request"]["inputs"]), 3)
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                record = load_batch_record(launched["record_path"])
                by_name = {Path(task["source"]).name: task for task in record["tasks"].values()}
                self.assertEqual(by_name["good.csv"]["status"], "succeeded")
                self.assertEqual(by_name["wrong.csv"]["status"], "failed")
                self.assertEqual(by_name["book.xlsx"]["status"], "failed")
                self.assertIn("signal", json.dumps(by_name["wrong.csv"].get("error")))
                self.assertIn("unknown source column", json.dumps(by_name["book.xlsx"].get("error")))
                details = [app.task_tree.item(iid, "values") for iid in app.task_tree.get_children()]
                self.assertTrue(any(str(row[0]).endswith("wrong.csv") and "失败" in row[1] and "signal" in row[2] for row in details))
                self.assertTrue(any(str(row[0]).endswith("book.xlsx") and "失败" in row[1] and "unknown source column" in row[2] for row in details))
                good_file = Path(by_name["good.csv"]["output"]["path"])
                self.assertTrue(good_file.is_file(), good_file)
                opened = load_workbook(good_file, read_only=True)
                try:
                    self.assertIn("good", opened.sheetnames)
                finally:
                    opened.close()
            finally:
                _shutdown(app)

    def test_missing_config_does_not_export_from_window_defaults(self):
        from origin_bridge import gui

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            record = folder / "batch-record.json"
            record.write_text('{"schema_version":1,"kind":"ori-batch-record","tasks":{}}\n', encoding="utf-8")
            app = self._app()
            try:
                app.explicit_y_var.set("guessed")
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", return_value=str(record)),
                    mock.patch.object(gui.messagebox, "askokcancel", return_value=False) as cancel,
                    mock.patch.object(app, "_launch_batch") as launch,
                ):
                    app.continue_batch()
                    cancel.assert_called()
                    launch.assert_not_called()
                self.assertIsNone(app._import_process)
                orphan = {
                    "schema_version": 1,
                    "kind": "ori-batch-config",
                    "name": "别处",
                    "record_path": str(folder / "other-record.json"),
                    "inputs": [str(folder / "a.csv")],
                    "request": {
                        "inputs": [str(folder / "a.csv")],
                        "record_path": str(folder / "other-record.json"),
                        "plot": {"kind": "none", "x": "auto", "y": "auto", "y_error": "auto", "title": "auto", "x_label": "auto", "y_label": "auto"},
                        "overwrite": False,
                        "resume": False,
                        "retry_task_ids": [],
                    },
                }
                given = folder / "given-config.json"
                given.write_text(json.dumps(orphan), encoding="utf-8")
                answers = iter((str(record), str(given)))
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", side_effect=lambda **_kwargs: next(answers)),
                    mock.patch.object(gui.messagebox, "askokcancel", return_value=True),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(app, "_launch_batch") as launch,
                ):
                    app.continue_batch()
                    showerror.assert_called()
                    launch.assert_not_called()
            finally:
                _shutdown(app)


    def test_controls_are_inside_the_client_and_buttons_run_their_commands(self):
        from origin_bridge import gui

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "sample.csv"
            source.write_text("x,y\n1,2\n", encoding="utf-8")
            out = folder / "out"
            app = self._app()
            try:
                app.root.deiconify()
                app._add_paths([source])
                _pump(app, 20)
                for geometry in ("1100x860+30+30", "1200x680+20+20"):
                    app.root.geometry(geometry)
                    app.root.update_idletasks()
                    app.root.update()
                    self.assertEqual(app.main_notebook.tab(app.main_notebook.select(), "text"), "批量")
                    for name, widget in (
                        ("create", app.export_button),
                        ("stop", app.stop_button),
                        ("status", app.status_label),
                        ("tasks", app.task_tree),
                    ):
                        box = widget_in_client(app.root, widget)
                        self.assertTrue(box["inside"], (geometry, name, box))
                    self.assertGreaterEqual(widget_in_client(app.root, app.task_tree)["height"], 40)
                    self.assertEqual(str(app.export_button.cget("state")), "normal")
                    self.assertIn("误差列需要Y使用原列名", str(app.batch_scope_label.cget("text")))
                app.main_notebook.select(1)
                app.root.update_idletasks()
                self.assertTrue(app.column_tree.winfo_exists())
                app.main_notebook.select(0)
                flag = _Flag()
                app._import_kind = "batch"
                app._import_cancel = flag
                app._set_busy(True)
                app.root.geometry("1200x680+20+20")
                app.root.update()
                stop_box = widget_in_client(app.root, app.stop_button)
                self.assertTrue(stop_box["inside"], stop_box)
                self.assertEqual(str(app.stop_button.cget("state")), "normal")
                app.stop_button.invoke()
                self.assertTrue(flag.flag)
                app._import_cancel = None
                app._import_kind = None
                app._set_busy(False)
                captured = {}
                with (
                    mock.patch.object(app, "_start_batch_process", lambda request: captured.setdefault("request", request)),
                    mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)) as ask,
                    mock.patch.object(gui.messagebox, "showerror"),
                ):
                    app.export_button.invoke()
                    ask.assert_called()
                self.assertIn("record_path", captured["request"])
                self.assertTrue(str(captured["request"]["record_path"]).startswith(str(out.resolve())))
            finally:
                _shutdown(app)

    def test_same_file_names_keep_their_own_paths(self):
        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            left = folder / "left"
            right = folder / "right"
            left.mkdir()
            right.mkdir()
            (left / "same.csv").write_text("x,y\n1,2\n", encoding="utf-8")
            (right / "same.csv").write_text("x,y\n3,4\n", encoding="utf-8")
            app = self._app()
            try:
                app._add_paths([left / "same.csv", right / "same.csv"])
                _pump(app, 20)
                rows = []
                for iid in app.table_tree.get_children():
                    app.table_tree.selection_set(iid)
                    app._on_table_selection()
                    rows.append((app.table_tree.item(iid, "values")[0], app.selection_detail_var.get()))
                self.assertEqual(len(rows), 2)
                self.assertNotEqual(rows[0][0], rows[1][0])
                self.assertNotEqual(rows[0][1], rows[1][1])
                self.assertEqual({Path(path).name for _label, path in rows}, {"same.csv"})
                self.assertTrue(all(label.startswith("same.csv") for label, _path in rows))
                self.assertIn(str((left / "same.csv").resolve()), {str(Path(path)) for _label, path in rows})
                self.assertIn(str((right / "same.csv").resolve()), {str(Path(path)) for _label, path in rows})
            finally:
                _shutdown(app)

    def test_long_task_text_and_warnings_open_in_one_detail(self):
        from origin_bridge import gui

        app = self._app()
        try:
            message = "Z" * 2000
            app._apply_batch_progress({
                "type": "task",
                "task_id": "long",
                "index": 0,
                "count": 1,
                "status": "succeeded",
                "source": str(Path("C:/data/deep/long.csv")),
                "error": None,
                "pdf_status": "failed",
                "pdf_errors": [
                    {"graph": "GraphA", "message": message},
                    {"graph": "GraphB", "message": "second figure failed"},
                ],
            })
            iid = app.task_tree.get_children()[0]
            cell = app.task_tree.item(iid, "values")[2]
            self.assertNotIn(message, cell)
            self.assertLess(len(cell), 200)
            self.assertTrue(str(app.task_tree.item(iid, "values")[0]).endswith("long.csv"))
            app.task_tree.selection_set(iid)
            app.detail_button.invoke()
            detail = app._detail_box.get("1.0", "end")
            self.assertIn(message, detail)
            self.assertIn("GraphB", detail)
            self.assertIn("second figure failed", detail)
            self.assertIn("long.csv", detail)
            self.assertIn(str(Path("C:/data/deep/long.csv")), detail)
            warnings = [
                "batch record was restored from the last complete backup",
                "batch record was corrupt and could not be recovered: torn json",
            ]
            with mock.patch.object(gui.messagebox, "showinfo") as showinfo, mock.patch.object(gui.messagebox, "showerror") as showerror:
                app._present_batch_result({
                    "ok": False,
                    "cancelled": False,
                    "counts": {"succeeded": 1, "skipped": 0, "failed": 0, "cancelled": 0, "blocked": 0, "pdf_failed": 1},
                    "tasks": [],
                    "warnings": warnings,
                })
                self.assertEqual(showinfo.call_count, 1)
                showerror.assert_not_called()
                body = showinfo.call_args[0][1]
            self.assertIn(warnings[0], body)
            self.assertIn(warnings[1], body)
            self.assertIn("提示 2 条", app.status_var.get())
            self.assertIn("backup", app.status_var.get())
            self.assertIn("corrupt", app.status_var.get())
            app.task_tree.selection_remove(iid)
            app.detail_button.invoke()
            follow = app._detail_box.get("1.0", "end")
            self.assertIn(warnings[0], follow)
            self.assertIn("torn json", follow)
        finally:
            _shutdown(app)

    def test_numeric_y_with_error_is_rejected_before_named_error_export(self):
        from origin_bridge import gui

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            source = folder / "signal.csv"
            source.write_text("time,signal,err\n0,1,0.1\n1,2,0.2\n", encoding="utf-8")
            out = folder / "out"
            app = self._app()
            try:
                app.format_var.set("xlsx")
                app._format_changed()
                with mock.patch.object(gui.messagebox, "showerror"), mock.patch.object(gui.messagebox, "showwarning"), mock.patch.object(gui.messagebox, "showinfo"):
                    app._add_paths([source])
                    _pump(app, 20)
                    app.table_tree.selection_set(app.table_tree.get_children()[0])
                    app.preview_selected()
                    _pump(app, 20)
                self.assertEqual(app._summaries.inspect_calls, 1)
                app.plot_scope_var.set("明确列")
                app.batch_kind_var.set("line")
                app.explicit_x_var.set("0")
                app.explicit_y_var.set("0")
                app.explicit_error_var.set("1")
                with (
                    mock.patch.object(gui.filedialog, "askdirectory") as ask,
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                ):
                    app.export_button.invoke()
                    ask.assert_not_called()
                    showerror.assert_called()
                    self.assertIn("误差列需要Y使用原列名", showerror.call_args[0][1])
                self.assertEqual(_batch_configs(out), [])
                app.explicit_x_var.set("time")
                app.explicit_y_var.set("signal")
                app.explicit_error_var.set("2")
                with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)), mock.patch.object(gui.messagebox, "showerror") as showerror, mock.patch.object(gui.messagebox, "showinfo"), mock.patch.object(gui.messagebox, "showwarning"):
                    app.export_button.invoke()
                    _pump(app, 60)
                    showerror.assert_not_called()
                self.assertFalse(app._busy, app.status_var.get())
                numeric_error = _one_config(out)
                self.assertEqual(numeric_error["request"]["plot"]["y_error"], {"signal": 2})
                produced = next(Path(task["output"]["path"]) for task in __import__("origin_bridge.batch", fromlist=["load_batch_record"]).load_batch_record(numeric_error["record_path"])["tasks"].values())
                self.assertIn("[[1,2]]", _provenance_text(produced).replace(" ", ""))
                name_out = folder / "named-error"
                app.explicit_error_var.set("err")
                with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(name_out)), mock.patch.object(gui.messagebox, "showerror") as showerror, mock.patch.object(gui.messagebox, "showinfo"), mock.patch.object(gui.messagebox, "showwarning"):
                    app.export_button.invoke()
                    _pump(app, 60)
                    showerror.assert_not_called()
                named = _one_config(name_out)
                self.assertEqual(named["request"]["plot"]["y_error"], {"signal": "err"})
                named_file = next(Path(task["output"]["path"]) for task in __import__("origin_bridge.batch", fromlist=["load_batch_record"]).load_batch_record(named["record_path"])["tasks"].values())
                self.assertIn("[[1,2]]", _provenance_text(named_file).replace(" ", ""))
                index_out = folder / "index-y"
                app.explicit_x_var.set("0")
                app.explicit_y_var.set("1")
                app.explicit_error_var.set("")
                with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(index_out)), mock.patch.object(gui.messagebox, "showerror") as showerror, mock.patch.object(gui.messagebox, "showinfo"), mock.patch.object(gui.messagebox, "showwarning"):
                    app.export_button.invoke()
                    _pump(app, 60)
                    showerror.assert_not_called()
                indexed = _one_config(index_out)
                self.assertEqual(indexed["request"]["plot"]["y"], [1])
                self.assertEqual(indexed["request"]["plot"]["y_error"], {})
                self.assertFalse(app._busy, app.status_var.get())
            finally:
                _shutdown(app)

    def test_second_batch_does_not_replace_the_first_records(self):
        from origin_bridge import gui
        from origin_bridge.batch import load_batch_record

        with tempfile.TemporaryDirectory() as raw:
            folder = Path(raw)
            first = folder / "first.csv"
            second = folder / "second.csv"
            first.write_text("x,y\n1,2\n", encoding="utf-8")
            second.write_text("a,b\n3,4\n", encoding="utf-8")
            out = folder / "out"
            app = self._app()
            try:
                app.format_var.set("xlsx")
                app._format_changed()
                app.batch_kind_var.set("none")
                quiet = (
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showerror"),
                    mock.patch.object(gui.messagebox, "showwarning"),
                )
                with quiet[0], quiet[1], quiet[2]:
                    app._add_paths([first])
                    _pump(app, 20)
                    with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)):
                        app.export_button.invoke()
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                config = _one_config(out)
                record_path = Path(config["record_path"])
                output = next(Path(task["output"]["path"]) for task in load_batch_record(record_path)["tasks"].values())
                output_hash = _sha(output)

                def snapshot() -> dict[str, bytes | None]:
                    paths = [
                        record_path,
                        Path(str(record_path) + ".bak"),
                        Path(str(record_path) + ".previous"),
                        record_path.with_suffix(".log.jsonl"),
                        config_path_for(record_path),
                        Path(str(config_path_for(record_path)) + ".bak"),
                        output,
                    ]
                    return {str(path): path.read_bytes() if path.is_file() else None for path in paths}

                before = snapshot()
                self.assertIsNotNone(before[str(record_path)])
                ctx = multiprocessing.get_context("spawn")
                original_event = ctx.Event
                seen: dict[str, object] = {}

                def make_event():
                    event = original_event()
                    seen["flag"] = event._flag
                    return event

                app._add_paths([second])
                _pump(app, 20)
                with (
                    mock.patch.object(ctx, "Event", make_event),
                    mock.patch.object(ctx.Process, "start", side_effect=OSError("spawn failed")),
                    mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showwarning"),
                ):
                    app.export_button.invoke()
                    showerror.assert_called()
                self.assertIsNone(app._import_cancel)
                self.assertIsNone(getattr(seen.get("flag"), "_semlock", "missing"))
                self.assertEqual(snapshot(), before)
                self.assertEqual([path.resolve() for path in _batch_configs(out)], [config_path_for(record_path)])

                notes = folder / "notes"
                notes.mkdir()
                (notes / "readme.md").write_text("ignore", encoding="utf-8")
                _shutdown(app)
                app = self._app()
                app.format_var.set("xlsx")
                app._format_changed()
                app.batch_kind_var.set("none")
                with quiet[0], quiet[1], quiet[2]:
                    app._add_paths([notes])
                    _pump(app, 20)
                    with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)):
                        app.export_button.invoke()
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                self.assertEqual(snapshot(), before)

                _shutdown(app)
                app = self._app()
                app.format_var.set("xlsx")
                app._format_changed()
                app.batch_kind_var.set("none")
                with quiet[0], quiet[1], quiet[2]:
                    app._add_paths([second])
                    _pump(app, 20)
                    with mock.patch.object(gui.filedialog, "askdirectory", return_value=str(out)):
                        app.export_button.invoke()
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                self.assertEqual(snapshot(), before)
                self.assertEqual(_sha(output), output_hash)
                configs = _batch_configs(out)
                self.assertGreaterEqual(len(configs), 2)
                _shutdown(app)
                app = self._app()
                with (
                    mock.patch.object(gui.filedialog, "askopenfilename", return_value=str(record_path)),
                    mock.patch.object(gui.messagebox, "showinfo"),
                    mock.patch.object(gui.messagebox, "showerror") as showerror,
                    mock.patch.object(gui.messagebox, "showwarning"),
                ):
                    app.continue_button.invoke()
                    showerror.assert_not_called()
                    _pump(app, 60)
                self.assertFalse(app._busy, app.status_var.get())
                self.assertEqual(_sha(output), output_hash)
                resumed = load_batch_record(record_path)
                self.assertEqual(next(iter(resumed["tasks"].values()))["status"], "succeeded")
            finally:
                _shutdown(app)


if __name__ == "__main__":
    unittest.main()
