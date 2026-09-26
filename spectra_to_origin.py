"""Merge 2-column 1D spectra into an Origin .opju project (XYYY / XYXY)."""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import tkinter as tk
import zipfile
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

TEMP_COPY_DIR = "按温度"
SPECTRUM_SUFFIXES = frozenset({".txt", ".dat", ".xy", ".csv", ".tsv"})
HEADER_ROWS = 3
X_LONG_NAME = "2theta"
X_UNITS = "deg"
Y_UNITS = "a.u."
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
EXCEL_HEADER_ROWS = 3
CSV_MANIFEST_NAME = ".spectra-to-origin.json"
__version__ = "1.1.0"

_HEADER_TEMP = re.compile(r"温度\s+([0-9.]+)\s*C")
_HEADER_TIME = re.compile(r"时效\s+([0-9.]+)\s*h")
_HEADER_ETA = re.compile(r"eta\s*=\s*([0-9.]+)", re.IGNORECASE)
_NAME_ETA = re.compile(r"_eta[0-9.]+$", re.IGNORECASE)
_NAME_TEMP = re.compile(r"(?<![0-9A-Za-z])(\d{2,4}C)(?![0-9A-Za-z])", re.IGNORECASE)
_TOKEN_SPLIT = re.compile(r"[_\-\s]+")
_SKIP_TOKEN = re.compile(
    r"^(?:eta[0-9.]+|\d{1,3}|t\d+\.?\d*h|\d+\.?\d*h)$",
    re.IGNORECASE,
)
_NUMERIC_TEXT = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
_NONFINITE_TEXT = re.compile(r"^[+-]?(?:nan|inf(?:inity)?)$", re.IGNORECASE)


class OriginExportError(RuntimeError):
    """Raised when Origin Pro did not produce a usable .opju file."""


@dataclass(frozen=True)
class Spectrum:
    path: Path
    x_text: tuple[str, ...]
    y_text: tuple[str, ...]
    long_name: str
    comment: str
    temperature_tag: str | None

    @property
    def n_points(self) -> int:
        return len(self.x_text)

    @property
    def x_values(self) -> list[float]:
        return [float(value) for value in self.x_text]

    @property
    def y_values(self) -> list[float]:
        return [float(value) for value in self.y_text]


@dataclass(frozen=True)
class SheetTable:
    name: str
    long_names: tuple[str, ...]
    units: tuple[str, ...]
    comments: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class SheetSpec:
    name: str
    spectra: tuple[Spectrum, ...]
    layout: str


@dataclass(frozen=True)
class AxisLabels:
    x_name: str = X_LONG_NAME
    x_unit: str = X_UNITS
    y_name: str = "Intensity"
    y_unit: str = Y_UNITS


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码：{path}")


def _long_name_from_stem(stem: str) -> str:
    return _NAME_ETA.sub("", stem)


def _temperature_tag(stem: str, header: str) -> str | None:
    match = _NAME_TEMP.search(stem)
    if match:
        return match.group(1).upper()
    match = _HEADER_TEMP.search(header)
    if match:
        return f"{match.group(1)}C"
    return None


def _comment_from_header(header: str, stem: str) -> str:
    parts: list[str] = []
    eta = _HEADER_ETA.search(header)
    temp = _HEADER_TEMP.search(header)
    time = _HEADER_TIME.search(header)
    if eta:
        parts.append(f"eta={eta.group(1)}")
    if temp:
        parts.append(f"{temp.group(1)} C")
    if time:
        parts.append(f"{time.group(1)} h")
    if not parts:
        parts.append(_long_name_from_stem(stem))
    return "; ".join(parts)


def parse_spectra(files: list[Path]) -> list[Spectrum]:
    return [parse_spectrum(path) for path in files]


def _split_spectrum_row(line: str) -> list[str]:
    delimiter = "," if "," in line else ";" if ";" in line else "\t" if "\t" in line else None
    if delimiter is None:
        return line.split()
    try:
        rows = list(csv.reader([line], delimiter=delimiter, skipinitialspace=True, strict=True))
    except csv.Error as exc:
        raise ValueError(f"分隔符格式无效：{line}") from exc
    return [cell.strip() for cell in rows[0]]


def _normalized_numeric_text(value: str, path: Path, line_number: int, raw_line: str) -> str:
    token = value.strip()
    if _NONFINITE_TEXT.fullmatch(token):
        raise ValueError(f"{path.name} 第 {line_number} 行包含非有限数值：{raw_line}")
    normalized = token.replace("d", "e").replace("D", "E")
    if not _NUMERIC_TEXT.fullmatch(normalized):
        raise ValueError(f"{path.name} 第 {line_number} 行不是数值：{raw_line}")
    if not math.isfinite(float(normalized)):
        raise ValueError(f"{path.name} 第 {line_number} 行包含非有限数值：{raw_line}")
    return normalized


def _natural_sort_key(value: str) -> tuple[tuple[int, int | str], ...]:
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.casefold())
        for part in re.split(r"(\d+)", value)
        if part
    )


def parse_spectrum(path: Path) -> Spectrum:
    try:
        text = read_text(path)
    except OSError as exc:
        raise ValueError(f"无法读取文件 {path}：{exc}") from exc
    header_lines: list[str] = []
    x_text: list[str] = []
    y_text: list[str] = []
    header_seen = False
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            header_lines.append(line.lstrip("# ").strip())
            continue
        try:
            parts = _split_spectrum_row(line)
        except ValueError as exc:
            raise ValueError(f"{path.name} 第 {line_number} 行分隔符格式无效：{raw_line}") from exc
        if len(parts) != 2:
            raise ValueError(f"{path.name} 第 {line_number} 行不是两列：{raw_line}")
        if not x_text and not header_seen and all(
            not _NUMERIC_TEXT.fullmatch(part.replace("d", "e").replace("D", "E"))
            and not _NONFINITE_TEXT.fullmatch(part)
            for part in parts
        ):
            if all(re.search(r"[A-Za-z\u0080-\uffff]", part) for part in parts):
                header_seen = True
                continue
        normalized = [
            _normalized_numeric_text(part, path, line_number, raw_line)
            for part in parts
        ]
        x_text.append(normalized[0])
        y_text.append(normalized[1])
    if len(x_text) < 2:
        raise ValueError(f"{path.name} 有效数据点不足：{len(x_text)}")
    header = "    ".join(header_lines)
    stem = path.stem
    return Spectrum(
        path=path,
        x_text=tuple(x_text),
        y_text=tuple(y_text),
        long_name=_long_name_from_stem(stem),
        comment=_comment_from_header(header, stem),
        temperature_tag=_temperature_tag(stem, header),
    )


def under_temp_copy_dir(path: Path) -> bool:
    return TEMP_COPY_DIR in path.parts


def collect_txt_files(root: Path) -> list[Path]:
    if root.is_file():
        if root.suffix.lower() not in SPECTRUM_SUFFIXES:
            raise ValueError(f"不是支持的谱线文件：{root}")
        return [root.resolve()]
    if not root.is_dir():
        raise ValueError(f"路径不存在：{root}")

    found: list[Path] = []
    walk_errors: list[OSError] = []
    for dirpath, dirnames, filenames in os.walk(root, onerror=walk_errors.append):
        dirnames.sort(key=_natural_sort_key)
        filenames.sort(key=_natural_sort_key)
        txt_here = [
            Path(dirpath) / name
            for name in filenames
            if Path(name).suffix.lower() in SPECTRUM_SUFFIXES
        ]
        found.extend(path.resolve() for path in txt_here)
    if walk_errors:
        error = walk_errors[0]
        raise ValueError(f"无法读取文件夹 {root}：{error}")
    resolved_root = root.resolve()
    found.sort(key=lambda path: _natural_sort_key(path.relative_to(resolved_root).as_posix()))
    return found


def _same_temp_copy(left: Path, right: Path) -> bool:
    left = left.resolve()
    right = right.resolve()

    def counterparts(path: Path) -> set[Path]:
        related: set[Path] = set()
        for parent in path.parents:
            if parent.name.casefold() == TEMP_COPY_DIR.casefold():
                relative_path = path.relative_to(parent)
                related.add((parent.parent / relative_path).resolve())
        return related

    if right not in counterparts(left) and left not in counterparts(right):
        return False
    try:
        return left.read_bytes() == right.read_bytes()
    except OSError:
        return False


def merge_file_list(existing: list[Path], incoming: list[Path]) -> list[Path]:
    out = list(existing)
    seen = {path.resolve() for path in existing}
    for path in incoming:
        resolved = path.resolve()
        if resolved in seen:
            continue
        duplicate_index = next(
            (
                index
                for index, old in enumerate(out)
                if old.name.casefold() == path.name.casefold() and _same_temp_copy(old, path)
            ),
            None,
        )
        if duplicate_index is not None:
            old = out[duplicate_index]
            if under_temp_copy_dir(old) and not under_temp_copy_dir(path):
                out[duplicate_index] = path
                seen.discard(old.resolve())
                seen.add(resolved)
            continue
        out.append(path)
        seen.add(resolved)
    return out


def collect_from_user_paths(paths: list[Path]) -> tuple[list[Path], list[str]]:
    incoming: list[Path] = []
    errors: list[str] = []
    for path in paths:
        try:
            incoming.extend(collect_txt_files(path))
        except ValueError as exc:
            errors.append(str(exc))
    return incoming, errors


def parse_drop_paths(payload: str | bytes | list[str] | list[Path]) -> list[Path]:
    if isinstance(payload, (list, tuple)):
        return [Path(str(item).strip().strip('"')) for item in payload if str(item).strip()]
    text = payload.decode("utf-16le") if isinstance(payload, (bytes, bytearray)) else str(payload)
    parts = re.split(r"[\r\n\x00]+", text.strip().strip("\x00"))
    return [Path(part.strip().strip('"')) for part in parts if part.strip()]


def load_from_drop_payload(existing: list[Path], payload: str | bytes | list[str] | list[Path]) -> list[Path]:
    incoming, _errors = collect_from_user_paths(parse_drop_paths(payload))
    return merge_file_list(existing, incoming)


def shared_x_grid(spectra: list[Spectrum]) -> bool:
    if not spectra:
        return False
    try:
        first = tuple(Decimal(value) for value in spectra[0].x_text)
        return all(
            tuple(Decimal(value) for value in item.x_text) == first
            for item in spectra[1:]
        )
    except (InvalidOperation, ValueError):
        return False


def infer_layout(spectra: list[Spectrum]) -> str:
    if shared_x_grid(spectra):
        return "XYYY"
    return "XYXY"


def resolve_layout(spectra: list[Spectrum], layout: str) -> str:
    text = (layout or "auto").strip().upper()
    if text in {"", "AUTO"}:
        return infer_layout(spectra)
    if text in {"XYYY", "XYXY"}:
        return text
    raise ValueError(f"不支持的布局：{layout}")


def _xyyy_table(name: str, spectra: list[Spectrum], axis_labels: AxisLabels) -> SheetTable:
    if not shared_x_grid(spectra):
        names = ", ".join(item.path.name for item in spectra[:4])
        raise ValueError(f"X 网格不一致，不能用 XYYY。请改选 XYXY。涉及文件：{names}")
    long_names = (axis_labels.x_name, *(item.long_name for item in spectra))
    units = (axis_labels.x_unit, *(axis_labels.y_unit for _ in spectra))
    comments = ("", *(item.comment for item in spectra))
    rows = []
    x_text = spectra[0].x_text
    for index, x_value in enumerate(x_text):
        rows.append((x_value, *(item.y_text[index] for item in spectra)))
    return SheetTable(name, long_names, units, comments, tuple(rows))


def _xyxy_table(name: str, spectra: list[Spectrum], axis_labels: AxisLabels) -> SheetTable:
    long_names: list[str] = []
    units: list[str] = []
    comments: list[str] = []
    for item in spectra:
        long_names.extend([axis_labels.x_name, item.long_name])
        units.extend([axis_labels.x_unit, axis_labels.y_unit])
        comments.extend(["", item.comment])
    n_rows = max(item.n_points for item in spectra)
    rows: list[tuple[str, ...]] = []
    for index in range(n_rows):
        cells: list[str] = []
        for item in spectra:
            if index < item.n_points:
                cells.extend([item.x_text[index], item.y_text[index]])
            else:
                cells.extend(["", ""])
        rows.append(tuple(cells))
    return SheetTable(name, tuple(long_names), tuple(units), tuple(comments), tuple(rows))


def build_table(
    name: str,
    spectra: list[Spectrum],
    layout: str,
    axis_labels: AxisLabels = AxisLabels(),
) -> SheetTable:
    if not spectra:
        raise ValueError("没有谱线")
    layout = resolve_layout(spectra, layout)
    if layout == "XYYY":
        return _xyyy_table(name, spectra, axis_labels)
    if layout == "XYXY":
        return _xyxy_table(name, spectra, axis_labels)
    raise ValueError(f"不支持的布局：{layout}")


def _stem_tokens(stem: str) -> list[str]:
    parts = [part for part in _TOKEN_SPLIT.split(stem) if part]
    for match in _NAME_TEMP.finditer(stem):
        token = match.group(1)
        if token not in parts:
            parts.append(token)
    return parts


def _usable_token(token: str) -> bool:
    return not _SKIP_TOKEN.match(token)


def _score_index_groups(groups: list[tuple[str, list[int]]], sample: str) -> int:
    sizes = [len(indexes) for _name, indexes in groups]
    n_items = sum(sizes)
    n_groups = len(groups)
    if n_items == 0 or n_groups < 2:
        return -10_000
    score = 0
    if any(name == "other" for name, _indexes in groups):
        other_n = next(len(indexes) for name, indexes in groups if name == "other")
        score -= 50 * other_n
    else:
        score += 25
    if min(sizes) >= 2:
        score += 20
    score -= 8 * sum(1 for size in sizes if size == 1)
    if n_groups <= max(2, n_items // 2):
        score += 12
    else:
        score -= 12
    if any(char.isalpha() for char in sample):
        score += 10
    score += min(n_groups, 8)
    mean = n_items / n_groups
    score -= int(sum(abs(size - mean) for size in sizes))
    return score


def _groups_from_assignment(assignment: list[str | None]) -> list[tuple[str, list[int]]]:
    order: list[str] = []
    buckets: dict[str, list[int]] = {}
    other: list[int] = []
    for index, name in enumerate(assignment):
        if name is None:
            other.append(index)
            continue
        if name not in buckets:
            order.append(name)
            buckets[name] = []
        buckets[name].append(index)
    groups = [(name, buckets[name]) for name in order]
    if other:
        groups.append(("other", other))
    return groups


def suggest_filename_groups(stems: list[str]) -> list[tuple[str, list[int]]]:
    n_items = len(stems)
    if n_items == 0:
        return []
    if n_items == 1:
        return [("all", [0])]

    token_lists = [_stem_tokens(stem) for stem in stems]
    candidates: list[tuple[int, list[tuple[str, list[int]]]]] = []

    lengths = {len(tokens) for tokens in token_lists}
    if len(lengths) == 1:
        width = next(iter(lengths))
        for column in range(width):
            values = [token_lists[index][column] for index in range(n_items)]
            if all(not _usable_token(value) for value in values):
                continue
            unique = list(dict.fromkeys(values))
            if len(unique) < 2 or len(unique) == n_items:
                continue
            groups = _groups_from_assignment(list(values))
            candidates.append((_score_index_groups(groups, values[0]), groups))

    freq: Counter[str] = Counter()
    for tokens in token_lists:
        for token in dict.fromkeys(tokens):
            if _usable_token(token):
                freq[token] += 1
    covering = {token for token, count in freq.items() if 2 <= count < n_items}
    if covering:
        assignment: list[str | None] = []
        for tokens in token_lists:
            found = [token for token in tokens if token in covering]
            assignment.append(found[0] if found else None)
        groups = _groups_from_assignment(assignment)
        named = [name for name, _indexes in groups if name != "other"]
        if len(named) >= 2:
            sample = named[0]
            candidates.append((_score_index_groups(groups, sample), groups))

    if not candidates:
        return [("all", list(range(n_items)))]
    candidates.sort(key=lambda item: item[0], reverse=True)
    best = candidates[0][1]
    named = [name for name, _indexes in best if name != "other"]
    if len(named) < 2:
        return [("all", list(range(n_items)))]
    return best


def groups_from_filenames(spectra: list[Spectrum]) -> list[tuple[str, list[Spectrum]]]:
    grouped = suggest_filename_groups([item.path.stem for item in spectra])
    return [(name, [spectra[index] for index in indexes]) for name, indexes in grouped]


def partition_even(n_items: int, n_groups: int, prefix: str = "group") -> list[tuple[str, list[int]]]:
    if n_items <= 0:
        return []
    if n_groups < 1:
        raise ValueError("分组数至少为 1")
    n_groups = min(n_groups, n_items)
    base, extra = divmod(n_items, n_groups)
    groups: list[tuple[str, list[int]]] = []
    start = 0
    for index in range(n_groups):
        size = base + (1 if index < extra else 0)
        end = start + size
        groups.append((f"{prefix}{index + 1}", list(range(start, end))))
        start = end
    return groups


def partition_spectra(spectra: list[Spectrum], n_groups: int, prefix: str = "group") -> list[tuple[str, list[Spectrum]]]:
    return [
        (name, [spectra[index] for index in indexes])
        for name, indexes in partition_even(len(spectra), n_groups, prefix)
    ]


def assignments_from_groups(n_items: int, groups: list[tuple[str, list[int]]]) -> list[str]:
    assigned = ["all"] * n_items
    for name, indexes in groups:
        for index in indexes:
            assigned[index] = name
    return assigned


def groups_from_assignments(items: list, assignments: list[str]) -> list[tuple[str, list]]:
    if len(items) != len(assignments):
        raise ValueError("分组标记与谱线数量不一致")
    order: list[str] = []
    buckets: dict[str, list] = {}
    for item, name in zip(items, assignments):
        label = name.strip() or "group"
        if label not in buckets:
            order.append(label)
            buckets[label] = []
        buckets[label].append(item)
    return [(name, buckets[name]) for name in order]


def resolve_groups(
    spectra: list[Spectrum],
    groups: list[tuple[str, list[Spectrum]]] | None = None,
    split_temp: bool = False,
    n_groups: int | None = None,
) -> list[tuple[str, list[Spectrum]]]:
    if n_groups is not None and n_groups < 1:
        raise ValueError("分组数至少为 1")
    if groups is not None:
        return [(name, list(items)) for name, items in groups if items]
    if n_groups is not None and n_groups > 1:
        return partition_spectra(spectra, n_groups)
    if split_temp:
        return groups_from_filenames(spectra)
    return [("spectra", list(spectra))]


def build_sheet_specs(
    groups: list[tuple[str, list[Spectrum]]],
    layout: str,
) -> list[SheetSpec]:
    if not groups:
        raise ValueError("没有分组")
    all_spectra = [item for _name, items in groups for item in items]
    if not all_spectra:
        raise ValueError("没有谱线")
    specs: list[SheetSpec] = []
    for name, items in groups:
        if not items:
            continue
        resolved = resolve_layout(items, layout)
        if resolved == "XYYY" and not shared_x_grid(items):
            names = ", ".join(item.path.name for item in items[:4])
            raise ValueError(f"X 网格不一致，不能用 XYYY。请改选 XYXY。涉及文件：{names}")
        specs.append(SheetSpec(name=name, spectra=tuple(items), layout=resolved))
    if not specs:
        raise ValueError("没有可导出的 sheet")
    return specs


def group_by_temperature(spectra: list[Spectrum]) -> dict[str, list[Spectrum]]:
    grouped: dict[str, list[Spectrum]] = {}
    for item in spectra:
        tag = item.temperature_tag or "other"
        grouped.setdefault(tag, []).append(item)
    return grouped


def sheet_name(raw: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", raw).strip() or "sheet"
    return cleaned[:31]


def origin_sheet_name(raw: str) -> str:
    token = re.sub(r"[^0-9A-Za-z]+", "", str(raw).strip()) or "sheet"
    if token[0].isdigit():
        token = "T" + token
    return token[:13]


def unique_origin_names(raw_names: list[str], limit: int = 13) -> list[str]:
    used: set[str] = set()
    names: list[str] = []
    for raw in raw_names:
        base = origin_sheet_name(raw)[:limit]
        candidate = base
        serial = 2
        while candidate.lower() in used:
            suffix = str(serial)
            candidate = (base[: max(1, limit - len(suffix))] + suffix)[:limit]
            serial += 1
        used.add(candidate.lower())
        names.append(candidate)
    return names


def readme_rows(
    spectra: list[Spectrum],
    layout: str,
    group_names: list[str],
    axis_labels: AxisLabels = AxisLabels(),
    group_layouts: list[tuple[str, str]] | None = None,
) -> list[list[str]]:
    shared = "yes" if shared_x_grid(spectra) else "no"
    rows = [
        ["layout", layout],
        ["n_spectra", str(len(spectra))],
        ["n_points_first", str(spectra[0].n_points if spectra else 0)],
        ["shared_x", shared],
        ["groups", ", ".join(group_names)],
        ["x_axis", f"{axis_labels.x_name} ({axis_labels.x_unit})"],
        ["y_axis", f"{axis_labels.y_name} ({axis_labels.y_unit})"],
        ["header_rows", "Long Name / Units / Comments"],
        [],
        ["index", "long_name", "n_points", "group_hint", "comment", "path"],
    ]
    if group_layouts:
        rows.insert(6, ["group_layouts", ", ".join(f"{name}: {value}" for name, value in group_layouts)])
    for index, item in enumerate(spectra, start=1):
        rows.append(
            [
                str(index),
                item.long_name,
                str(item.n_points),
                item.temperature_tag or "",
                item.comment,
                str(item.path),
            ]
        )
    return rows


def _append_metadata_row(sheet, values: list[str] | tuple[str, ...]) -> None:
    sheet.append(list(values))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _write_data_sheet(workbook: Workbook, table: SheetTable, worksheet_name: str | None = None) -> None:
    sheet = workbook.create_sheet(worksheet_name or sheet_name(table.name))
    _append_metadata_row(sheet, table.long_names)
    _append_metadata_row(sheet, table.units)
    _append_metadata_row(sheet, table.comments)
    for row in table.rows:
        sheet.append([_excel_value(cell) for cell in row])
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = "A4"
    for index, name in enumerate(table.long_names, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(36, max(12, len(name) + 2))


def _excel_value(cell: str):
    if cell == "":
        return None
    return float(cell)


def unique_sheet_names(raw_names: list[str], limit: int = 31) -> list[str]:
    used = {"readme"}
    names: list[str] = []
    for raw in raw_names:
        base = sheet_name(raw)[:limit]
        candidate = base
        serial = 2
        while candidate.casefold() in used:
            suffix = f"_{serial}"
            candidate = f"{base[: max(1, limit - len(suffix))]}{suffix}"
            serial += 1
        used.add(candidate.casefold())
        names.append(candidate)
    return names


def _safe_csv_stem(raw: str) -> str:
    cleaned = "".join(char if (char.isalnum() or char in "-_.") else "_" for char in str(raw).strip())
    cleaned = cleaned.strip(" ._") or "group"
    if cleaned in {".", ".."}:
        cleaned = "group"
    if cleaned.split(".", 1)[0].upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        cleaned = f"_{cleaned}"
    return cleaned[:100]


def unique_csv_names(raw_names: list[str]) -> list[str]:
    used: set[str] = set()
    names: list[str] = []
    for raw in raw_names:
        base = _safe_csv_stem(raw)
        candidate = base
        serial = 2
        while f"{candidate}.csv".casefold() in used:
            suffix = f"_{serial}"
            candidate = f"{base[:100 - len(suffix)]}{suffix}"
            serial += 1
        filename = f"{candidate}.csv"
        used.add(filename.casefold())
        names.append(filename)
    return names


def _csv_output_plan(output_xlsx: Path, specs: list[SheetSpec]) -> tuple[Path, list[Path]]:
    output_xlsx = Path(output_xlsx)
    if len(specs) == 1:
        csv_path = output_xlsx.with_suffix(".csv")
        return csv_path, [csv_path]
    directory = output_xlsx.with_name(f"{output_xlsx.stem}_csv")
    paths = [directory / name for name in unique_csv_names([spec.name for spec in specs])]
    return directory, paths


def _validate_excel_dimensions(specs: list[SheetSpec], n_spectra: int) -> None:
    for spec in specs:
        if spec.layout == "XYYY":
            columns = 1 + len(spec.spectra)
            rows = spec.spectra[0].n_points + EXCEL_HEADER_ROWS
        else:
            columns = 2 * len(spec.spectra)
            rows = max(item.n_points for item in spec.spectra) + EXCEL_HEADER_ROWS
        if columns > EXCEL_MAX_COLUMNS:
            raise ValueError(f"分组“{spec.name}”有 {columns} 列，超过 Excel 上限 {EXCEL_MAX_COLUMNS}")
        if rows > EXCEL_MAX_ROWS:
            raise ValueError(f"分组“{spec.name}”有 {rows} 行，超过 Excel 上限 {EXCEL_MAX_ROWS}")
    readme_rows = n_spectra + 11
    if readme_rows > EXCEL_MAX_ROWS:
        raise ValueError(f"readme 工作表有 {readme_rows} 行，超过 Excel 上限 {EXCEL_MAX_ROWS}")


def _path_key(path: Path) -> str:
    return os.path.normcase(str(Path(path).expanduser().resolve(strict=False)))


def _validate_target_paths(input_paths: list[Path], target_paths: list[Path]) -> None:
    input_by_key = {_path_key(path): Path(path) for path in input_paths}
    collisions = [input_by_key[_path_key(target)] for target in target_paths if _path_key(target) in input_by_key]
    if collisions:
        paths = ", ".join(str(path) for path in dict.fromkeys(collisions))
        raise ValueError(f"导出目标与输入文件冲突：{paths}")


def _validate_no_input_in_directory(input_paths: list[Path], directory: Path) -> None:
    base = Path(directory).expanduser().resolve(strict=False)
    for path in input_paths:
        candidate = Path(path).expanduser().resolve(strict=False)
        try:
            candidate.relative_to(base)
        except ValueError:
            continue
        raise ValueError(f"多组 CSV 输出目录包含输入文件，拒绝替换目录：{candidate}")


def _atomic_write_file(path: Path, writer, *, validate_xlsx: bool = False) -> None:
    path = Path(path).expanduser().resolve(strict=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, raw_temp = tempfile.mkstemp(
        prefix=f".{path.stem}_",
        suffix=f".tmp{path.suffix}",
        dir=str(path.parent),
    )
    os.close(handle)
    temp_path = Path(raw_temp)
    try:
        writer(temp_path)
        if not temp_path.is_file() or temp_path.stat().st_size == 0:
            raise ValueError(f"导出文件不存在或为空：{path}")
        if validate_xlsx and not zipfile.is_zipfile(temp_path):
            raise ValueError(f"生成的工作簿无效：{path}")
        os.replace(temp_path, path)
    finally:
        try:
            temp_path.unlink()
        except FileNotFoundError:
            pass


def _csv_directory_is_managed(path: Path) -> bool:
    if not path.exists():
        return True
    if path.is_symlink():
        raise ValueError(f"多组 CSV 目标是符号链接，拒绝替换：{path}")
    if not path.is_dir():
        raise ValueError(f"多组 CSV 目标不是目录：{path}")
    children = list(path.iterdir())
    if not children:
        return True
    manifest_path = path / CSV_MANIFEST_NAME
    if not manifest_path.is_file():
        raise ValueError(f"多组 CSV 目录已存在且不属于本工具，拒绝覆盖：{path}")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        files = manifest["files"]
        if not isinstance(files, list) or any(not isinstance(name, str) or Path(name).name != name for name in files):
            raise ValueError
        expected = {CSV_MANIFEST_NAME, *files}
    except (OSError, ValueError, KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"多组 CSV 目录清单无效，拒绝覆盖：{path}") from exc
    actual = {child.name for child in children}
    if actual != expected or any(not child.is_file() for child in children):
        raise ValueError(f"多组 CSV 目录包含清单之外的文件，拒绝覆盖：{path}")
    return True


def _write_csv_directory(path: Path, tables: list[SheetTable], filenames: list[str]) -> Path:
    path = Path(path).expanduser().resolve(strict=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    _csv_directory_is_managed(path)
    stage = Path(tempfile.mkdtemp(prefix=f".{path.name}_stage_", dir=str(path.parent)))
    backup: Path | None = None
    preserve_backup = False
    try:
        for table, filename in zip(tables, filenames):
            _write_csv_contents(stage / filename, table)
        (stage / CSV_MANIFEST_NAME).write_text(
            json.dumps({"version": 1, "files": filenames}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        if path.exists():
            _csv_directory_is_managed(path)
            backup = Path(tempfile.mkdtemp(prefix=f".{path.name}_backup_", dir=str(path.parent)))
            backup.rmdir()
            os.replace(path, backup)
        try:
            os.replace(stage, path)
            stage = Path()
        except OSError as stage_error:
            if backup is not None and backup.exists():
                try:
                    os.replace(backup, path)
                    backup = None
                except OSError as rollback_error:
                    preserve_backup = True
                    raise OriginExportError(
                        f"新 CSV 目录安装失败，旧目录保留在 {backup}；恢复失败：{rollback_error}"
                    ) from rollback_error
            raise stage_error
        if backup is not None:
            try:
                try:
                    _csv_directory_is_managed(backup)
                except ValueError as exc:
                    preserve_backup = True
                    raise OriginExportError(
                        f"旧 CSV 目录在替换期间出现清单之外的文件，已保留在 {backup}：{exc}"
                    ) from exc
                shutil.rmtree(backup)
                backup = None
            except OSError as exc:
                preserve_backup = True
                raise OriginExportError(f"CSV 已写出，但旧目录清理失败，仍保留在 {backup}：{exc}") from exc
        return path
    finally:
        if stage != Path() and stage.exists():
            shutil.rmtree(stage, ignore_errors=True)
        if backup is not None and backup.exists() and not preserve_backup:
            shutil.rmtree(backup, ignore_errors=True)


def _write_csv_contents(path: Path, table: SheetTable) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(table.long_names)
        writer.writerow(table.units)
        writer.writerow(table.comments)
        writer.writerows(table.rows)


def write_xlsx(
    path: Path,
    tables: list[SheetTable],
    spectra: list[Spectrum],
    layout: str,
    group_names: list[str],
    axis_labels: AxisLabels = AxisLabels(),
    group_layouts: list[tuple[str, str]] | None = None,
) -> None:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    names = unique_sheet_names([table.name for table in tables])
    for table, name in zip(tables, names):
        _write_data_sheet(workbook, table, name)
    readme = workbook.create_sheet("readme")
    for row in readme_rows(spectra, layout, group_names, axis_labels, group_layouts):
        _append_metadata_row(readme, row)
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 28
    readme.column_dimensions["F"].width = 80
    _atomic_write_file(path, workbook.save, validate_xlsx=True)


def write_csv(path: Path, table: SheetTable) -> None:
    _atomic_write_file(path, lambda temp_path: _write_csv_contents(temp_path, table))


def build_export_tables(
    spectra: list[Spectrum],
    layout: str,
    groups: list[tuple[str, list[Spectrum]]],
    axis_labels: AxisLabels = AxisLabels(),
) -> list[SheetTable]:
    specs = build_sheet_specs(groups, layout)
    _validate_excel_dimensions(specs, len(spectra))
    return [build_table(spec.name, list(spec.spectra), spec.layout, axis_labels) for spec in specs]


def export_spectra(
    files: list[Path],
    output_xlsx: Path,
    layout: str = "auto",
    split_temp: bool = False,
    n_groups: int | None = None,
    groups: list[tuple[str, list[Spectrum]]] | None = None,
    write_csv_copy: bool = True,
    axis_labels: AxisLabels = AxisLabels(),
) -> tuple[Path, Path | None]:
    if not files:
        raise ValueError("没有谱线文件")
    output_xlsx = Path(output_xlsx).expanduser()
    if output_xlsx.suffix.lower() != ".xlsx":
        output_xlsx = output_xlsx.with_suffix(".xlsx")
    if groups:
        spectra = [item for _name, items in groups for item in items]
    else:
        spectra = parse_spectra(files)
    if not spectra:
        raise ValueError("没有谱线文件")
    resolved_groups = resolve_groups(spectra, groups=groups, split_temp=split_temp, n_groups=n_groups)
    specs = build_sheet_specs(resolved_groups, layout)
    _validate_excel_dimensions(specs, len(spectra))
    tables = [build_table(spec.name, list(spec.spectra), spec.layout, axis_labels) for spec in specs]
    csv_target, csv_files = _csv_output_plan(output_xlsx, specs) if write_csv_copy else (None, [])
    target_paths = [Path(output_xlsx), *csv_files]
    _validate_target_paths([item.path for item in spectra], target_paths)
    if write_csv_copy and len(specs) > 1:
        _validate_no_input_in_directory([item.path for item in spectra], csv_target)
        _csv_directory_is_managed(csv_target)
    resolved_layouts = {spec.layout for spec in specs}
    resolved_layout = next(iter(resolved_layouts)) if len(resolved_layouts) == 1 else "mixed"
    write_xlsx(
        output_xlsx,
        tables,
        spectra,
        resolved_layout,
        [spec.name for spec in specs],
        axis_labels,
        [(spec.name, spec.layout) for spec in specs],
    )
    csv_path = None
    if write_csv_copy:
        if len(specs) == 1:
            csv_path = csv_files[0]
            write_csv(csv_path, tables[0])
        else:
            csv_path = _write_csv_directory(csv_target, tables, [path.name for path in csv_files])
    return output_xlsx, csv_path


def origin_process_running() -> bool:
    if sys.platform != "win32":
        return False
    for executable in ("Origin64.exe", "Origin.exe"):
        try:
            output = subprocess.check_output(
                ["tasklist", "/FI", f"IMAGENAME eq {executable}", "/NH"],
                text=True,
                errors="ignore",
            )
        except (OSError, subprocess.CalledProcessError):
            continue
        if executable.casefold() in output.casefold():
            return True
    return False


def close_origin_app(op, *, started: bool, timeout: float = 8.0) -> None:
    """Close the COM session only when this export started Origin itself."""
    if not started:
        return
    try:
        op.exit()
    except Exception:
        pass
    if sys.platform != "win32":
        return
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline and origin_process_running():
        time.sleep(0.2)


def _load_originpro():
    try:
        import originpro as op
    except ImportError as exc:
        raise OriginExportError(
            "当前 Python 无法 import originpro，不能生成 .opju。"
            "请在已安装 Origin Pro 的环境里运行，并确认 originpro / OriginExt 可用。"
        ) from exc
    return op


def _pad_column(values: list[float], n_rows: int) -> list[float]:
    if len(values) >= n_rows:
        return values[:n_rows]
    return values + [float("nan")] * (n_rows - len(values))


def _fill_origin_sheet(
    sheet,
    spectra: list[Spectrum],
    layout: str,
    axis_labels: AxisLabels = AxisLabels(),
) -> None:
    layout = resolve_layout(spectra, layout)
    if layout == "XYYY":
        if not shared_x_grid(spectra):
            names = ", ".join(item.path.name for item in spectra[:4])
            raise ValueError(f"X 网格不一致，不能用 XYYY。请改选 XYXY。涉及文件：{names}")
        columns = [spectra[0].x_values] + [item.y_values for item in spectra]
        longs = [axis_labels.x_name, *(item.long_name for item in spectra)]
        units = [axis_labels.x_unit, *(axis_labels.y_unit for _ in spectra)]
        comments = ["", *(item.comment for item in spectra)]
        axes = "x" + "y" * len(spectra)
    elif layout == "XYXY":
        n_rows = max(item.n_points for item in spectra)
        columns = []
        longs: list[str] = []
        units: list[str] = []
        comments: list[str] = []
        for item in spectra:
            columns.append(_pad_column(item.x_values, n_rows))
            columns.append(_pad_column(item.y_values, n_rows))
            longs.extend([axis_labels.x_name, item.long_name])
            units.extend([axis_labels.x_unit, axis_labels.y_unit])
            comments.extend(["", item.comment])
        axes = "xy" * len(spectra)
    else:
        raise ValueError(f"不支持的布局：{layout}")
    sheet.cols = len(columns)
    sheet.from_list2(columns, 0, 0)
    sheet.set_labels(longs, "L")
    sheet.set_labels(units, "U")
    sheet.set_labels(comments, "C")
    sheet.cols_axis(axes)


def _plot_origin_sheet(
    op,
    sheet,
    spectra: list[Spectrum],
    layout: str,
    graph_name: str,
    axis_labels: AxisLabels = AxisLabels(),
):
    graph = op.new_graph(lname=graph_name, hidden=False)
    if graph is None:
        raise OriginExportError(f"Origin 未能创建图：{graph_name}")
    layer = graph[0]
    layout = resolve_layout(spectra, layout)
    if layout == "XYYY":
        for y_col in range(1, len(spectra) + 1):
            plot = layer.add_plot(sheet, coly=y_col, colx=0, type="l")
            if plot is None:
                raise OriginExportError(f"Origin 未能把第 {y_col} 列画进 {graph_name}")
    else:
        for index in range(len(spectra)):
            x_col = 2 * index
            y_col = x_col + 1
            plot = layer.add_plot(sheet, coly=y_col, colx=x_col, type="l")
            if plot is None:
                raise OriginExportError(f"Origin 未能把 {spectra[index].long_name} 画进 {graph_name}")
    n_plots = len(list(layer.obj.DataPlots))
    if n_plots != len(spectra):
        raise OriginExportError(
            f"{graph_name} 应有 {len(spectra)} 条曲线，Origin 实际只画了 {n_plots} 条。"
        )
    layer.axis("x").title = f"{axis_labels.x_name} ({axis_labels.x_unit})"
    layer.axis("y").title = f"{axis_labels.y_name} ({axis_labels.y_unit})"
    layer.rescale()
    layer.lt_exec("legend -r")
    return graph


def commit_exported_file(src: Path, dest: Path) -> Path:
    """Install a finished Origin file atomically, including cross-volume saves."""
    src = src.expanduser().resolve()
    dest = dest.expanduser().resolve()
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not src.exists() or src.stat().st_size == 0:
        raise OriginExportError(f"Origin 临时文件不存在或为空：{src}")
    handle, raw_temp = tempfile.mkstemp(
        prefix=f".{dest.stem}_",
        suffix=f".tmp{dest.suffix}",
        dir=str(dest.parent),
    )
    os.close(handle)
    staged = Path(raw_temp)
    try:
        shutil.copy2(src, staged)
        if not staged.exists() or staged.stat().st_size != src.stat().st_size:
            raise OriginExportError(f"无法完整复制 Origin 临时文件到 {dest}")
        os.replace(staged, dest)
        try:
            src.unlink()
        except OSError:
            # Origin may still hold the temporary save open until the session closes.
            pass
        return dest
    except OSError as exc:
        raise OriginExportError(f"无法把 Origin 工程安全写到 {dest}：{exc}") from exc
    finally:
        try:
            staged.unlink()
        except FileNotFoundError:
            pass


def _save_origin_project(
    op,
    path: Path,
    pending_cleanup: list[Path] | None = None,
) -> Path:
    path = path.expanduser().resolve(strict=False)
    if path.suffix.lower() != ".opju":
        path = path.with_suffix(".opju")
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_dirs: list[Path] = []
    try:
        tmp_dirs.append(Path(tempfile.mkdtemp(prefix="spectra_opju_", dir=str(path.parent))))
    except OSError:
        pass
    tmp_dirs.append(Path(tempfile.mkdtemp(prefix="spectra_opju_")))
    last_saved = None
    errors: list[str] = []
    try:
        for tmp_dir in tmp_dirs:
            tmp_path = tmp_dir / "export.opju"
            try:
                last_saved = op.save(str(tmp_path))
            except Exception as exc:
                errors.append(str(exc))
                continue
            if tmp_path.exists() and tmp_path.is_file() and tmp_path.stat().st_size > 0:
                return commit_exported_file(tmp_path, path)
        details = f"；错误：{' | '.join(errors)}" if errors else ""
        raise OriginExportError(
            f"Origin 没有写出 .opju 文件（save 返回 {last_saved!r}）{details}。"
            "请确认 Origin Pro 已启动且未被其他对话框挡住。"
        )
    finally:
        for tmp_dir in tmp_dirs:
            shutil.rmtree(tmp_dir, ignore_errors=True)
            if tmp_dir.exists() and pending_cleanup is not None:
                pending_cleanup.append(tmp_dir)


def _cleanup_origin_temp_dirs(directories: list[Path]) -> None:
    for directory in directories:
        shutil.rmtree(directory, ignore_errors=True)


def write_origin_project(
    spectra: list[Spectrum],
    output_opju: Path,
    layout: str = "auto",
    split_temp: bool = False,
    n_groups: int | None = None,
    groups: list[tuple[str, list[Spectrum]]] | None = None,
    show_origin: bool = False,
    keep_open: bool = False,
    axis_labels: AxisLabels = AxisLabels(),
) -> Path:
    if not spectra:
        raise ValueError("没有谱线")
    resolved_groups = resolve_groups(spectra, groups=groups, split_temp=split_temp, n_groups=n_groups)
    specs = build_sheet_specs(resolved_groups, layout)
    target_path = Path(output_opju).expanduser()
    if target_path.suffix.lower() != ".opju":
        target_path = target_path.with_suffix(".opju")
    _validate_target_paths([item.path for item in spectra], [target_path])
    if origin_process_running():
        raise OriginExportError("检测到 Origin 已运行。为保护当前工程，请先关闭 Origin 再生成新工程。")
    op = _load_originpro()
    if origin_process_running():
        raise OriginExportError("Origin 在启动导出期间已运行；为保护当前工程，本次未写出。")
    closed = False
    saved_successfully = False
    pending_cleanup: list[Path] = []
    started = True
    layouts = {spec.layout for spec in specs}
    resolved_layout = next(iter(layouts)) if len(layouts) == 1 else "mixed"
    try:
        op.set_show(bool(show_origin) or bool(keep_open))
        op.new(False)
        books = list(op.pages("w"))
        if not books:
            raise OriginExportError("Origin 新建工程后没有工作簿。")
        book = books[0]
        book.lname = f"spectra_{resolved_layout}"
        for extra in books[1:]:
            extra.destroy()
        shorts = unique_origin_names([spec.name for spec in specs])
        for index, (spec, short) in enumerate(zip(specs, shorts)):
            sheet = book[0] if index == 0 else book.add_sheet(short)
            sheet.name = short
            sheet.lname = spec.name
            items = list(spec.spectra)
            _fill_origin_sheet(sheet, items, spec.layout, axis_labels)
            _plot_origin_sheet(op, sheet, items, spec.layout, f"{short}_plot", axis_labels)
        saved = _save_origin_project(op, output_opju, pending_cleanup)
        if keep_open:
            open_project = getattr(op, "open", None)
            if callable(open_project):
                try:
                    open_project(str(saved))
                except Exception:
                    pass
            op.set_show(True)
        else:
            close_origin_app(op, started=started)
            closed = True
        if not saved.exists() or saved.stat().st_size == 0:
            raise OriginExportError(f"Origin 保存后文件仍不存在或为空：{saved}")
        saved_successfully = True
        return saved
    except OriginExportError:
        raise
    except Exception as exc:
        raise OriginExportError(f"Origin 工程生成失败：{exc}") from exc
    finally:
        if not closed and (not keep_open or not saved_successfully):
            close_origin_app(op, started=started)
        _cleanup_origin_temp_dirs(pending_cleanup)


def export_origin_project(
    files: list[Path],
    output_opju: Path,
    layout: str = "auto",
    split_temp: bool = False,
    n_groups: int | None = None,
    groups: list[tuple[str, list[Spectrum]]] | None = None,
    show_origin: bool = False,
    keep_open: bool = False,
    also_xlsx: bool = False,
    axis_labels: AxisLabels = AxisLabels(),
) -> tuple[Path, Path | None, Path | None]:
    if not files and not groups:
        raise ValueError("没有谱线文件")
    if groups:
        spectra = [item for _name, items in groups for item in items]
    else:
        spectra = parse_spectra(files)
    if not spectra:
        raise ValueError("没有谱线文件")
    resolved_groups = resolve_groups(spectra, groups=groups, split_temp=split_temp, n_groups=n_groups)
    specs = build_sheet_specs(resolved_groups, layout)
    output_path = Path(output_opju).expanduser()
    if output_path.suffix.lower() != ".opju":
        output_path = output_path.with_suffix(".opju")
    target_paths = [output_path]
    tables: list[SheetTable] = []
    csv_target: Path | None = None
    csv_files: list[Path] = []
    xlsx_output = output_path.with_suffix(".xlsx")
    if also_xlsx:
        _validate_excel_dimensions(specs, len(spectra))
        tables = [build_table(spec.name, list(spec.spectra), spec.layout, axis_labels) for spec in specs]
        csv_target, csv_files = _csv_output_plan(xlsx_output, specs)
        target_paths.extend([xlsx_output, *csv_files])
        if len(specs) > 1:
            _validate_no_input_in_directory([item.path for item in spectra], csv_target)
            _csv_directory_is_managed(csv_target)
    _validate_target_paths([item.path for item in spectra], target_paths)
    opju_path = write_origin_project(
        spectra,
        output_path,
        layout=layout,
        groups=resolved_groups,
        show_origin=show_origin,
        keep_open=keep_open,
        axis_labels=axis_labels,
    )
    xlsx_path = None
    csv_path = None
    if also_xlsx:
        layouts = {spec.layout for spec in specs}
        layout_summary = next(iter(layouts)) if len(layouts) == 1 else "mixed"
        write_xlsx(
            xlsx_output,
            tables,
            spectra,
            layout_summary,
            [spec.name for spec in specs],
            axis_labels,
            [(spec.name, spec.layout) for spec in specs],
        )
        if len(specs) == 1:
            csv_path = csv_files[0]
            write_csv(csv_path, tables[0])
        else:
            csv_path = _write_csv_directory(csv_target, tables, [path.name for path in csv_files])
        xlsx_path = xlsx_output
    return opju_path, xlsx_path, csv_path


def enable_windows_file_drop(widget: tk.Misc, on_paths) -> bool:
    """Hook WM_DROPFILES so Explorer can drop files or folders onto the window."""
    if sys.platform != "win32":
        return False
    try:
        return _enable_windows_file_drop(widget, on_paths)
    except Exception:
        return False


def _enable_windows_file_drop(widget: tk.Misc, on_paths) -> bool:
    import ctypes
    from ctypes import wintypes

    WM_DROPFILES = 0x0233
    GWLP_WNDPROC = -4
    GA_ROOT = 2
    LRESULT = ctypes.c_int64 if ctypes.sizeof(ctypes.c_void_p) == 8 else ctypes.c_long
    user32 = ctypes.windll.user32
    shell32 = ctypes.windll.shell32

    user32.GetParent.argtypes = [wintypes.HWND]
    user32.GetParent.restype = wintypes.HWND
    user32.GetAncestor.argtypes = [wintypes.HWND, ctypes.c_uint]
    user32.GetAncestor.restype = wintypes.HWND
    DragAcceptFiles = shell32.DragAcceptFiles
    DragAcceptFiles.argtypes = [wintypes.HWND, wintypes.BOOL]
    DragQueryFileW = shell32.DragQueryFileW
    DragQueryFileW.argtypes = [wintypes.HANDLE, wintypes.UINT, wintypes.LPWSTR, wintypes.UINT]
    DragQueryFileW.restype = wintypes.UINT
    DragFinish = shell32.DragFinish
    DragFinish.argtypes = [wintypes.HANDLE]

    widget.update_idletasks()
    hwnd = wintypes.HWND(int(widget.winfo_id()))
    parent = user32.GetParent(hwnd)
    if parent:
        hwnd = parent
    root = user32.GetAncestor(hwnd, GA_ROOT)
    if root:
        hwnd = root

    WNDPROC = ctypes.WINFUNCTYPE(LRESULT, wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM)
    GetWindowLongPtr = user32.GetWindowLongPtrW
    SetWindowLongPtr = user32.SetWindowLongPtrW
    CallWindowProc = user32.CallWindowProcW
    GetWindowLongPtr.argtypes = [wintypes.HWND, ctypes.c_int]
    GetWindowLongPtr.restype = ctypes.c_void_p
    SetWindowLongPtr.argtypes = [wintypes.HWND, ctypes.c_int, ctypes.c_void_p]
    SetWindowLongPtr.restype = ctypes.c_void_p
    CallWindowProc.argtypes = [ctypes.c_void_p, wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
    CallWindowProc.restype = LRESULT

    old_proc = GetWindowLongPtr(hwnd, GWLP_WNDPROC)
    if not old_proc:
        return False

    def _paths_from_hdrop(hdrop) -> list[str]:
        count = DragQueryFileW(hdrop, 0xFFFFFFFF, None, 0)
        paths: list[str] = []
        for index in range(count):
            length = DragQueryFileW(hdrop, index, None, 0) + 1
            buffer = ctypes.create_unicode_buffer(length)
            DragQueryFileW(hdrop, index, buffer, length)
            paths.append(buffer.value)
        DragFinish(hdrop)
        return paths

    def _wndproc(hwnd_value, message, wparam, lparam):
        if message == WM_DROPFILES:
            try:
                paths = _paths_from_hdrop(wparam)
                widget.after(0, lambda: on_paths(paths))
            except Exception:
                pass
            return 0
        return CallWindowProc(old_proc, hwnd_value, message, wparam, lparam)

    new_proc = WNDPROC(_wndproc)
    widget._drop_wndproc = new_proc  # noqa: SLF001 — keep callback alive
    widget._drop_oldproc = old_proc  # noqa: SLF001
    widget._drop_hwnd = hwnd  # noqa: SLF001

    def _restore(_event=None):
        try:
            SetWindowLongPtr(hwnd, GWLP_WNDPROC, old_proc)
        except Exception:
            pass

    widget.bind("<Destroy>", _restore, add="+")
    SetWindowLongPtr(hwnd, GWLP_WNDPROC, ctypes.cast(new_proc, ctypes.c_void_p))
    DragAcceptFiles(hwnd, True)
    return True


def _status_text(spectra: list[Spectrum], layout: str, assignments: list[str]) -> str:
    if not spectra:
        return "把两列谱线文件或文件夹拖进窗口，或点添加。各组 X 相同使用 XYYY，不同使用 XYXY。"
    valid_assignments = assignments if len(assignments) == len(spectra) else ["all"] * len(spectra)
    grouped = groups_from_assignments(spectra, valid_assignments)
    details: list[str] = []
    for name, items in grouped:
        try:
            resolved = resolve_layout(items, layout)
            n_columns = 1 + len(items) if resolved == "XYYY" else 2 * len(items)
            point_counts = [item.n_points for item in items]
            points = str(point_counts[0]) if min(point_counts) == max(point_counts) else f"{min(point_counts)}–{max(point_counts)}"
            detail = f"{name}: {resolved}, {n_columns} 列, {len(items)} 条, {points} 点"
            if resolved == "XYYY" and not shared_x_grid(items):
                detail += "（X 网格不一致）"
        except ValueError as exc:
            detail = f"{name}: 无法确定布局（{exc}）"
        details.append(detail)
    mode = "自动布局" if (layout or "auto").strip().upper() in {"", "AUTO"} else f"布局 {layout}"
    return f"{len(spectra)} 条谱线 | {mode} | " + "；".join(details)


class SpectraToOriginApp:
    def __init__(self, initial_files: list[Path] | None = None) -> None:
        self.files: list[Path] = []
        self.assignments: list[str] = []
        self._user_grouped = False
        self._cached_key: tuple[tuple[Path, int | None, int | None], ...] | None = None
        self._cached_spectra: list[Spectrum] = []
        self._busy_widgets: list[tk.Widget] = []
        self._widget_states: dict[tk.Widget, str] = {}
        self._export_process = None
        self._export_recv = None
        self._export_result: dict | None = None
        self._export_kind: str | None = None
        self._preview_spectrum: Spectrum | None = None
        self.root = tk.Tk()
        self.root.title("谱线 → Origin 工程")
        self.root.minsize(860, 740)
        self.layout_var = tk.StringVar(value="auto")
        self.n_groups_var = tk.StringVar(value="3")
        self.keep_open_var = tk.BooleanVar(value=False)
        self.show_origin_var = tk.BooleanVar(value=False)
        self.also_xlsx_var = tk.BooleanVar(value=False)
        self.group_choice = tk.StringVar(value="")
        self.rename_var = tk.StringVar(value="")
        self.x_name_var = tk.StringVar(value=X_LONG_NAME)
        self.x_unit_var = tk.StringVar(value=X_UNITS)
        self.y_name_var = tk.StringVar(value="Intensity")
        self.y_unit_var = tk.StringVar(value=Y_UNITS)
        self._build()
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.root.after(200, self._hook_drop)
        if initial_files:
            self._add_paths(initial_files)

    def _track_busy_widget(self, widget: tk.Widget):
        self._busy_widgets.append(widget)
        self._widget_states[widget] = str(widget.cget("state")) if "state" in widget.keys() else "normal"
        return widget

    def _build(self) -> None:
        frame = ttk.Frame(self.root, padding=10)
        frame.pack(fill=tk.BOTH, expand=True)

        hint = ttk.Label(frame, text="把谱线文件或文件夹拖到这个窗口即可加载")
        hint.pack(fill=tk.X)

        list_frame = ttk.Frame(frame)
        list_frame.pack(fill=tk.BOTH, expand=True, pady=(6, 0))
        scroll = ttk.Scrollbar(list_frame)
        self.listbox = tk.Listbox(list_frame, selectmode=tk.EXTENDED, height=16)
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        hscroll = ttk.Scrollbar(list_frame, orient=tk.HORIZONTAL, command=self.listbox.xview)
        hscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.listbox.config(yscrollcommand=scroll.set, xscrollcommand=hscroll.set)
        scroll.config(command=self.listbox.yview)
        self._track_busy_widget(self.listbox)
        self.listbox.bind("<<ListboxSelect>>", self._on_list_selection)

        preview = ttk.LabelFrame(frame, text="谱线预览", padding=(6, 4))
        preview.pack(fill=tk.X, pady=(8, 0))
        self.preview_text = ttk.Label(preview, text="选择一条谱线查看预览和完整路径", wraplength=820, justify=tk.LEFT)
        self.preview_text.pack(fill=tk.X)
        self.preview_canvas = tk.Canvas(preview, height=155, background="white", highlightthickness=1, highlightbackground="#b8b8b8")
        self.preview_canvas.pack(fill=tk.X, pady=(4, 0))
        self.preview_canvas.bind("<Configure>", lambda _event: self._draw_preview())
        self.preview_canvas.create_text(10, 12, anchor=tk.NW, text="选择一条谱线查看预览", fill="#666666")

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X, pady=(8, 0))
        for text, command in (
            ("添加文件", self.add_files),
            ("添加文件夹", self.add_folder),
            ("删除选中", self.remove_selected),
            ("上移", lambda: self.move_selected(-1)),
            ("下移", lambda: self.move_selected(1)),
            ("清空", self.clear_files),
        ):
            self._track_busy_widget(ttk.Button(buttons, text=text, command=command)).pack(side=tk.LEFT, padx=(0, 6))

        options = ttk.Frame(frame)
        options.pack(fill=tk.X, pady=(10, 0))
        ttk.Label(options, text="布局").pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Radiobutton(
            options, text="自动", value="auto", variable=self.layout_var, command=self._refresh_status
        )).pack(side=tk.LEFT, padx=(8, 0))
        self._track_busy_widget(ttk.Radiobutton(
            options, text="XYYY（共用 X）", value="XYYY", variable=self.layout_var, command=self._refresh_status
        )).pack(side=tk.LEFT, padx=(8, 0))
        self._track_busy_widget(ttk.Radiobutton(
            options, text="XYXY（每条自己的 X）", value="XYXY", variable=self.layout_var, command=self._refresh_status
        )).pack(side=tk.LEFT, padx=(8, 0))

        groups = ttk.Frame(frame)
        groups.pack(fill=tk.X, pady=(8, 0))
        self._track_busy_widget(ttk.Button(groups, text="按文件名分组", command=self.apply_filename_groups)).pack(side=tk.LEFT)
        ttk.Label(groups, text="均分成").pack(side=tk.LEFT, padx=(12, 4))
        self._track_busy_widget(ttk.Spinbox(groups, from_=1, to=12, width=4, textvariable=self.n_groups_var)).pack(side=tk.LEFT)
        ttk.Label(groups, text="张").pack(side=tk.LEFT, padx=(4, 6))
        self._track_busy_widget(ttk.Button(groups, text="应用均分", command=self.apply_even_split)).pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Button(groups, text="全部一张表", command=self.apply_single_sheet)).pack(side=tk.LEFT, padx=(8, 0))

        move = ttk.Frame(frame)
        move.pack(fill=tk.X, pady=(6, 0))
        ttk.Label(move, text="选中移入").pack(side=tk.LEFT)
        self.group_combo = ttk.Combobox(move, textvariable=self.group_choice, width=14, state="readonly")
        self.group_combo.pack(side=tk.LEFT, padx=(6, 6))
        self._track_busy_widget(self.group_combo)
        self._track_busy_widget(ttk.Button(move, text="移入", command=self.move_selected_to_group)).pack(side=tk.LEFT)
        ttk.Label(move, text="组改名").pack(side=tk.LEFT, padx=(12, 4))
        self._track_busy_widget(ttk.Entry(move, textvariable=self.rename_var, width=16)).pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Button(move, text="改名", command=self.rename_selected_group)).pack(side=tk.LEFT, padx=(6, 0))

        origin_opts = ttk.Frame(frame)
        origin_opts.pack(fill=tk.X, pady=(8, 0))
        self._track_busy_widget(ttk.Checkbutton(origin_opts, text="完成后保持 Origin 打开", variable=self.keep_open_var)).pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Checkbutton(origin_opts, text="显示 Origin 窗口", variable=self.show_origin_var)).pack(side=tk.LEFT, padx=(16, 0))
        self._track_busy_widget(ttk.Checkbutton(origin_opts, text="同时导出 xlsx/csv", variable=self.also_xlsx_var)).pack(side=tk.LEFT, padx=(16, 0))

        axes = ttk.Frame(frame)
        axes.pack(fill=tk.X, pady=(7, 0))
        ttk.Label(axes, text="X 名称").pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Entry(axes, textvariable=self.x_name_var, width=13)).pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(axes, text="X 单位").pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Entry(axes, textvariable=self.x_unit_var, width=9)).pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(axes, text="Y 名称").pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Entry(axes, textvariable=self.y_name_var, width=13)).pack(side=tk.LEFT, padx=(4, 10))
        ttk.Label(axes, text="Y 单位").pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Entry(axes, textvariable=self.y_unit_var, width=9)).pack(side=tk.LEFT, padx=(4, 0))

        export_row = ttk.Frame(frame)
        export_row.pack(fill=tk.X, pady=(10, 0))
        self._track_busy_widget(ttk.Button(export_row, text="生成 Origin 工程 (.opju)", command=self.export)).pack(side=tk.LEFT)
        self._track_busy_widget(ttk.Button(export_row, text="只导出 Excel/CSV", command=self.export_xlsx)).pack(side=tk.LEFT, padx=(8, 0))

        self.status = ttk.Label(frame, text="", wraplength=780, justify=tk.LEFT)
        self.status.pack(fill=tk.X, pady=(10, 0))
        self._refresh_status()

    def _hook_drop(self) -> None:
        enable_windows_file_drop(self.root, self._on_drop_paths)

    def _on_drop_paths(self, paths: list[str]) -> None:
        if self._export_process is not None:
            self.status.config(text="后台导出中；请等待完成后再添加文件。")
            return
        self._add_paths([Path(item) for item in paths])

    @staticmethod
    def _list_labels(files: list[Path], assignments: list[str]) -> list[str]:
        basename_counts = Counter(path.name.casefold() for path in files)
        labels = []
        for index, path in enumerate(files):
            group = assignments[index] if index < len(assignments) else "all"
            suffix = f" — {path.parent}" if basename_counts[path.name.casefold()] > 1 else ""
            labels.append(f"{group}  |  {path.name}{suffix}")
        return labels

    def _refresh_list(self) -> None:
        self.listbox.delete(0, tk.END)
        for label in self._list_labels(self.files, self.assignments):
            self.listbox.insert(tk.END, label)
        names = list(dict.fromkeys(self.assignments))
        self.group_combo["values"] = names
        if names and self.group_choice.get() not in names:
            self.group_choice.set(names[0])
        self._refresh_status()
        self._refresh_preview()

    @staticmethod
    def _file_signature(path: Path) -> tuple[Path, int | None, int | None]:
        try:
            resolved = path.resolve()
            stat = resolved.stat()
            return resolved, stat.st_mtime_ns, stat.st_size
        except OSError:
            return path.absolute(), None, None

    def _parsed_spectra(self, *, force: bool = False) -> list[Spectrum]:
        key = tuple(self._file_signature(path) for path in self.files)
        if force or self._cached_key != key:
            parsed = parse_spectra(self.files) if self.files else []
            self._cached_spectra = parsed
            self._cached_key = key
        return self._cached_spectra

    def _refresh_status(self) -> None:
        try:
            spectra = self._parsed_spectra()
        except (OSError, ValueError) as exc:
            self.status.config(text=f"读取失败：{exc}")
            return
        self.status.config(text=_status_text(spectra, self.layout_var.get(), self.assignments))

    def _set_assignments(self, assignments: list[str], *, user: bool) -> None:
        self.assignments = assignments
        if user:
            self._user_grouped = True
        self._refresh_list()

    def _auto_assign(self) -> None:
        if not self.files:
            self.assignments = []
            return
        grouped = suggest_filename_groups([path.stem for path in self.files])
        self.assignments = assignments_from_groups(len(self.files), grouped)

    def _add_paths(self, paths: list[Path]) -> None:
        if self._export_process is not None:
            return
        try:
            incoming, errors = collect_from_user_paths(paths)
        except OSError as exc:
            incoming, errors = [], [str(exc)]
        self.files = merge_file_list(self.files, incoming)
        if not self._user_grouped:
            self._auto_assign()
        else:
            while len(self.assignments) < len(self.files):
                self.assignments.append(self.assignments[-1] if self.assignments else "all")
            self.assignments = self.assignments[: len(self.files)]
        self._refresh_list()
        if errors:
            messagebox.showwarning("部分路径未加入", "\n".join(errors[:8]))

    def _on_list_selection(self, _event=None) -> None:
        self._refresh_preview()

    def _refresh_preview(self) -> None:
        self._preview_spectrum = None
        selected = self.listbox.curselection()
        if not selected:
            self.preview_text.config(text="选择一条谱线查看预览和完整路径")
            self._draw_preview()
            return
        index = selected[0]
        try:
            spectra = self._parsed_spectra()
            spectrum = spectra[index]
        except (IndexError, OSError, ValueError) as exc:
            self.preview_text.config(text=f"无法读取选中谱线：{exc}")
            self._draw_preview()
            return
        self._preview_spectrum = spectrum
        count = spectrum.n_points
        if count > 1600:
            preview_count = min(1600, count)
            detail = f"{count:,} 点；图中均匀抽样显示 {preview_count:,} 点"
        else:
            detail = f"{count:,} 点；完整显示"
        self.preview_text.config(text=f"{spectrum.path}\n{detail}")
        self._draw_preview()

    def _draw_preview(self) -> None:
        canvas = getattr(self, "preview_canvas", None)
        if canvas is None:
            return
        canvas.delete("all")
        width = max(100, canvas.winfo_width())
        height = max(80, canvas.winfo_height())
        spectrum = self._preview_spectrum
        if spectrum is None:
            canvas.create_text(12, 12, anchor=tk.NW, text="选择一条谱线查看预览", fill="#666666")
            return
        pairs = []
        for x_text, y_text in zip(spectrum.x_text, spectrum.y_text):
            x_value, y_value = float(x_text), float(y_text)
            if math.isfinite(x_value) and math.isfinite(y_value):
                pairs.append((x_value, y_value))
        if not pairs:
            canvas.create_text(12, 12, anchor=tk.NW, text="没有可绘制的有限数值", fill="#666666")
            return
        left, top, right, bottom = 48, 10, width - 12, height - 26
        if right <= left or bottom <= top:
            return
        x_min = min(x for x, _y in pairs)
        x_max = max(x for x, _y in pairs)
        y_min = min(y for _x, y in pairs)
        y_max = max(y for _x, y in pairs)
        if x_min == x_max:
            x_min, x_max = x_min - 0.5, x_max + 0.5
        if y_min == y_max:
            margin = abs(y_min) * 0.05 or 0.5
            y_min, y_max = y_min - margin, y_max + margin
        canvas.create_line(left, top, left, bottom, fill="#555555")
        canvas.create_line(left, bottom, right, bottom, fill="#555555")
        canvas.create_text(left - 6, top, anchor=tk.NE, text=f"{y_max:.4g}", fill="#555555")
        canvas.create_text(left - 6, bottom, anchor=tk.SE, text=f"{y_min:.4g}", fill="#555555")
        canvas.create_text(left, bottom + 4, anchor=tk.NW, text=f"{x_min:.4g}", fill="#555555")
        canvas.create_text(right, bottom + 4, anchor=tk.NE, text=f"{x_max:.4g}", fill="#555555")
        preview_count = min(1600, len(pairs))
        sampled = (
            [pairs[index * (len(pairs) - 1) // (preview_count - 1)] for index in range(preview_count)]
            if preview_count > 1 else pairs
        )
        coords = []
        for x_value, y_value in sampled:
            px = left + (x_value - x_min) / (x_max - x_min) * (right - left)
            py = bottom - (y_value - y_min) / (y_max - y_min) * (bottom - top)
            coords.extend((px, py))
        if len(coords) >= 4:
            canvas.create_line(*coords, fill="#1769aa", width=1.5)
        elif coords:
            canvas.create_oval(coords[0] - 2, coords[1] - 2, coords[0] + 2, coords[1] + 2, fill="#1769aa", outline="")

    def add_files(self) -> None:
        selected = filedialog.askopenfilenames(
            title="选择谱线文件",
            filetypes=[("谱线数据", "*.txt *.dat *.xy *.csv *.tsv"), ("所有文件", "*.*")],
        )
        if selected:
            self._add_paths([Path(item) for item in selected])

    def add_folder(self) -> None:
        selected = filedialog.askdirectory(title="选择谱线文件夹")
        if selected:
            self._add_paths([Path(selected)])

    def remove_selected(self) -> None:
        indexes = sorted(self.listbox.curselection(), reverse=True)
        for index in indexes:
            del self.files[index]
            del self.assignments[index]
        self._refresh_list()

    def move_selected(self, step: int) -> None:
        indexes = list(self.listbox.curselection())
        if len(indexes) != 1:
            return
        index = indexes[0]
        target = index + step
        if target < 0 or target >= len(self.files):
            return
        self.files[index], self.files[target] = self.files[target], self.files[index]
        self.assignments[index], self.assignments[target] = self.assignments[target], self.assignments[index]
        self._refresh_list()
        self.listbox.selection_set(target)

    def clear_files(self) -> None:
        self.files = []
        self.assignments = []
        self._user_grouped = False
        self._refresh_list()

    def apply_filename_groups(self) -> None:
        self._user_grouped = False
        self._auto_assign()
        self._refresh_list()

    def apply_even_split(self) -> None:
        try:
            n_groups = int(self.n_groups_var.get())
            grouped = partition_even(len(self.files), n_groups)
        except ValueError as exc:
            messagebox.showinfo("分组数无效", str(exc) or "均分数必须是正整数。")
            return
        if not self.files:
            return
        self._set_assignments(assignments_from_groups(len(self.files), grouped), user=True)

    def apply_single_sheet(self) -> None:
        if not self.files:
            return
        self._set_assignments(["all"] * len(self.files), user=True)

    def move_selected_to_group(self) -> None:
        name = self.group_choice.get().strip()
        if not name:
            messagebox.showinfo("没有目标组", "先按文件名分组或均分，再把选中项移入某组。")
            return
        indexes = list(self.listbox.curselection())
        if not indexes:
            return
        for index in indexes:
            self.assignments[index] = name
        self._set_assignments(self.assignments, user=True)
        for index in indexes:
            self.listbox.selection_set(index)

    def rename_selected_group(self) -> None:
        new_name = self.rename_var.get().strip()
        if not new_name:
            messagebox.showinfo("组名是空的", "输入新的 sheet 名再改名。")
            return
        indexes = list(self.listbox.curselection())
        if indexes:
            old_names = {self.assignments[index] for index in indexes}
        else:
            current = self.group_choice.get().strip()
            old_names = {current} if current else set()
        if not old_names:
            return
        self.assignments = [new_name if name in old_names else name for name in self.assignments]
        self._set_assignments(self.assignments, user=True)

    def _export_groups(self) -> list[tuple[str, list[Spectrum]]]:
        spectra = self._parsed_spectra(force=True)
        grouped_paths = groups_from_assignments(self.files, self.assignments)
        by_path = {item.path.resolve(): item for item in spectra}
        return [
            (name, [by_path[path.resolve()] for path in paths])
            for name, paths in grouped_paths
        ]

    def export(self) -> None:
        if not self.files:
            messagebox.showinfo("没有谱线", "先拖入或添加两列谱线文件。")
            return
        if self._export_process is not None:
            return
        if origin_process_running():
            messagebox.showwarning(
                "请先关闭 Origin",
                "检测到 Origin Pro 正在运行。请先保存当前工程并关闭 Origin，再生成新工程。",
            )
            return
        layout = self.layout_var.get()
        default_dir = self.files[0].parent
        try:
            spectra = self._parsed_spectra(force=True)
            resolved = resolve_layout(spectra, layout)
            default_name = f"origin_{resolved}.opju"
        except (OSError, ValueError) as exc:
            if isinstance(exc, OSError):
                messagebox.showerror("读取谱线失败", str(exc))
                return
            default_name = "origin_spectra.opju"
        output = filedialog.asksaveasfilename(
            title="保存 Origin 工程",
            defaultextension=".opju",
            initialdir=str(default_dir),
            initialfile=default_name,
            filetypes=[("Origin 工程", "*.opju")],
        )
        if not output:
            return
        self._start_export_job("opju", Path(output))

    def export_xlsx(self) -> None:
        if not self.files:
            messagebox.showinfo("没有谱线", "先拖入或添加两列谱线文件。")
            return
        if self._export_process is not None:
            return
        try:
            self._parsed_spectra(force=True)
        except (OSError, ValueError) as exc:
            messagebox.showerror("读取谱线失败", str(exc))
            return
        output = filedialog.asksaveasfilename(
            title="导出 Excel/CSV",
            defaultextension=".xlsx",
            initialdir=str(self.files[0].parent),
            initialfile="spectra.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if output:
            self._start_export_job("xlsx", Path(output))

    def _axis_label_values(self) -> tuple[str, str, str, str]:
        return (
            self.x_name_var.get().strip(),
            self.x_unit_var.get().strip(),
            self.y_name_var.get().strip(),
            self.y_unit_var.get().strip(),
        )

    def _start_export_job(self, kind: str, output: Path) -> None:
        import multiprocessing

        try:
            groups = self._export_groups()
        except (OSError, ValueError, OriginExportError) as exc:
            self._refresh_status()
            messagebox.showerror("读取谱线失败", str(exc))
            return
        if kind == "opju" and origin_process_running():
            messagebox.showwarning(
                "请先关闭 Origin",
                "Origin Pro 已在后台导出前启动。请保存并关闭 Origin，再重新导出。",
            )
            return
        request = {
            "kind": kind,
            "files": list(self.files),
            "groups": groups,
            "output": str(output),
            "layout": self.layout_var.get(),
            "axis_labels": self._axis_label_values(),
            "show_origin": self.show_origin_var.get(),
            "keep_open": self.keep_open_var.get(),
            "also_xlsx": self.also_xlsx_var.get() if kind == "opju" else False,
        }
        from gui_export import export_worker

        context = multiprocessing.get_context("spawn")
        recv_conn, send_conn = context.Pipe(duplex=False)
        process = context.Process(target=export_worker, args=(send_conn, request))
        try:
            process.start()
        except (OSError, RuntimeError) as exc:
            recv_conn.close()
            send_conn.close()
            messagebox.showerror("无法启动后台导出", str(exc))
            return
        send_conn.close()
        self._export_process = process
        self._export_recv = recv_conn
        self._export_result = None
        self._export_kind = kind
        self._set_busy(True)
        label = "Origin 工程" if kind == "opju" else "Excel/CSV"
        self.status.config(text=f"正在后台生成{label}；窗口保持响应，完成前请勿关闭。")
        self.root.after(100, self._poll_export)

    def _set_busy(self, busy: bool) -> None:
        for widget in self._busy_widgets:
            try:
                widget.configure(state="disabled" if busy else self._widget_states[widget])
            except tk.TclError:
                continue

    def _poll_export(self) -> None:
        process = self._export_process
        recv_conn = self._export_recv
        if process is None or recv_conn is None:
            return
        try:
            while recv_conn.poll():
                message = recv_conn.recv()
                if message.get("type") == "progress":
                    self.status.config(text=f"{message.get('text', '后台正在导出…')} 请等待完成。")
                elif message.get("type") == "result":
                    self._export_result = message.get("result")
        except (EOFError, OSError):
            pass
        if process.is_alive():
            self.root.after(150, self._poll_export)
            return
        process.join()
        exit_code = process.exitcode
        try:
            recv_conn.close()
        except OSError:
            pass
        try:
            process.close()
        except (OSError, ValueError):
            pass
        self._export_process = None
        self._export_recv = None
        self._set_busy(False)
        self._refresh_status()
        self._refresh_preview()
        result = self._export_result
        self._export_result = None
        self._export_kind = None
        if result is None:
            result = {"ok": False, "error": f"后台导出进程异常结束（退出码 {exit_code}）"}
        if not result.get("ok"):
            self.status.config(text="导出失败；请查看错误信息。")
            messagebox.showerror("导出失败", result.get("error", "后台导出失败"))
            return
        paths = result["paths"]
        lines = []
        for key, title in (("opju", "OPJU"), ("xlsx", "XLSX"), ("csv", "CSV")):
            if paths.get(key):
                lines.append(f"{title}：{paths[key]}")
        self.status.config(text="导出完成：" + "；".join(lines))
        messagebox.showinfo("导出完成", "\n".join(lines))

    def _on_close(self) -> None:
        process = self._export_process
        if process is not None and process.is_alive():
            messagebox.showinfo(
                "导出仍在进行",
                "后台导出正在使用 Origin 或写入文件。请等待它完成后再关闭窗口。",
            )
            return
        if process is not None:
            self._poll_export()
            return
        self.root.destroy()

    def run(self) -> None:
        self.root.mainloop()


def _enable_windows_dpi() -> None:
    if sys.platform != "win32":
        return
    try:
        from ctypes import windll

        windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        return


def _positive_int(value: str) -> int:
    try:
        number = int(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("分组数必须是正整数") from exc
    if number < 1:
        raise argparse.ArgumentTypeError("分组数必须至少为 1")
    return number


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="批量整理两列谱线，导出 Origin 工程或 Excel/CSV")
    parser.add_argument("--version", action="version", version=f"SpectraToOrigin {__version__}")
    parser.add_argument("--cli", action="store_true", help="不打开窗口，直接导出")
    parser.add_argument("-i", "--input", nargs="+", help="txt/dat/xy/csv/tsv 文件或文件夹")
    parser.add_argument("-o", "--output", help="输出 .opju 或 .xlsx 路径")
    parser.add_argument("--layout", choices=["auto", "XYYY", "XYXY"], default="auto")
    grouping = parser.add_mutually_exclusive_group()
    grouping.add_argument("--split-temp", "--group-by-name", action="store_true", help="按文件名自动分组（不限于温度）")
    grouping.add_argument("--n-groups", type=_positive_int, default=None, help="按列表顺序均分成 N 张工作表")
    parser.add_argument("--also-xlsx", action="store_true", help="同时写出 xlsx/csv 备份")
    parser.add_argument("--show-origin", action="store_true", help="生成时显示 Origin 窗口")
    parser.add_argument("--keep-open", action="store_true", help="保存后不关闭 Origin")
    parser.add_argument("--xlsx-only", action="store_true", help="只写 xlsx/csv，不启动 Origin")
    parser.add_argument("--check", action="store_true", help="只检查输入并列出每组布局，不导出、不启动 Origin")
    parser.add_argument("--x-name", default=X_LONG_NAME, help="X 轴名称，默认 2theta")
    parser.add_argument("--x-unit", default=X_UNITS, help="X 轴单位，默认 deg；可传空字符串")
    parser.add_argument("--y-name", default="Intensity", help="Y 轴名称，默认 Intensity")
    parser.add_argument("--y-unit", default=Y_UNITS, help="Y 轴单位，默认 a.u.；可传空字符串")
    return parser


def run_cli(args: argparse.Namespace) -> int:
    if not args.input or (not args.check and not args.output):
        print("需要 --input；导出时还需要 --output", file=sys.stderr)
        return 2
    try:
        incoming, errors = collect_from_user_paths([Path(raw) for raw in args.input])
        if errors:
            print("输入路径错误：\n" + "\n".join(errors), file=sys.stderr)
            return 2
        files = merge_file_list([], incoming)
        if not files:
            print("没有找到可读取的两列谱线文件", file=sys.stderr)
            return 2
        spectra = parse_spectra(files)
        groups = resolve_groups(spectra, split_temp=args.split_temp, n_groups=args.n_groups)
        specs = build_sheet_specs(groups, args.layout)
        if args.check:
            print(f"检查通过：{len(spectra)} 条谱线，{len(specs)} 组")
            for spec in specs:
                n_columns = 1 + len(spec.spectra) if spec.layout == "XYYY" else 2 * len(spec.spectra)
                points = [item.n_points for item in spec.spectra]
                print(f"{spec.name}: {len(spec.spectra)} 条 | {spec.layout} | {n_columns} 列 | {min(points)}–{max(points)} 点")
            return 0
        output = Path(args.output)
        labels = AxisLabels(args.x_name, args.x_unit, args.y_name, args.y_unit)
        if args.xlsx_only:
            xlsx_path, csv_path = export_spectra(
                files,
                output.with_suffix(".xlsx"),
                layout=args.layout,
                groups=groups,
                write_csv_copy=True,
                axis_labels=labels,
            )
            print(xlsx_path)
            if csv_path:
                print(csv_path)
            return 0
        opju_path, xlsx_path, csv_path = export_origin_project(
            files,
            output,
            layout=args.layout,
            groups=groups,
            show_origin=args.show_origin,
            keep_open=args.keep_open,
            also_xlsx=args.also_xlsx,
            axis_labels=labels,
        )
    except (ValueError, OSError, OriginExportError) as exc:
        operation = "输入检查失败" if args.check else "表格导出失败" if args.xlsx_only else "Origin 工程生成失败"
        print(f"{operation}：{exc}", file=sys.stderr)
        return 4
    print(opju_path)
    if xlsx_path:
        print(xlsx_path)
    if csv_path:
        print(csv_path)
    return 0


def _attach_console_if_cli(argv: list[str] | None) -> None:
    args = list(argv if argv is not None else sys.argv[1:])
    wants_console = any(flag in args for flag in ("--cli", "--check", "--version", "-h", "--help"))
    if not wants_console or not getattr(sys, "frozen", False) or sys.platform != "win32":
        return
    try:
        from ctypes import windll

        attached = windll.kernel32.AttachConsole(-1)
        if attached == 0:
            windll.kernel32.AllocConsole()
        sys.stdin = open("CONIN$", "r", encoding="utf-8", errors="replace")
        sys.stdout = open("CONOUT$", "w", encoding="utf-8", errors="replace")
        sys.stderr = open("CONOUT$", "w", encoding="utf-8", errors="replace")
    except Exception:
        return


def main(argv: list[str] | None = None) -> int:
    _attach_console_if_cli(argv)
    args = build_parser().parse_args(argv)
    if args.cli or args.check:
        return run_cli(args)
    _enable_windows_dpi()
    initial: list[Path] = []
    if args.input:
        initial = load_from_drop_payload([], args.input)
    SpectraToOriginApp(initial).run()
    return 0


if __name__ == "__main__":
    import multiprocessing

    multiprocessing.freeze_support()
    raise SystemExit(main())
