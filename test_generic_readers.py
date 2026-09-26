"""Contract tests for the loss-aware generic data readers."""

from __future__ import annotations

import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from origin_bridge.models import DataImportError
from origin_bridge.readers import discover_files, read_tables


class GenericReaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def write(self, name: str, content: str | bytes) -> Path:
        path = (self.root / name).resolve()
        path.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, bytes):
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8")
        return path

    def test_discover_files_recurses_naturally_and_deduplicates(self) -> None:
        first = self.write("nested/run10.csv", "x,y\n1,2\n")
        second = self.write("nested/run2.TSV", "x\ty\n1\t2\n")
        found = discover_files([self.root, first])
        self.assertEqual(found, [second, first])
        self.assertTrue(all(path.is_absolute() for path in found))

    def test_discover_files_rejects_empty_missing_unknown_and_empty_directory(self) -> None:
        with self.assertRaises(DataImportError):
            discover_files([])
        with self.assertRaises(DataImportError):
            discover_files([self.root / "missing.csv"])
        unsupported = self.write("notes.md", "hello")
        with self.assertRaises(DataImportError):
            discover_files([unsupported])
        empty = self.root / "empty"
        empty.mkdir()
        with self.assertRaises(DataImportError):
            discover_files([empty])

    def test_csv_auto_header_units_duplicate_names_missing_and_d_exponent(self) -> None:
        path = self.write(
            "measurements.csv",
            "Time (s),Signal[mV],Signal[mV],Comment\n"
            "0.0,1.25D+2,,first\n"
            "0.5,NA,3.0,second\n",
        )
        table = read_tables(path)[0]
        self.assertEqual([column.name for column in table.columns], ["Time", "Signal", "Signal_2", "Comment"])
        self.assertEqual([column.unit for column in table.columns], ["s", "mV", "mV", ""])
        self.assertEqual(table.columns[0].kind, "number")
        self.assertEqual(table.columns[1].values, ("1.25E+2", "NA"))
        self.assertEqual(table.columns[2].values, (None, "3.0"))
        self.assertEqual(table.columns[3].kind, "text")
        self.assertTrue(any("数值和文本" in warning for warning in table.warnings))
        self.assertEqual(table.n_rows, 2)
        self.assertEqual(table.read_options["header"], True)
        self.assertEqual(table.read_options["delimiter"], ",")
        self.assertEqual(table.read_options["encoding"], "utf-8-sig")
        self.assertEqual(table.source, path.resolve())
        self.assertEqual(len(table.source_hash), 64)

    def test_auto_header_keeps_all_string_first_row_and_warns(self) -> None:
        path = self.write("labels.tsv", "red\tblue\ngreen\tyellow\n")
        table = read_tables(path)[0]
        self.assertEqual(table.n_rows, 2)
        self.assertEqual(table.columns[0].values, ("red", "green"))
        self.assertTrue(table.warnings)
        explicit = read_tables(path, header=True)[0]
        self.assertEqual(explicit.columns[0].name, "red")
        self.assertEqual(explicit.n_rows, 1)

    def test_auto_header_recognizes_text_names_followed_by_numeric_records(self) -> None:
        path = self.write("series.txt", "Elapsed (s),Signal (mV)\n0,2.5\n1,3.5\n")
        table = read_tables(path)[0]
        self.assertEqual([column.name for column in table.columns], ["Elapsed", "Signal"])
        self.assertEqual([column.kind for column in table.columns], ["number", "number"])
        self.assertEqual(table.n_rows, 2)

        one_column = self.write("single.csv", "Temperature[K]\n300\n301\n")
        single_table = read_tables(one_column)[0]
        self.assertEqual(single_table.columns[0].name, "Temperature")
        self.assertEqual(single_table.n_rows, 2)

    def test_leading_zero_identifiers_stay_text_and_values_are_not_deleted(self) -> None:
        path = self.write("ids.csv", "ID,value\n00123,1\n00124,\n00125,3\n")
        table = read_tables(path)[0]
        self.assertEqual(table.columns[0].kind, "text")
        self.assertEqual(table.columns[0].values, ("00123", "00124", "00125"))
        self.assertEqual(table.columns[1].values, ("1", None, "3"))
        self.assertEqual(table.n_rows, 3)

    def test_skip_rows_and_explicit_delimiter_are_recorded_for_replay(self) -> None:
        path = self.write("semicolon.dat", "metadata\nX;Y\n1;2\n3;4\n")
        table = read_tables(path, skip_rows=1, delimiter=";")[0]
        self.assertEqual(table.read_options["skip_rows"], 1)
        self.assertEqual(table.read_options["delimiter"], ";")
        replay = read_tables(path, **table.read_options)[0]
        self.assertEqual(replay.columns, table.columns)

    def test_gbk_is_detected_and_kept_as_actual_option(self) -> None:
        path = self.write("chinese.dat", "温度,K\n室温,1.25\n".encode("gbk"))
        table = read_tables(path, header=True)[0]
        self.assertEqual(table.read_options["encoding"], "gb18030")
        self.assertEqual(table.columns[0].values, ("室温",))
        self.assertEqual(table.columns[1].kind, "number")

    def test_nonfinite_number_has_physical_line_and_column(self) -> None:
        path = self.write("bad.csv", "skip\nx,y\n0,1\n1,NaN\n")
        with self.assertRaisesRegex(DataImportError, r"第 4 行第 2 列.*NaN"):
            read_tables(path, skip_rows=1)

    def test_ragged_delimited_rows_are_rejected(self) -> None:
        path = self.write("ragged.csv", "a,b\n1,2\n3\n")
        with self.assertRaisesRegex(DataImportError, "矩形表"):
            read_tables(path)

    def test_datetime_is_inferred_without_treating_numeric_date_as_datetime(self) -> None:
        path = self.write("dates.csv", "timestamp,count\n2025-03-04T05:06:07,2\n2025-03-05T05:06:07,3\n")
        table = read_tables(path)[0]
        self.assertEqual(table.columns[0].kind, "datetime")
        self.assertEqual(table.columns[0].values[0], "2025-03-04T05:06:07")
        self.assertEqual(table.columns[1].kind, "number")

        single_date = self.write(
            "timestamps.csv",
            "timestamp\n2025-03-04T05:06:07\n2025-03-05T05:06:07\n",
        )
        date_table = read_tables(single_date)[0]
        self.assertEqual(date_table.columns[0].name, "timestamp")
        self.assertEqual(date_table.columns[0].kind, "datetime")
        self.assertEqual(date_table.n_rows, 2)

    def test_json_record_array_and_column_units(self) -> None:
        path = self.write(
            "records.json",
            json.dumps([{"time (s)": "2025-01-01", "signal[mV]": 1.5}, {"time (s)": "2025-01-02", "signal[mV]": 2.5}]),
        )
        table = read_tables(path)[0]
        self.assertEqual([column.unit for column in table.columns], ["s", "mV"])
        self.assertEqual(table.columns[0].kind, "datetime")
        self.assertEqual(table.columns[1].values, ("1.5", "2.5"))
        self.assertTrue(table.read_options["header"])

    def test_json_column_arrays_and_matrix_form(self) -> None:
        columns_path = self.write("columns.json", '{"x":[1,2],"y":[3,4]}')
        columns_table = read_tables(columns_path)[0]
        self.assertEqual(columns_table.n_rows, 2)
        self.assertEqual(columns_table.columns[1].values, ("3", "4"))

        matrix_path = self.write(
            "matrix.json",
            '{"columns":["x","temperature (K)"],"data":[[0,300],[1,301]]}',
        )
        matrix_table = read_tables(matrix_path)[0]
        self.assertEqual(matrix_table.columns[1].unit, "K")
        self.assertEqual(matrix_table.columns[0].kind, "number")

    def test_json_decimal_precision_is_preserved(self) -> None:
        exact = "0.123456789012345678901234567890123456789"
        path = self.write("precision.json", '{"x":[' + exact + "," + exact + "]}")
        table = read_tables(path)[0]
        self.assertEqual(table.columns[0].kind, "number")
        self.assertEqual(table.columns[0].values, (exact, exact))

        jsonl_path = self.write("precision.jsonl", '{"x":' + exact + "}\n")
        jsonl_table = read_tables(jsonl_path)[0]
        self.assertEqual(jsonl_table.columns[0].values, (exact,))

    def test_jsonl_records_preserve_rows_and_reject_nested_or_ragged_data(self) -> None:
        path = self.write("events.jsonl", '{"t":1,"label":"a"}\n{"t":2,"label":"b"}\n')
        table = read_tables(path)[0]
        self.assertEqual(table.columns[0].values, ("1", "2"))
        self.assertEqual(table.n_rows, 2)

        nested = self.write("nested.json", '{"x":[{"a":1}]}')
        with self.assertRaisesRegex(DataImportError, "嵌套 JSON"):
            read_tables(nested)
        ragged = self.write("ragged.json", '[{"x":1,"y":2},{"x":3}]')
        with self.assertRaisesRegex(DataImportError, "字段不一致"):
            read_tables(ragged)

        nonfinite = self.write("nonfinite.json", '{"x":[1,1e999]}')
        with self.assertRaisesRegex(DataImportError, r"第 2 行第 1 列.*非有限"):
            read_tables(nonfinite)

    def test_json_and_jsonl_duplicate_keys_are_rejected_with_location(self) -> None:
        duplicate_json = self.write("duplicate.json", '[{"x":1,"x":2}]')
        with self.assertRaisesRegex(DataImportError, r"duplicate\.json 第 1 条记录.*重复 JSON 键 'x'"):
            read_tables(duplicate_json)

        duplicate_jsonl = self.write(
            "duplicate.jsonl",
            '{"x":1}\n{"x":2,"x":3}\n',
        )
        with self.assertRaisesRegex(DataImportError, r"duplicate\.jsonl 第 2 行.*重复 JSON 键 'x'"):
            read_tables(duplicate_jsonl)

    def test_excel_reads_nonempty_sheets_and_iso_dates(self) -> None:
        import openpyxl

        path = self.root / "book.xlsx"
        workbook = openpyxl.Workbook()
        first = workbook.active
        first.title = "Run 1"
        first.append(["Time", "Temperature (K)"])
        first.append([datetime(2025, 1, 2, 3, 4), 300.0])
        workbook.create_sheet("Empty")
        second = workbook.create_sheet("Run 2")
        second.append(["x", "y"])
        second.append([0, 1])
        workbook.save(path)

        tables = read_tables(path)
        self.assertEqual([table.source_sheet for table in tables], ["Run 1", "Run 2"])
        self.assertEqual(tables[0].columns[0].kind, "datetime")
        self.assertEqual(tables[0].columns[0].values, ("2025-01-02T03:04:00",))
        self.assertEqual(tables[0].columns[1].unit, "K")
        self.assertEqual(tables[0].read_options["sheet"], "Run 1")
        self.assertEqual(read_tables(path, sheet="Run 2")[0].source_sheet, "Run 2")
        with self.assertRaisesRegex(DataImportError, "不存在工作表"):
            read_tables(path, sheet="Missing")

    def test_excel_formula_policy_requires_cache_or_preserves_formula_text(self) -> None:
        import openpyxl

        path = self.root / "formula.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["x", "y"])
        sheet.append([1, "=1+2"])
        workbook.save(path)
        with self.assertRaisesRegex(DataImportError, "没有缓存结果"):
            read_tables(path)
        table = read_tables(path, header=True, formula_policy="text")[0]
        self.assertEqual(table.columns[1].values, ("=1+2",))
        self.assertEqual(table.read_options["formula_policy"], "text")

    def test_excel_formula_view_is_streamed_once_for_large_sheets(self) -> None:
        import openpyxl
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        path = self.root / "many_rows.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["x", "y"])
        for index in range(2000):
            sheet.append([index, index * 2])
        workbook.save(path)

        calls = []
        original_iter_rows = ReadOnlyWorksheet.iter_rows

        def counted_iter_rows(instance, *args, **kwargs):
            calls.append(instance)
            return original_iter_rows(instance, *args, **kwargs)

        with patch.object(ReadOnlyWorksheet, "iter_rows", counted_iter_rows):
            table = read_tables(path)[0]
        self.assertEqual(table.n_rows, 2000)
        self.assertEqual(len(calls), 2)  # one streaming pass for cached values and one for formulas

    def test_options_are_validated_instead_of_silently_ignored(self) -> None:
        path = self.write("simple.csv", "x,y\n1,2\n")
        with self.assertRaisesRegex(DataImportError, "未知的读取选项"):
            read_tables(path, guess=True)
        with self.assertRaisesRegex(DataImportError, "skip_rows"):
            read_tables(path, skip_rows=-1)
        with self.assertRaisesRegex(DataImportError, "delimiter"):
            read_tables(path, delimiter="pipe")


if __name__ == "__main__":
    unittest.main()
