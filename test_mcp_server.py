"""Optional end-to-end tests for the agent-facing MCP stdio server."""

from __future__ import annotations

import asyncio
import importlib.util
import os
import sys
import tempfile
import unittest
from pathlib import Path

try:
    from openpyxl import load_workbook, Workbook
except ImportError:  # pragma: no cover - covered by base requirements
    load_workbook = None  # type: ignore[assignment]
    Workbook = None  # type: ignore[assignment,misc]


MCP_AVAILABLE = importlib.util.find_spec("mcp") is not None


@unittest.skipUnless(MCP_AVAILABLE, "install requirements-agent.txt for MCP tests")
class McpStdioRoundTripTests(unittest.TestCase):
    def test_stdio_client_inspect_plan_validate_and_import_xlsx(self) -> None:
        self.assertIsNotNone(Workbook, "install requirements.txt for XLSX tests")
        from mcp import Client, StdioServerParameters

        repo_root = Path(__file__).resolve().parent

        async def run_round_trip(root: Path) -> None:
            source = root / "measurement.xlsx"
            output = root / "origin-ready.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Sweep"
            sheet.append(["Time (s)", "Voltage (V)", "State"])
            sheet.append([0, 1.25, "start"])
            sheet.append([1, 2.50, "end"])
            book.save(source)

            parameters = StdioServerParameters(
                command=sys.executable,
                args=["-m", "origin_bridge.mcp_server"],
                env={
                    "PATH": os.environ.get("PATH", ""),
                    "PYTHONPATH": str(repo_root),
                    "PYTHONIOENCODING": "utf-8",
                },
            )
            async with Client(parameters) as client:
                listed = await client.list_tools()
                names = {tool.name for tool in listed.tools}
                self.assertTrue(
                    {
                        "inspect_data",
                        "create_import_plan",
                        "validate_import_plan",
                        "execute_import_plan",
                    }.issubset(names)
                )

                inspected = await client.call_tool(
                    "inspect_data", {"paths": [str(source)]}
                )
                self.assertFalse(inspected.is_error, _tool_error(inspected))
                inspect_result = inspected.structured_content
                self.assertIsInstance(inspect_result, dict)
                self.assertTrue(inspect_result["tables"])

                planned = await client.call_tool(
                    "create_import_plan",
                    {
                        "paths": [str(source)],
                        "output_path": str(output),
                        "output_format": "xlsx",
                    },
                )
                self.assertFalse(planned.is_error, _tool_error(planned))
                plan = planned.structured_content
                self.assertIsInstance(plan, dict)
                self.assertFalse(output.exists(), "plan creation must not write output")

                validated = await client.call_tool(
                    "validate_import_plan", {"plan": plan}
                )
                self.assertFalse(validated.is_error, _tool_error(validated))
                self.assertTrue(validated.structured_content["valid"])
                self.assertFalse(output.exists(), "validation must not write output")

                declined = await client.call_tool(
                    "execute_import_plan", {"plan": plan, "confirm": False}
                )
                self.assertTrue(declined.is_error)
                self.assertFalse(output.exists(), "execution requires explicit confirm")

                exported = await client.call_tool(
                    "execute_import_plan", {"plan": plan, "confirm": True}
                )
                self.assertFalse(exported.is_error, _tool_error(exported))
                self.assertTrue(output.is_file())
                self.assertGreater(output.stat().st_size, 0)

            exported_book = load_workbook(output, data_only=True, read_only=True)
            try:
                self.assertIn("measurement_Sweep", exported_book.sheetnames)
                out_sheet = exported_book["measurement_Sweep"]
                self.assertEqual(out_sheet.cell(1, 1).value, "Time")
                self.assertEqual(out_sheet.cell(1, 2).value, "Voltage")
                self.assertEqual(out_sheet.cell(2, 2).value, 1.25)
                self.assertEqual(out_sheet.cell(3, 3).value, "end")
            finally:
                exported_book.close()

        with tempfile.TemporaryDirectory(prefix="mcp-stdio-import-") as temp:
            asyncio.run(run_round_trip(Path(temp)))


def _tool_error(result: object) -> str:
    content = getattr(result, "content", ())
    return "; ".join(str(getattr(item, "text", item)) for item in content)


if __name__ == "__main__":
    unittest.main()
